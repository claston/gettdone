"""Strict, offline Free characterization used by check_pdf_compatibility.py.

No production imports at module scope: each worker imports the selected revision.
Only the existing fictional corpus is accepted, never a directory of user PDFs.
"""
from __future__ import annotations

import csv
import hashlib
import json
import platform
import re
import sys
import unicodedata
from dataclasses import asdict
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from time import perf_counter, process_time

SCHEMA_VERSION = 1
FROZEN_BASELINE = Path(__file__).resolve().parents[1] / "backend/tests/fixtures/free_converter_baseline.json"
# Exhaustive dynamic-field catalogue, applied only to pdf_processing_metrics.
TIMING_FIELDS = {"total_ms", "parse_ms", "classify_ms", "normalize_ms", "reconcile_ms"}


def cents(value):
    if value is None:
        return None
    exact = Decimal(str(value)) * 100
    if not exact.is_finite() or exact != exact.to_integral_value():
        raise ValueError(f"Non-exact cent value: {value!r}")
    return int(exact)


def normalized_description(value):
    return " ".join(unicodedata.normalize("NFC", value).split())


def diff_values(before, after, path="$"):
    """Ordered, multiplicity-preserving differences with explicit old/new values."""
    if isinstance(before, dict) and isinstance(after, dict):
        differences = []
        for key in sorted(before.keys() | after.keys()):
            location = f"{path}.{key}"
            if key not in before or key not in after:
                differences.append({"path": location, "before": before.get(key), "after": after.get(key),
                                    "change": "added" if key not in before else "removed"})
            else:
                differences.extend(diff_values(before[key], after[key], location))
        return differences
    if isinstance(before, list) and isinstance(after, list):
        differences = []
        for index in range(max(len(before), len(after))):
            if index >= min(len(before), len(after)):
                differences.append({"path": f"{path}[{index}]", "before": before[index] if index < len(before) else None,
                                    "after": after[index] if index < len(after) else None,
                                    "change": "added" if index >= len(before) else "removed"})
            else:
                differences.extend(diff_values(before[index], after[index], f"{path}[{index}]"))
        return differences
    if before != after or type(before) is not type(after):
        return [{"path": path, "before": before, "after": after, "change": "changed"}]
    return []


def compare_snapshots(baseline, candidate, *, approvals=(), reference_commit=None, candidate_app_tree=None):
    errors, differences = [], []
    if not baseline or set(baseline) != set(candidate):
        errors.append("Missing/extra cases or empty baseline; comparison cannot pass")
    for key in sorted(baseline.keys() & candidate.keys()):
        old, new = baseline[key], candidate[key]
        if new.get("technical_failure") or old.get("technical_failure"):
            errors.append(f"{key}: technical failure is not a financial rejection")
        if "snapshot" not in old or "snapshot" not in new:
            errors.append(f"{key}: missing strict snapshot")
            continue
        changes = diff_values(old["snapshot"], new["snapshot"])
        unapproved = 0
        for change in changes:
            difference = {"case": key, **change}
            approval = next((item for item in approvals if item.get("difference") == difference
                             and reference_commit and item.get("reference_commit") == reference_commit
                             and candidate_app_tree and item.get("candidate_app_tree") == candidate_app_tree
                             and all(item.get(field) for field in ("approved_by", "approval_reference", "reason"))), None)
            if approval:
                difference["approval"] = {field: approval[field] for field in ("approved_by", "approval_reference", "reason")}
            else:
                unapproved += 1
            differences.append(difference)
        if unapproved:
            errors.append(f"{key}: {unapproved} unapproved characterization differences")
    return {"errors": errors, "differences": differences, "cases": len(candidate)}


def normalize_ofx(content, *, empty):
    # Empty statements use datetime.now() for DTASOF; nonempty dates are financial.
    if empty:
        content = re.sub(r"(<DTASOF>)\d{14}(\[-3:BRT\])", r"\1<generation-date>\2", content)
    return content


def snapshot_csv(content):
    # Ignore only empty physical records, including Windows CR/CRLF duplication.
    return [row for row in csv.reader(StringIO(content)) if row]


def snapshot_workbook(path):
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False)
    try:
        # Preserve sheet order, names, cell types, formulas and financial precision.
        return [{"name": sheet.title, "rows": [[
            {"type": cell.data_type, "value": str(Decimal(str(cell.value))) if cell.data_type == "n"
             and isinstance(cell.value, (int, float)) else cell.value}
            for cell in row] for row in sheet.iter_rows()]} for sheet in book]
    finally:
        book.close()


def snapshot_rows(rows):
    result = []
    for row in rows or []:
        item = asdict(row)
        item["description"] = normalized_description(item["description"])
        item["amount_cents"] = cents(item.pop("amount"))
        item["running_balance_cents"] = cents(item.pop("running_balance"))
        result.append(item)
    return result


def snapshot_data(data, parsed):
    metrics = {key: value for key, value in (data.pdf_processing_metrics or {}).items() if key not in TIMING_FIELDS}
    return {
        "status": "review" if metrics.get("export_recommendation") == "review_recommended" else "success",
        "transactions_total": data.transactions_total,
        "rows": snapshot_rows(data.report_transactions), "preview": snapshot_rows(data.preview_transactions),
        "opening_balance_cents": cents(data.opening_balance), "closing_balance_cents": cents(data.closing_balance),
        "inflows_cents": cents(data.total_inflows), "outflows_cents": cents(data.total_outflows), "net_cents": cents(data.net_total),
        "bank": {key: getattr(data, key) for key in ("bank_name", "bank_code", "bank_branch", "account_number", "ofx_account_type")},
        "document_type": data.document_type, "layout": data.layout_inference_name,
        "metrics": metrics, "quality_issues": data.quality_issues, "document_parts": data.document_parts,
        "financial_document": data.financial_document,
        "reconciliation": {"matched_groups": data.matched_groups, "reversed_entries": data.reversed_entries,
                           "potential_duplicates": data.potential_duplicates},
        "canonical": [{**asdict(row), "description": normalized_description(row.description),
                       "amount": cents(row.amount), "running_balance": cents(row.running_balance)}
                      for row in (getattr(parsed, "canonical_transactions", None) or [])],
    }


def extra_cases():
    """Reuse existing fixtures, with a small independently transcribed source oracle."""
    import base64

    from synthetic_pdf_corpus.catalog import load_scenarios
    from synthetic_pdf_corpus.native_text_pdf import generate_native_text_pdf
    from tests.test_document_preflight_service import _build_pdf_with_pages
    from tests.test_empty_statements import statement_lines
    from tests.test_pdf_text_extraction import _build_encrypted_pdf

    def case(key, raw, expected, **conditions):
        return {"id": key, "pdf": base64.b64encode(raw).decode(), "ocr": None, "expected": expected, **conditions}

    cases = [
        case("free_long_multiline_transfer", generate_native_text_pdf(((
            "EXTRATO CONTA CORRENTE", "PERIODO 01/07/2026 A 31/07/2026",
            "15/07/2026", "TRANSFERENCIA ENTRE CONTAS PROPRIAS", "CLIENTE FICTICIO ALFA E BETA CANAL INTERNET",
            "REFERENCIA OPERACIONAL DO DOCUMENTO FICTICIO XYZ", "1250,00 C", "8430,20",
            "16/07/2026", "PAGAMENTO FORNECEDOR BETA", "200,00 D", "8230,20",
        ),)), {"movements": [["2026-07-15", 1250], ["2026-07-16", -200]],
               "descriptions": ["TRANSFERENCIA ENTRE CONTAS PROPRIAS CLIENTE FICTICIO ALFA E BETA CANAL INTERNET "
                                "REFERENCIA OPERACIONAL DO DOCUMENTO FICTICIO XYZ", "PAGAMENTO FORNECEDOR BETA"],
               "balances": [718020, 823020], "export_oracle": True}),
        case("free_empty_statement", generate_native_text_pdf((statement_lines("Sem movimentação no período."),)),
             {"movements": [], "balances": [15000, 15000]}),
        case("free_invalid_pdf", b"%PDF-invalid-fictional", {"movements": [], "rejected": True}),
        case("free_password_protected", _build_encrypted_pdf(password="fictional-test-only"),
             {"movements": [], "rejected": True}),
        case("free_unknown_layout", generate_native_text_pdf((("Relatório fictício sem informação financeira recuperável",),)),
             {"movements": [], "rejected": True}),
        case("free_ocr_page_limit", _build_pdf_with_pages(11), {"movements": [], "rejected": True}, preflight=True),
        case("free_native_page_limit", _build_pdf_with_pages(251), {"movements": [], "rejected": True},
             preflight=True, scanned_likely=False),
        case("free_upload_size_limit", _build_pdf_with_pages(1), {"movements": [], "rejected": True},
             preflight=True, upload_size_bytes=10 * 1024 * 1024 + 1, scanned_likely=False),
    ]
    for scenario in load_scenarios(FROZEN_BASELINE.parent / "pdf_scenarios"):
        if not scenario.expected.enforce_success:
            continue  # Existing known gaps remain correctness tests, not immutable truths.
        cases.append(case("free_catalog_" + scenario.scenario_id, generate_native_text_pdf(scenario.pages),
                          {"movements": [[row.date, row.amount] for row in scenario.expected.transactions],
                           "descriptions": [row.description_contains for row in scenario.expected.transactions],
                           "export_oracle": True}))
    return cases


def source_checks(case, result):
    """Additional independent truth checks; never inferred from the frozen snapshot."""
    errors = []
    if result.get("technical_failure"):
        errors.append(f"{case['id']}: unexpected technical failure")
    expected = case["expected"]
    snapshot = result.get("snapshot", {})
    if "balances" in expected and [snapshot.get("opening_balance_cents"), snapshot.get("closing_balance_cents")] != expected["balances"]:
        errors.append(f"{case['id']}: balances disagree with independently transcribed source")
    if expected.get("descriptions"):
        actual = [row["description"] for row in snapshot.get("rows", [])]
        if len(actual) != len(expected["descriptions"]) or any(
            token not in description for token, description in zip(expected["descriptions"], actual)
        ):
            errors.append(f"{case['id']}: descriptions disagree with source oracle")
    if expected.get("export_oracle") and result.get("success"):
        expected_rows = [(day, cents(amount)) for day, amount in expected["movements"]]
        artifacts = snapshot["artifacts"]
        csv_rows = artifacts["converted.csv"]
        # These catalogue fixtures use the existing plain bank-statement exporters.
        actual_csv = [(row[0], cents(row[2])) for row in csv_rows[1:]]
        ofx = artifacts["converted.ofx"]
        actual_ofx = [(f"{day[:4]}-{day[4:6]}-{day[6:8]}", cents(amount)) for day, amount in
                      zip(re.findall(r"<DTPOSTED>(\d{8})", ofx), re.findall(r"<TRNAMT>([^\n<]+)", ofx), strict=True)]
        xlsx_rows = artifacts["converted.xlsx"][0]["rows"][1:]
        actual_xlsx = [("-".join(row[0]["value"].split("-")[::-1]),
                        cents(row[2]["value"] or 0) - cents(row[3]["value"] or 0)) for row in xlsx_rows]
        if any(rows != expected_rows for rows in (actual_csv, actual_ofx, actual_xlsx)):
            errors.append(f"{case['id']}: exported financial rows differ from independent source truth")
        fit_ids = re.findall(r"<FITID>([^\n<]+)", ofx)
        if len(fit_ids) != len(expected_rows) or len(set(fit_ids)) != len(fit_ids):
            errors.append(f"{case['id']}: missing or duplicate FITID")
    return errors


def peak_rss_bytes():
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD),
                        *[(name, ctypes.c_size_t) for name in ("PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage",
                          "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage",
                          "PagefileUsage", "PeakPagefileUsage")]]

        counters = Counters()
        counters.cb = ctypes.sizeof(counters)
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.GetCurrentProcess.restype = wintypes.HANDLE
        psapi = ctypes.WinDLL("psapi", use_last_error=True)
        psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
        if not psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
            raise ctypes.WinError(ctypes.get_last_error())
        return counters.PeakWorkingSetSize
    import resource

    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == "darwin" else 1024)


def percentile(values, fraction):
    ordered = sorted(values)
    rank = (len(ordered) - 1) * fraction
    low = int(rank)
    return ordered[low] + (ordered[min(low + 1, len(ordered) - 1)] - ordered[low]) * (rank - low)


def performance_report(results, *, repeat, profile):
    import importlib.metadata

    groups = {}
    for case in results.values():
        group = case["input"]["strategy"]
        if (case["input"].get("pages") or 0) >= 50 and group == "native":
            group = "native_large"
        groups.setdefault(group, []).extend(case["performance"])
    return {
        "environment": {"profile": profile, "platform": platform.platform(), "python": platform.python_version(),
                        "machine": platform.machine(), "processor": platform.processor(),
                        "dependencies": {name: importlib.metadata.version(name) for name in ("pypdf", "openpyxl", "pydantic")}},
        "repeat": repeat, "process_peak_rss_bytes": peak_rss_bytes(),
        "groups": {key: {"runs": len(runs), "p50_ms": percentile([r["total_ms"] for r in runs], .5),
                         "p95_ms": percentile([r["total_ms"] for r in runs], .95),
                         "cpu_p95_ms": percentile([r["cpu_ms"] for r in runs], .95),
                         "export_p95_ms": percentile([r["export_ms"] for r in runs], .95),
                         "technical_failures": sum(r["technical_failure"] for r in runs), "timeouts": 0,
                         "rejections": sum(r["rejected"] for r in runs),
                         "artifact_bytes": sum(r["artifact_bytes"] for r in runs)} for key, runs in groups.items()},
    }


def compare_performance(baseline, candidate):
    errors = []
    if baseline["environment"] != candidate["environment"] or min(baseline["repeat"], candidate["repeat"]) < 5:
        errors.append("Performance requires the same environment/profile and at least five repeats")
    if set(baseline["groups"]) != set(candidate["groups"]):
        errors.append("Performance groups differ")
    for group in baseline["groups"].keys() & candidate["groups"].keys():
        old, new = baseline["groups"][group], candidate["groups"][group]
        # Reject-only policy decisions take microseconds; relative time budgets
        # there are timer noise. Their outcome/failure gates still apply.
        for metric in (() if group == "preflight" else ("p95_ms", "cpu_p95_ms", "export_p95_ms")):
            if new[metric] > old[metric] * 1.10:
                errors.append(f"{group}.{metric}: exceeds +10% budget ({old[metric]} -> {new[metric]})")
        for metric in ("technical_failures", "timeouts"):
            if new[metric] > old[metric]:
                errors.append(f"{group}.{metric}: increased")
    if candidate["process_peak_rss_bytes"] > baseline["process_peak_rss_bytes"] * 1.10:
        errors.append("Worker peak RSS exceeds +10% budget")
    return {"errors": errors}


def worker(backend, corpus, output, *, repeat=1, profile="offline-local"):
    """Actual parsing -> pipeline -> filesystem persistence -> production exporters."""
    import base64
    import tempfile
    from types import SimpleNamespace

    sys.path.insert(0, str(backend.resolve()))
    from app.api.conversion.conversion_observability import _resolve_failed_conversion_code
    from app.application import pdf_parser
    from app.application.conversion.document_preflight_service import DocumentPreflightResult, DocumentPreflightService
    from app.application.default_conversion_pipeline import build_default_conversion_pipeline
    from app.application.document_extraction_models import ExtractedLine, ExtractedPage, ExtractedWord, RawDocumentExtraction
    from app.application.errors import FileTooLargeError, InvalidFileContentError, MaxPagesPerFileExceededError
    from app.application.storage_service import TempAnalysisStorage
    from pypdf import PdfReader

    results = {}
    for case in corpus:
        raw = base64.b64decode(case["pdf"])
        source = case.get("ocr")
        try:
            pages_count = len(PdfReader(BytesIO(raw)).pages)
        except Exception:
            pages_count = None
        entry = {"input": {"type": "ocr_evidence" if source else "native_or_vector_pdf", "pages": pages_count,
                           "strategy": "ocr_evidence" if source else "preflight" if case.get("preflight") else "native",
                           "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
                           "conditions": {k: v for k, v in case.items() if k not in {"pdf", "ocr"}}}, "performance": []}
        results[case["id"]] = entry
        nondeterministic = False
        for iteration in range(repeat):
            start, cpu_start = perf_counter(), process_time()
            export_ms, artifact_bytes = 0.0, 0
            try:
                if case.get("preflight"):
                    DocumentPreflightService().build_policy(
                        identity=SimpleNamespace(max_pages_per_file=10, max_pages_per_file_ocr=10,
                                                 max_upload_size_bytes=5 * 1024 * 1024),
                        filename="fictional.pdf", staged_upload_size_bytes=case.get("upload_size_bytes", len(raw)),
                        preflight_result=DocumentPreflightResult(scanned_likely=case.get("scanned_likely", True),
                                                               estimated_pages_count=pages_count))
                if source:
                    pages = [ExtractedPage(page_number=p["page_number"], lines=[ExtractedLine(**line) for line in p["lines"]],
                                          words=[ExtractedWord(**w) for w in p["words"]]) for p in source["pages"]]
                    extracted = RawDocumentExtraction(source["provider"], source["document_hash"], pages, {})
                    parsed = pdf_parser._parse_scanned_pdf_with_textract_text_pages(extracted)
                else:
                    parsed = pdf_parser.parse_pdf_transactions(raw)
                data = build_default_conversion_pipeline().run(filename="fictional.pdf", raw_bytes=raw, analysis_id="an_pdf_gate",
                                                               pdf_parser=lambda *a, **kw: parsed).analysis_data
                with tempfile.TemporaryDirectory(prefix="gettdone-free-artifacts-") as temporary:
                    storage = TempAnalysisStorage(Path(temporary), now_provider=lambda: datetime(2026, 1, 1, tzinfo=timezone.utc))
                    export_start = perf_counter()
                    storage.save_analysis(data)
                    export_ms = (perf_counter() - export_start) * 1000
                    elapsed_ms, cpu_ms = (perf_counter() - start) * 1000, (process_time() - cpu_start) * 1000
                    snapshot = snapshot_data(data, parsed)
                    directory = Path(temporary) / data.analysis_id
                    artifacts = {}
                    for path in sorted(directory.iterdir()):
                        artifact_bytes += path.stat().st_size
                        if path.suffix == ".xlsx":
                            artifacts[path.name] = snapshot_workbook(path)
                        elif path.suffix == ".ofx":
                            artifacts[path.name] = normalize_ofx(path.read_text(encoding="utf-8"), empty=data.transactions_total == 0)
                        elif path.suffix == ".csv":
                            artifacts[path.name] = snapshot_csv(path.read_text(encoding="utf-8"))
                    snapshot["artifacts"] = artifacts
                    persisted = json.loads((directory / "analysis.json").read_text(encoding="utf-8"))
                    snapshot["persisted_rows"] = persisted["report_transactions"]
                    snapshot["persisted_preview"] = persisted["preview_transactions"]
                doc = data.financial_document
                result = {"success": True, "movements": [[r.date, r.amount] for r in data.report_transactions],
                          "bank_code": data.bank_code, "document_type": data.document_type, "origin": doc and doc.get("origin"),
                          "formats": doc and doc["export"]["available_formats"], "review": snapshot["status"] == "review",
                          "snapshot": snapshot, "technical_failure": False}
            except (InvalidFileContentError, FileTooLargeError, MaxPagesPerFileExceededError) as exc:
                elapsed_ms, cpu_ms = (perf_counter() - start) * 1000, (process_time() - cpu_start) * 1000
                result = {"success": False, "error": type(exc).__name__, "technical_failure": False,
                          "snapshot": {"status": "rejected", "error": type(exc).__name__, "code": _resolve_failed_conversion_code(exc)}}
            except Exception as exc:
                elapsed_ms, cpu_ms = (perf_counter() - start) * 1000, (process_time() - cpu_start) * 1000
                result = {"success": False, "error": type(exc).__name__, "technical_failure": True,
                          "snapshot": {"status": "technical_failure", "error": type(exc).__name__}}
            if iteration and entry["snapshot"] != result["snapshot"]:
                nondeterministic = True
            result["technical_failure"] = result["technical_failure"] or nondeterministic
            entry["performance"].append({"total_ms": elapsed_ms,
                                         "cpu_ms": cpu_ms, "export_ms": export_ms,
                                         "artifact_bytes": artifact_bytes, "technical_failure": result["technical_failure"],
                                         "rejected": not result["success"]})
            entry.update(result)
    report = {"schema_version": SCHEMA_VERSION, "cases": results,
              "performance": performance_report(results, repeat=repeat, profile=profile)}
    output.write_text(json.dumps(report, ensure_ascii=True, indent=2), encoding="utf-8")
