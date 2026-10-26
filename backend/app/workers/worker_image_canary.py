"""Opt-in, fixed synthetic checks for an isolated worker image, never a queue consumer."""

from __future__ import annotations

import os
import tempfile
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from app.workers.worker_image_contract import verify_worker_image


def lambda_handler(event, context):
    if os.environ.get("WORKER_IMAGE_CANARY") != "1" or event != {"contract": "worker-image-v1"}:
        raise ValueError("Only the explicitly enabled synthetic canary is accepted.")
    verify_worker_image()
    return run_checks()


def _pdf_fixture() -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    lines = [
        "ref:synthetic-wise 1 / 1",
        "Wise Brasil Instituição de Pagamento Ltda.",
        "Extrato em BRL",
        "1 de abril de 2026 [GMT-03:00] - 30 de abril de 2026 [GMT-03:00]",
        "Gerado em: 13 de outubro de 2026",
        "Titular da Conta",
        "EMPRESA SINTETICA LTDA",
        "Número da conta",
        "12345678",
        "BRL em 30 de abril de 2026 [GMT-03:00] 0,00 BRL",
        "Descrição Entrada Saída Valor",
        "Obtenção dos dados bancários da Wise -40,00 0,00",
        "7 de abril de 2026 | Transação: BANK_DETAILS_ORDER_CHECKOUT-invoice-100002",
        "Dinheiro adicionado à conta 40,00 40,00",
        "7 de abril de 2026 | Transação: TRANSFER-100001",
        "A Wise Brasil Pagamentos LTDA é uma Instituição de Pagamentos.",
        "Precisa de ajuda? Visite wise.com/help",
    ]
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    }))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    commands = ["BT /F1 10 Tf 48 790 Td 14 TL"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj T*")
    commands.append("ET")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("cp1252"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def check_postgres(url: str) -> str:
    parsed = urlsplit(url)
    if (parsed.scheme not in {"postgresql", "postgres"} or parsed.hostname not in {"localhost", "127.0.0.1"}
            or parsed.username != "worker_canary" or parsed.path != "/worker_canary" or parsed.query or parsed.fragment):
        raise ValueError("Canary only accepts a local synthetic database.")
    import psycopg

    with psycopg.connect(url, connect_timeout=5) as connection:
        with connection.cursor() as cursor:
            cursor.execute("CREATE TEMP TABLE worker_image_canary (amount numeric NOT NULL)")
            cursor.execute("INSERT INTO worker_image_canary VALUES (40), (-40)")
            cursor.execute("SELECT count(*), sum(amount) FROM worker_image_canary")
            if cursor.fetchone() != (2, 0):
                raise RuntimeError("Synthetic PostgreSQL round trip failed.")
    return "passed"


def run_checks() -> dict:
    import psycopg
    import pypdfium2
    from openpyxl import Workbook, load_workbook

    from app.application.default_conversion_pipeline import build_default_conversion_pipeline
    from app.application.export_artifact_service import ExportArtifactService

    xlsx = BytesIO()
    workbook = Workbook()
    workbook.active.append(["date", "description", "amount"])
    workbook.active.append(["2026-04-07", "SYNTHETIC CREDIT", 40])
    workbook.active.append(["2026-04-07", "SYNTHETIC DEBIT", -40])
    workbook.save(xlsx)
    workbook.close()
    ofx = b'''OFXHEADER:100
DATA:OFXSGML
VERSION:102
SECURITY:NONE
ENCODING:USASCII
CHARSET:1252
COMPRESSION:NONE
OLDFILEUID:NONE
NEWFILEUID:NONE

<OFX><BANKMSGSRSV1><STMTTRNRS><TRNUID>1<STATUS><CODE>0<SEVERITY>INFO</STATUS><STMTRS>
<CURDEF>BRL<BANKACCTFROM><BANKID>001<ACCTID>SYNTHETIC<ACCTTYPE>CHECKING</BANKACCTFROM>
<BANKTRANLIST><DTSTART>20260407000000<DTEND>20260407000000
<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20260407000000<TRNAMT>40.00<FITID>canary1<MEMO>SYNTHETIC CREDIT</STMTTRN>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260407000000<TRNAMT>-40.00<FITID>canary2<MEMO>SYNTHETIC DEBIT</STMTTRN>
</BANKTRANLIST><LEDGERBAL><BALAMT>0.00<DTASOF>20260407000000</LEDGERBAL>
</STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>'''
    pdf = _pdf_fixture()
    fixtures = {
        "csv": b"date,description,amount\n2026-04-07,SYNTHETIC CREDIT,40.00\n2026-04-07,SYNTHETIC DEBIT,-40.00\n",
        "xlsx": xlsx.getvalue(), "ofx": ofx, "pdf": pdf,
    }
    checks = []
    with tempfile.TemporaryDirectory(prefix="worker-canary-") as directory:
        for suffix, raw in fixtures.items():
            analysis = build_default_conversion_pipeline().run(
                filename=f"synthetic.{suffix}", raw_bytes=raw, analysis_id=f"an_canary_{suffix}",
            ).analysis_data
            amounts = sorted(row.amount for row in analysis.report_transactions)
            if analysis.transactions_total != 2 or amounts != [-40, 40]:
                raise RuntimeError(f"{suffix}: incorrect synthetic transactions: {amounts}")
            target = Path(directory) / suffix
            target.mkdir()
            ExportArtifactService().write_convert_artifacts(
                target, report_rows=analysis.report_transactions, layout_inference_name=analysis.layout_inference_name,
                opening_balance=analysis.opening_balance, closing_balance=analysis.closing_balance,
            )
            exported = (target / "converted.ofx").read_text(encoding="utf-8")
            if exported.count("<STMTTRN>") != 2 or "<TRNAMT>40.00" not in exported or "<TRNAMT>-40.00" not in exported:
                raise RuntimeError(f"{suffix}: incorrect synthetic OFX export")
            # Also read the generated spreadsheet through its native compression/XML dependencies.
            with (target / "converted.xlsx").open("rb") as source:
                exported_book = load_workbook(source, read_only=True)
                if exported_book.active.max_row < 3:
                    raise RuntimeError(f"{suffix}: incomplete spreadsheet export")
                exported_book.close()
            checks.append(suffix)
    with pypdfium2.PdfDocument(pdf) as document:
        page = document[0]
        try:
            bitmap = page.render(scale=0.5)
            try:
                if bitmap.width <= 0 or bitmap.height <= 0:
                    raise RuntimeError("PDFium produced an empty bitmap.")
            finally:
                bitmap.close()
        finally:
            page.close()
    checks.append("pdfium")
    if psycopg.pq.__impl__ != "binary" or psycopg.pq.version() <= 0:
        raise RuntimeError("The binary PostgreSQL client is unavailable.")
    checks.append("psycopg")
    database_url = os.environ.get("WORKER_CANARY_DATABASE_URL")
    return {"status": "passed", "checks": checks, "postgres": check_postgres(database_url) if database_url else "not_configured"}
