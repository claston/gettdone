import hashlib
import json
import os
import re
from dataclasses import asdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Callable
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.application.analysis_lock import locked_analysis
from app.application.balance_checks import parse_balance_check
from app.application.bank_identity import resolve_conversion_model_label
from app.application.errors import AnalysisAccessDeniedError, AnalysisEditConflictError, AnalysisNotFoundError, InvalidFileContentError
from app.application.export_artifact_service import ExportArtifactService
from app.application.financial_document_export import assert_financial_export
from app.application.models import AnalysisData, BeforeAfterRow, TransactionRow
from app.application.source_regions import parse_source_regions, validate_review_regions
from app.application.statement_account_export import assert_account_overrides, refresh_edited_account_rows, statement_scopes
from app.application.structured_conversion import build_structured_conversion_result_from_analysis_data

_SAFE_ANALYSIS_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,127}$")


def _is_confirmed_empty_statement(content: dict[str, object]) -> bool:
    metrics = content.get("pdf_processing_metrics")
    return (
        content.get("transactions_total") == 0
        and isinstance(metrics, dict)
        and metrics.get("statement_no_activity") == 1
        and content.get("preview_transactions") == []
        and content.get("report_transactions") == []
    )


class TempAnalysisStorage:
    def __init__(
        self,
        root_dir: Path,
        ttl_seconds: int = 24 * 60 * 60,
        now_provider: Callable[[], datetime] | None = None,
        export_artifact_service: ExportArtifactService | None = None,
    ) -> None:
        self.root_dir = root_dir
        self.ttl_seconds = ttl_seconds
        self.now_provider = now_provider or (lambda: datetime.now(timezone.utc))
        self.export_artifact_service = export_artifact_service or ExportArtifactService()
        self.root_dir.mkdir(parents=True, exist_ok=True)

    def save_analysis(self, data: AnalysisData) -> str:
        analysis_dir = self._resolve_analysis_dir(data.analysis_id)
        analysis_dir.mkdir(parents=True, exist_ok=True)
        now = self.now_provider()
        expires_at = now + timedelta(seconds=self.ttl_seconds)
        report_rows = data.report_transactions or data.preview_transactions
        updated_at = data.updated_at or now.isoformat()
        data.updated_at = updated_at
        data.structured_result = build_structured_conversion_result_from_analysis_data(
            data,
            expires_at=expires_at.isoformat(),
        )

        json_path = analysis_dir / "analysis.json"
        json_path.write_text(
            json.dumps(
                {
                    **asdict(data),
                    "created_at": now.isoformat(),
                    "updated_at": updated_at,
                    "expires_at": expires_at.isoformat(),
                    "preview_transactions": [asdict(item) for item in data.preview_transactions],
                    "report_transactions": [asdict(item) for item in report_rows],
                },
                ensure_ascii=True,
                indent=2,
            ),
            encoding="utf-8",
        )

        self.export_artifact_service.write_analysis_report_workbook(
            analysis_dir / "report.xlsx",
            report_rows=report_rows,
            snapshot=data,
        )
        preserve_bank_metadata = data.file_type == "ofx" or (data.file_type == "pdf" and data.financial_document is None)
        self.export_artifact_service.write_convert_artifacts(
            analysis_dir,
            report_rows=report_rows,
            ofx_account_type=data.ofx_account_type,
            layout_inference_name=data.layout_inference_name,
            opening_balance=data.opening_balance,
            closing_balance=data.closing_balance,
            balance_bounds_unavailable=bool((data.pdf_processing_metrics or {}).get("balance_bounds_unavailable")),
            bank_branch=data.bank_branch if preserve_bank_metadata else None,
            account_number=data.account_number if preserve_bank_metadata else None,
            bank_code=data.bank_code if preserve_bank_metadata else None,
            **({"financial_document": data.financial_document} if data.financial_document is not None else {}),
            **({"statement_accounts": statement_scopes(data.document_parts)} if statement_scopes(data.document_parts) else {}),
        )
        return expires_at.isoformat()

    def get_report_path(self, analysis_id: str) -> Path:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        report_path = analysis_dir / "report.xlsx"
        if not report_path.exists():
            raise AnalysisNotFoundError
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError
        return report_path

    def get_convert_report_path(
        self,
        analysis_id: str,
        file_format: str,
        *,
        closing_balance: float | None = None,
        bank_branch: str | None = None,
        account_number: str | None = None,
        bank_code: str | None = None,
    ) -> Path:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        content, _ = self._read_analysis_metadata(analysis_id=analysis_id, metadata_filename="analysis.json")
        financial_document = content.get("financial_document")
        if financial_document is not None:
            persisted = self._extract_convert_overrides(content)
            assert_financial_export(financial_document, file_format,
                bank_code=bank_code or persisted.get("bank_code"), bank_branch=bank_branch or persisted.get("bank_branch"),
                account_number=account_number or persisted.get("account_number"))
        if file_format == "ofx" and (
            closing_balance is not None
            or (str(bank_branch or "").strip() != "")
            or (str(account_number or "").strip() != "")
            or (str(bank_code or "").strip() != "")
        ):
            self._regenerate_ofx_with_overrides(
                analysis_dir,
                closing_balance=closing_balance,
                bank_branch=bank_branch,
                account_number=account_number,
                bank_code=bank_code,
            )
        if file_format == "ofx":
            suffix = "ofx"
        elif file_format == "xlsx":
            suffix = "xlsx"
        else:
            suffix = "csv"
        report_path = analysis_dir / f"converted.{suffix}"
        if not report_path.exists():
            raise AnalysisNotFoundError
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError
        return report_path

    def get_upload_filename(self, analysis_id: str) -> str | None:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError

        metadata_path = analysis_dir / "analysis.json"
        if not metadata_path.exists():
            raise AnalysisNotFoundError
        try:
            content = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None

        raw_name = content.get("upload_filename")
        if not isinstance(raw_name, str):
            return None
        name = raw_name.strip()
        return name or None

    def set_report_owner(self, analysis_id: str, identity_type: str, identity_id: str) -> None:
        self._set_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="analysis.json",
            identity_type=identity_type,
            identity_id=identity_id,
        )

    def assert_report_owner(
        self,
        analysis_id: str,
        identity_type: str,
        identity_id: str,
        *,
        allow_unowned: bool = False,
    ) -> None:
        self._assert_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="analysis.json",
            identity_type=identity_type,
            identity_id=identity_id,
            allow_unowned=allow_unowned,
        )

    def set_convert_owner(self, analysis_id: str, identity_type: str, identity_id: str) -> None:
        self._set_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="analysis.json",
            identity_type=identity_type,
            identity_id=identity_id,
        )

    def assert_convert_owner(self, analysis_id: str, identity_type: str, identity_id: str) -> None:
        self._assert_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="analysis.json",
            identity_type=identity_type,
            identity_id=identity_id,
            allow_unowned=True,
        )

    def set_reconcile_owner(self, analysis_id: str, identity_type: str, identity_id: str) -> None:
        self._set_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="reconcile.json",
            identity_type=identity_type,
            identity_id=identity_id,
        )

    def assert_reconcile_owner(
        self,
        analysis_id: str,
        identity_type: str,
        identity_id: str,
        *,
        allow_unowned: bool = False,
    ) -> None:
        self._assert_analysis_owner(
            analysis_id=analysis_id,
            metadata_filename="reconcile.json",
            identity_type=identity_type,
            identity_id=identity_id,
            allow_unowned=allow_unowned,
        )

    def list_convert_history(self, identity_type: str, identity_id: str, limit: int = 20) -> list[dict[str, str]]:
        items: list[dict[str, str]] = []
        try:
            analysis_dirs = list(self.root_dir.iterdir())
        except OSError:
            analysis_dirs = []

        for analysis_dir in analysis_dirs:
            if not analysis_dir.is_dir():
                continue
            if self._is_expired(analysis_dir):
                try:
                    self._cleanup_analysis(analysis_dir)
                except OSError:
                    pass
                continue

            metadata_path = analysis_dir / "analysis.json"
            if not metadata_path.exists():
                continue

            try:
                content = json.loads(metadata_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue

            owner_type = str(content.get("owner_identity_type") or "").strip()
            owner_id = str(content.get("owner_identity_id") or "").strip()
            if owner_type != identity_type or owner_id != identity_id:
                continue

            processing_id = str(content.get("analysis_id") or analysis_dir.name).strip()
            created_at = str(content.get("created_at") or content.get("updated_at") or "").strip()
            filename = str(content.get("upload_filename") or "").strip() or f"{processing_id}.pdf"
            model = resolve_conversion_model_label(
                layout_inference_name=str(content.get("layout_inference_name") or "").strip() or None,
                bank_name=str(content.get("bank_name") or "").strip() or None,
            )
            file_type = str(content.get("file_type") or "").strip().lower()
            conversion_type = f"{file_type}-ofx" if file_type else "pdf-ofx"
            status = "Sucesso" if (analysis_dir / "converted.ofx").exists() else "Processando"
            if content.get("financial_document") is not None:
                conversion_type = f"{file_type}-xlsx"
                status = "Sucesso" if (analysis_dir / "converted.xlsx").exists() else "Processando"
                model = {"receivables_report": "Títulos recebidos", "accounts_payable_report": "Contas a pagar",
                         "accounts_receivable_report": "Contas a receber", "investment_statement": "Investimentos CDB",
                         "fiscal_summary": "Resumo fiscal"}.get(str(content.get("document_type")), "Relatório financeiro")

            items.append(
                {
                    "processing_id": processing_id,
                    "created_at": created_at,
                    "filename": filename,
                    "model": model,
                    "conversion_type": conversion_type,
                    "status": status,
                }
            )

        items.sort(key=lambda item: item["created_at"], reverse=True)
        return items[: max(1, min(limit, 100))]

    @locked_analysis
    def get_convert_edit_operation(self, analysis_id: str, operation_id: str) -> dict[str, object] | None:
        content, _ = self._read_analysis_metadata(analysis_id=analysis_id, metadata_filename="analysis.json")
        receipts = content.get("edit_operations") or {}
        receipt = receipts.get(operation_id)
        if receipt is None:
            return None
        if receipt["updated_at"] != content.get("updated_at"):
            raise AnalysisEditConflictError
        return self._edit_result(analysis_id, content, operation_id)

    @staticmethod
    def _edit_result(analysis_id: str, content: dict, operation_id: str | None = None) -> dict[str, object]:
        return {
            "processing_id": analysis_id,
            "operation_id": operation_id,
            **{key: content.get(key) for key in (
                "transactions_total", "total_inflows", "total_outflows", "net_total",
                "preview_transactions", "updated_at", "opening_balance", "closing_balance",
                "statement_balance_check",
            )},
        }

    @staticmethod
    def _write_metadata(path: Path, content: dict) -> None:
        # Readers must never interpret a partially written JSON document as a 404.
        with NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent, suffix=".tmp", delete=False) as target:
            temporary = Path(target.name)
            json.dump(content, target, ensure_ascii=True, indent=2)
        try:
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)

    @locked_analysis
    def apply_convert_edits(
        self,
        analysis_id: str,
        edits: list[dict[str, object]],
        expected_updated_at: str | None = None,
        opening_balance: float | None = None,
        closing_balance: float | None = None,
        bank_branch: str | None = None,
        account_number: str | None = None,
        bank_code: str | None = None,
        operation_id: str | None = None,
    ) -> dict[str, object]:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError

        metadata_path = analysis_dir / "analysis.json"
        if not metadata_path.exists():
            raise AnalysisNotFoundError

        try:
            content = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raise AnalysisNotFoundError from None

        if content.get("financial_document") is not None:
            raise InvalidFileContentError(
                "A edição de relatórios especializados ainda não está disponível. "
                "Os campos extraídos estão preservados no CSV e no Excel."
            )

        operation_hash = hashlib.sha256(json.dumps({
            "edits": edits, "expected_updated_at": expected_updated_at,
            "opening_balance": opening_balance, "closing_balance": closing_balance,
            "bank_branch": bank_branch, "account_number": account_number, "bank_code": bank_code,
        }, sort_keys=True).encode()).hexdigest()
        receipts = content.get("edit_operations") or {}
        if operation_id and operation_id in receipts:
            if receipts[operation_id]["hash"] != operation_hash:
                raise AnalysisEditConflictError
            return self.get_convert_edit_operation(analysis_id, operation_id)

        preview_rows = self._parse_transaction_rows(content.get("preview_transactions", []))
        report_rows_raw = content.get("report_transactions") or content.get("preview_transactions", [])
        report_rows = self._parse_transaction_rows(report_rows_raw)
        current_updated_at = str(content.get("updated_at") or "").strip()
        assert_account_overrides(report_rows, opening_balance=opening_balance, closing_balance=closing_balance,
                                 bank_branch=bank_branch, account_number=account_number, bank_code=bank_code,
                                 statement_accounts=statement_scopes(content.get("document_parts")))

        if (len(preview_rows) == 0 or len(report_rows) == 0) and not _is_confirmed_empty_statement(content):
            raise AnalysisNotFoundError
        if expected_updated_at and current_updated_at and expected_updated_at != current_updated_at:
            raise AnalysisEditConflictError

        edited_account_blocks = set()
        for edit in edits:
            action = str(edit.get("action") or "update").strip().lower()
            if action not in {"update", "delete", "restore", "insert", "annotate"}:
                raise ValueError("Invalid edit action.")
            review_regions = None
            if "review_regions" in edit:
                if content.get("file_type") != "pdf":
                    raise ValueError("Marcações estão disponíveis apenas para PDFs.")
                review_regions = validate_review_regions(edit["review_regions"],
                    (content.get("pdf_processing_metrics") or {}).get("page_count"))

            if action == "insert":
                if statement_scopes(content.get("document_parts")) and not preview_rows:
                    raise InvalidFileContentError("A inserção exige um lançamento de referência que identifique a conta de origem.")
                amount = self._build_amount_from_credit_debit(edit.get("credit"), edit.get("debit"))
                description = str(edit.get("description") or "").strip()
                date = str(edit.get("date") or "").strip()
                if not description or not date:
                    raise ValueError("Date and description are required.")
                insert_position = self._resolve_insert_position(edit.get("insert_position"), max_index=len(preview_rows))
                new_row = TransactionRow(
                    date=date,
                    description=description,
                    amount=amount,
                    category="Outros",
                    reconciliation_status="unmatched",
                    is_deleted=False,
                    review_regions=review_regions or [],
                )
                if all(row.source_sequence is not None for row in preview_rows):
                    anchor_id = edit.get("anchor_transaction_id")
                    anchor = next((row for row in preview_rows if row.transaction_id == anchor_id), None) if anchor_id else None
                    if anchor_id and anchor is None:
                        raise ValueError("Anchor transaction not found.")
                    if any(row.statement_account for row in preview_rows):
                        if anchor is None or not anchor.statement_account:
                            raise InvalidFileContentError("Escolha um lançamento da conta de origem para inserir a nova linha.")
                        new_row.statement_account = dict(anchor.statement_account)
                        edited_account_blocks.add(anchor.statement_account["block_id"])
                    new_row.transaction_id = str(edit.get("transaction_id") or f"txn_{uuid4().hex}")
                    if any(row.transaction_id == new_row.transaction_id for row in preview_rows):
                        raise ValueError("Transaction ID already exists.")
                    new_row.source_sequence = anchor.source_sequence if anchor else 0
                    new_row.source_subsequence = 1 + max(
                        (row.source_subsequence for row in preview_rows if row.source_sequence == new_row.source_sequence), default=0,
                    )
                preview_rows.insert(insert_position, new_row)
                report_rows.insert(insert_position, TransactionRow(**asdict(new_row)))
                continue

            row_id_raw = edit.get("row_id")
            stable_index = next((idx for idx, row in enumerate(preview_rows)
                                 if row.transaction_id is not None and row.transaction_id == row_id_raw), None)
            row_index = stable_index if stable_index is not None else self._resolve_row_index(row_id_raw, max_index=len(preview_rows))

            if review_regions is not None:
                preview_rows[row_index].review_regions = list(review_regions)
                report_rows[row_index].review_regions = list(review_regions)
            if action == "annotate":
                if review_regions is None:
                    raise ValueError("Informe as marcações do lançamento.")
                if preview_rows[row_index].is_deleted:
                    raise ValueError("Restaure o lançamento antes de marcar.")
                continue

            if action == "delete":
                if preview_rows[row_index].statement_account:
                    edited_account_blocks.add(preview_rows[row_index].statement_account["block_id"])
                if row_index >= len(report_rows):
                    raise ValueError("row_id out of bounds.")
                preview_rows[row_index].is_deleted = True
                report_rows[row_index].is_deleted = True
                continue

            # Zero-valued entries can be present in imported statements. Editing
            # their text/date must preserve the amount when both cells are empty.
            preserve_zero = (
                preview_rows[row_index].amount == 0
                and edit.get("credit") is None
                and edit.get("debit") is None
            )
            amount = 0.0 if preserve_zero else self._build_amount_from_credit_debit(edit.get("credit"), edit.get("debit"))
            description = str(edit.get("description") or "").strip()
            date = str(edit.get("date") or "").strip()
            if not description or not date:
                raise ValueError("Date and description are required.")

            report_target = report_rows[row_index]
            if report_target.statement_account:
                edited_account_blocks.add(report_target.statement_account["block_id"])
            report_target.date = date
            report_target.description = description
            report_target.amount = amount
            report_target.is_deleted = False

            preview_target = preview_rows[row_index]
            preview_target.date = date
            preview_target.description = description
            preview_target.amount = amount
            preview_target.is_deleted = False

        if edits and all(str(edit.get("action") or "").strip().lower() == "annotate" for edit in edits) and all(value is None for value in (
                opening_balance, closing_balance, bank_branch, account_number, bank_code)):
            # Reviewer geometry must never recalculate balances or regenerate
            # financial artifacts. Versioning/receipts still protect concurrent edits.
            content["preview_transactions"] = [asdict(item) for item in preview_rows]
            content["report_transactions"] = [asdict(item) for item in report_rows]
            content["updated_at"] = self.now_provider().isoformat()
            if operation_id:
                receipts[operation_id] = {"hash": operation_hash, "updated_at": content["updated_at"]}
                content["edit_operations"] = dict(list(receipts.items())[-64:])
            self._write_metadata(metadata_path, content)
            return self._edit_result(analysis_id, content, operation_id)

        requested_overrides = self._normalize_convert_overrides(
            opening_balance=opening_balance,
            closing_balance=closing_balance,
            bank_branch=bank_branch,
            account_number=account_number,
            bank_code=bank_code,
        )
        persisted_overrides = self._extract_convert_overrides(content)
        opening_balance_seed = self._resolve_recalculation_opening_balance(
            original_rows=report_rows_raw,
            current_rows=report_rows,
            content=content,
            requested_opening_balance=requested_overrides.get("opening_balance"),
            persisted_opening_balance=persisted_overrides.get("opening_balance"),
        )
        if any(row.statement_account for row in report_rows):
            opening_balance_seed = None
            if edited_account_blocks:
                refresh_edited_account_rows(preview_rows, report_rows, edited_account_blocks)
                content["pdf_processing_metrics"].update(financial_controls_validated=0, export_recommendation="review_recommended",
                    export_recommendation_reason="statement_account_edited", balance_validation_status="review_required")
        if opening_balance_seed is not None or not (content.get("pdf_processing_metrics") or {}).get("balance_bounds_unavailable"):
            self._recalculate_running_balances(
                preview_rows=preview_rows,
                report_rows=report_rows,
                opening_balance=opening_balance_seed,
            )

        active_rows = self._active_rows(report_rows)
        total_inflows = round(sum(item.amount for item in active_rows if item.amount > 0), 2)
        total_outflows = round(sum(item.amount for item in active_rows if item.amount < 0), 2)
        net_total = round(total_inflows + total_outflows, 2)
        effective_overrides = {
            "opening_balance": (
                requested_overrides.get("opening_balance")
                if requested_overrides.get("opening_balance") is not None
                else persisted_overrides.get("opening_balance")
            ),
            "closing_balance": (
                requested_overrides.get("closing_balance")
                if requested_overrides.get("closing_balance") is not None
                else persisted_overrides.get("closing_balance")
            ),
            "bank_branch": (
                requested_overrides.get("bank_branch")
                if requested_overrides.get("bank_branch") is not None
                else persisted_overrides.get("bank_branch")
            ),
            "account_number": (
                requested_overrides.get("account_number")
                if requested_overrides.get("account_number") is not None
                else persisted_overrides.get("account_number")
            ),
            "bank_code": (
                requested_overrides.get("bank_code")
                if requested_overrides.get("bank_code") is not None
                else persisted_overrides.get("bank_code")
            ),
        }
        if requested_overrides:
            self._apply_convert_overrides_to_content(content, requested_overrides)

        content["transactions_total"] = len(active_rows)
        content["total_inflows"] = total_inflows
        content["total_outflows"] = total_outflows
        content["net_total"] = net_total
        opening_balance_value, closing_balance_value = self._resolve_balance_bounds(
            active_rows,
            opening_balance_hint=effective_overrides.get("opening_balance"),
        )
        if (content.get("pdf_processing_metrics") or {}).get("balance_bounds_unavailable"):
            opening_balance_value = self._coerce_balance_value(effective_overrides.get("opening_balance"))
            closing_balance_value = self._coerce_balance_value(effective_overrides.get("closing_balance"))
        content["opening_balance"] = opening_balance_value
        content["closing_balance"] = closing_balance_value
        content["preview_transactions"] = [asdict(item) for item in preview_rows]
        content["report_transactions"] = [asdict(item) for item in report_rows]
        if "ocr_balance_arithmetic_matches" in (content.get("pdf_processing_metrics") or {}):
            content["pdf_processing_metrics"]["ocr_balance_arithmetic_matches"] = 0
        new_updated_at = self.now_provider().isoformat()
        content["updated_at"] = new_updated_at
        self._refresh_structured_result_content(content, preview_rows=preview_rows, report_rows=report_rows)

        self._write_report_workbook(analysis_dir, content=content, report_rows=report_rows, preview_rows=preview_rows)
        self.export_artifact_service.write_convert_artifacts(
            analysis_dir,
            report_rows=report_rows,
            ofx_account_type=str(content.get("ofx_account_type") or "").strip() or None,
            layout_inference_name=str(content.get("layout_inference_name") or "").strip() or None,
            opening_balance=opening_balance_value,
            closing_balance=effective_overrides.get("closing_balance"),
            balance_bounds_unavailable=bool((content.get("pdf_processing_metrics") or {}).get("balance_bounds_unavailable")),
            bank_branch=effective_overrides.get("bank_branch"),
            account_number=effective_overrides.get("account_number"),
            bank_code=effective_overrides.get("bank_code"),
            **({"financial_document": content["financial_document"]} if content.get("financial_document") is not None else {}),
            **({"statement_accounts": statement_scopes(content.get("document_parts"))}
               if statement_scopes(content.get("document_parts")) else {}),
        )

        if operation_id:
            receipts[operation_id] = {"hash": operation_hash, "updated_at": new_updated_at}
            content["edit_operations"] = dict(list(receipts.items())[-64:])
        self._write_metadata(metadata_path, content)
        return self._edit_result(analysis_id, content, operation_id)

    def save_reconcile_report(
        self,
        summary: dict[str, int],
        reconciliation_rows: list[dict[str, str | float | None]],
        problems: list[dict[str, str]],
    ) -> tuple[str, str]:
        analysis_id = self._build_analysis_id(prefix="rc")
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        analysis_dir.mkdir(parents=True, exist_ok=True)
        now = self.now_provider()
        expires_at = now + timedelta(seconds=self.ttl_seconds)

        json_path = analysis_dir / "reconcile.json"
        json_path.write_text(
            json.dumps(
                {
                    "analysis_id": analysis_id,
                    "created_at": now.isoformat(),
                    "expires_at": expires_at.isoformat(),
                    "summary": summary,
                    "reconciliation_rows": reconciliation_rows,
                    "problems": problems,
                },
                ensure_ascii=True,
                indent=2,
            ),
            encoding="utf-8",
        )

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Resumo"
        summary_sheet.append(["Metrica", "Valor"])
        summary_sheet.append(["Total linhas extrato", summary.get("total_bank_rows", 0)])
        summary_sheet.append(["Total linhas planilha", summary.get("total_sheet_rows", 0)])
        summary_sheet.append(["Conciliados", summary.get("conciliated_count", 0)])
        summary_sheet.append(["Pendentes", summary.get("pending_count", 0)])
        summary_sheet.append(["Divergentes", summary.get("divergent_count", 0)])
        self._format_transacoes_sheet(summary_sheet)

        detail_sheet = workbook.create_sheet(title="Conciliacao_Detalhada")
        detail_headers = [
            "Linha",
            "Fonte",
            "Data",
            "Descricao",
            "Valor",
            "Status",
            "Regra de match",
            "Motivo",
            "Linha pareada",
        ]
        detail_sheet.append(detail_headers)
        for row in reconciliation_rows:
            detail_sheet.append(
                [
                    row.get("row_id"),
                    self._translate_source_label(row.get("source")),
                    row.get("date"),
                    row.get("description"),
                    row.get("amount"),
                    self._translate_status_label(row.get("status")),
                    self._translate_match_rule_label(row.get("match_rule")),
                    self._translate_reason_label(row.get("reason")),
                    row.get("matched_row_id"),
                ]
            )
        self._format_transacoes_sheet(detail_sheet)

        problems_sheet = workbook.create_sheet(title="Problemas")
        problem_row_headers = [
            "Linha",
            "Fonte",
            "Data",
            "Descricao",
            "Valor",
            "Status",
            "Motivo",
            "Linha pareada",
        ]
        problems_sheet.append(problem_row_headers)
        problematic_rows = [
            row for row in reconciliation_rows if row.get("status") in {"pendente", "divergente"}
        ]
        for row in problematic_rows:
            problems_sheet.append(
                [
                    row.get("row_id"),
                    self._translate_source_label(row.get("source")),
                    row.get("date"),
                    row.get("description"),
                    row.get("amount"),
                    self._translate_status_label(row.get("status")),
                    self._translate_reason_label(row.get("reason")),
                    row.get("matched_row_id"),
                ]
            )
        if len(problematic_rows) == 0:
            problems_sheet.append(
                [
                    "nenhum",
                    "Sistema",
                    "",
                    "Nenhuma pendencia ou divergencia foi identificada.",
                    "",
                    "-",
                    "-",
                    "",
                ]
            )
        self._format_transacoes_sheet(problems_sheet)

        workbook.save(analysis_dir / "reconcile_report.xlsx")

        csv_headers = detail_headers
        csv_lines = [",".join(csv_headers)]
        for row in reconciliation_rows:
            values = [
                self._escape_csv_value(value)
                for value in [
                    row.get("row_id"),
                    self._translate_source_label(row.get("source")),
                    row.get("date"),
                    row.get("description"),
                    row.get("amount"),
                    self._translate_status_label(row.get("status")),
                    self._translate_match_rule_label(row.get("match_rule")),
                    self._translate_reason_label(row.get("reason")),
                    row.get("matched_row_id"),
                ]
            ]
            csv_lines.append(",".join(values))
        (analysis_dir / "reconcile_report.csv").write_text("\n".join(csv_lines), encoding="utf-8")

        return analysis_id, expires_at.isoformat()

    def get_reconcile_report_path(self, analysis_id: str, file_format: str) -> Path:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError

        suffix = "xlsx" if file_format == "xlsx" else "csv"
        report_path = analysis_dir / f"reconcile_report.{suffix}"
        if not report_path.exists():
            raise AnalysisNotFoundError
        return report_path

    def _regenerate_ofx_with_overrides(
        self,
        analysis_dir: Path,
        *,
        closing_balance: float | None = None,
        bank_branch: str | None = None,
        account_number: str | None = None,
        bank_code: str | None = None,
    ) -> None:
        metadata_path = analysis_dir / "analysis.json"
        if not metadata_path.exists():
            raise AnalysisNotFoundError
        try:
            content = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raise AnalysisNotFoundError from None

        report_rows_raw = content.get("report_transactions") or content.get("preview_transactions", [])
        report_rows = self._parse_transaction_rows(report_rows_raw)
        if not report_rows and not _is_confirmed_empty_statement(content):
            raise AnalysisNotFoundError

        assert_account_overrides(report_rows, closing_balance=closing_balance, bank_branch=bank_branch,
                                 account_number=account_number, bank_code=bank_code,
                                 statement_accounts=statement_scopes(content.get("document_parts")))

        persisted_overrides = self._extract_convert_overrides(content)
        requested_overrides = self._normalize_convert_overrides(
            opening_balance=None,
            closing_balance=closing_balance,
            bank_branch=bank_branch,
            account_number=account_number,
            bank_code=bank_code,
        )
        effective_overrides = {
            "opening_balance": persisted_overrides.get("opening_balance"),
            "closing_balance": (
                requested_overrides.get("closing_balance")
                if requested_overrides.get("closing_balance") is not None
                else persisted_overrides.get("closing_balance")
            ),
            "bank_branch": (
                requested_overrides.get("bank_branch")
                if requested_overrides.get("bank_branch") is not None
                else persisted_overrides.get("bank_branch")
            ),
            "account_number": (
                requested_overrides.get("account_number")
                if requested_overrides.get("account_number") is not None
                else persisted_overrides.get("account_number")
            ),
            "bank_code": (
                requested_overrides.get("bank_code")
                if requested_overrides.get("bank_code") is not None
                else persisted_overrides.get("bank_code")
            ),
        }
        metrics = content.get("pdf_processing_metrics") or {}
        if metrics.get("financial_controls_validated") and not metrics.get("balance_bounds_unavailable"):
            # Changing the destination account must preserve confirmed statement
            # boundaries. Explicit user overrides still take precedence.
            for key in ("opening_balance", "closing_balance"):
                if effective_overrides.get(key) is None:
                    effective_overrides[key] = content.get(key)
        if requested_overrides:
            self._apply_convert_overrides_to_content(content, requested_overrides)
            self._refresh_structured_result_content(content, preview_rows=report_rows, report_rows=report_rows)
            self._write_metadata(metadata_path, content)

        self.export_artifact_service.write_convert_artifacts(
            analysis_dir,
            report_rows=report_rows,
            ofx_account_type=str(content.get("ofx_account_type") or "").strip() or None,
            layout_inference_name=str(content.get("layout_inference_name") or "").strip() or None,
            opening_balance=effective_overrides.get("opening_balance"),
            closing_balance=effective_overrides.get("closing_balance"),
            balance_bounds_unavailable=bool((content.get("pdf_processing_metrics") or {}).get("balance_bounds_unavailable")),
            bank_branch=effective_overrides.get("bank_branch"),
            account_number=effective_overrides.get("account_number"),
            bank_code=effective_overrides.get("bank_code"),
            **({"financial_document": content["financial_document"]} if content.get("financial_document") is not None else {}),
            **({"statement_accounts": statement_scopes(content.get("document_parts"))}
               if statement_scopes(content.get("document_parts")) else {}),
        )

    def _normalize_convert_overrides(
        self,
        *,
        opening_balance: float | None,
        closing_balance: float | None,
        bank_branch: str | None,
        account_number: str | None,
        bank_code: str | None,
    ) -> dict[str, float | str]:
        normalized: dict[str, float | str] = {}
        if opening_balance is not None:
            normalized["opening_balance"] = float(opening_balance)
        if closing_balance is not None:
            normalized["closing_balance"] = float(closing_balance)
        branch = str(bank_branch or "").strip()
        if branch:
            normalized["bank_branch"] = branch
        account = str(account_number or "").strip()
        if account:
            normalized["account_number"] = account
        code = str(bank_code or "").strip()
        if code:
            normalized["bank_code"] = code
        return normalized

    def _extract_convert_overrides(self, content: dict[str, object]) -> dict[str, float | str | None]:
        empty_statement = _is_confirmed_empty_statement(content)
        opening_raw = content.get("convert_opening_balance")
        if opening_raw is None and empty_statement:
            opening_raw = content.get("opening_balance")
        opening_balance = float(opening_raw) if opening_raw is not None else None
        closing_raw = content.get("convert_closing_balance")
        if closing_raw is None and empty_statement:
            closing_raw = content.get("closing_balance")
        closing_balance = float(closing_raw) if closing_raw is not None else None
        preserve_bank_metadata = (content.get("file_type") == "ofx"
                                  or content.get("file_type") == "pdf" and content.get("financial_document") is None)
        source = content if preserve_bank_metadata or empty_statement else {}
        branch = str(content.get("convert_bank_branch") or source.get("bank_branch") or "").strip() or None
        account = str(content.get("convert_account_number") or source.get("account_number") or "").strip() or None
        code = str(content.get("convert_bank_code") or source.get("bank_code") or "").strip() or None
        return {
            "opening_balance": opening_balance,
            "closing_balance": closing_balance,
            "bank_branch": branch,
            "account_number": account,
            "bank_code": code,
        }

    def _apply_convert_overrides_to_content(
        self,
        content: dict[str, object],
        overrides: dict[str, float | str],
    ) -> None:
        if "opening_balance" in overrides:
            content["convert_opening_balance"] = float(overrides["opening_balance"])
        if "closing_balance" in overrides:
            content["convert_closing_balance"] = float(overrides["closing_balance"])
        if "bank_branch" in overrides:
            content["convert_bank_branch"] = str(overrides["bank_branch"])
        if "account_number" in overrides:
            content["convert_account_number"] = str(overrides["account_number"])
        if "bank_code" in overrides:
            content["convert_bank_code"] = str(overrides["bank_code"])

    def _write_report_workbook(
        self,
        analysis_dir: Path,
        content: dict[str, object],
        report_rows: list[TransactionRow],
        preview_rows: list[TransactionRow],
    ) -> None:
        snapshot = self._build_analysis_snapshot_from_content(
            content,
            preview_rows=preview_rows,
            report_rows=report_rows,
        )
        self.export_artifact_service.write_analysis_report_workbook(
            analysis_dir / "report.xlsx",
            report_rows=report_rows,
            snapshot=snapshot,
        )

    def _parse_transaction_rows(self, rows_raw: object) -> list[TransactionRow]:
        parsed: list[TransactionRow] = []
        if not isinstance(rows_raw, list):
            return parsed
        for item in rows_raw:
            if not isinstance(item, dict):
                continue
            parsed.append(
                TransactionRow(
                    statement_account=dict(item["statement_account"]) if isinstance(item.get("statement_account"), dict) else None,
                    balance_check=parse_balance_check(item.get("balance_check")),
                    date=str(item.get("date") or ""),
                    description=str(item.get("description") or ""),
                    amount=float(item.get("amount") or 0.0),
                    category=str(item.get("category") or "Outros"),
                    reconciliation_status=str(item.get("reconciliation_status") or "unmatched"),
                    source_page=item.get("source_page") if isinstance(item.get("source_page"), int) else None,
                    source_regions=parse_source_regions(item.get("source_regions")),
                    review_regions=parse_source_regions(item.get("review_regions")),
                    source_line=item.get("source_line") if isinstance(item.get("source_line"), int) else None,
                    transaction_id=item.get("transaction_id") if isinstance(item.get("transaction_id"), str) else None,
                    source_sequence=item.get("source_sequence") if isinstance(item.get("source_sequence"), int) else None,
                    source_subsequence=int(item.get("source_subsequence") or 0),
                    running_balance=(
                        float(item.get("running_balance"))
                        if item.get("running_balance") is not None
                        else None
                    ),
                    is_deleted=bool(item.get("is_deleted") or False),
                    warning_types=[
                        str(entry).strip()
                        for entry in (item.get("warning_types") or [])
                        if str(entry).strip()
                    ]
                    if isinstance(item.get("warning_types"), list)
                    else [],
                )
            )
        return parsed

    def _parse_before_after_rows(self, rows_raw: object) -> list[BeforeAfterRow]:
        parsed: list[BeforeAfterRow] = []
        if not isinstance(rows_raw, list):
            return parsed
        for item in rows_raw:
            if not isinstance(item, dict):
                continue
            parsed.append(
                BeforeAfterRow(
                    date=str(item.get("date") or ""),
                    description_before=str(item.get("description_before") or ""),
                    description_after=str(item.get("description_after") or ""),
                    amount_before=float(item.get("amount_before") or 0.0),
                    amount_after=float(item.get("amount_after") or 0.0),
                )
            )
        return parsed

    def _active_rows(self, rows: list[TransactionRow]) -> list[TransactionRow]:
        return [item for item in rows if not item.is_deleted]

    def _build_analysis_snapshot_from_content(
        self,
        content: dict[str, object],
        *,
        preview_rows: list[TransactionRow],
        report_rows: list[TransactionRow],
    ) -> AnalysisData:
        return AnalysisData(
            analysis_id=str(content.get("analysis_id", "")),
            file_type=str(content.get("file_type", "")),
            upload_filename=content.get("upload_filename") if isinstance(content.get("upload_filename"), str) else None,
            transactions_total=int(content.get("transactions_total", len(self._active_rows(report_rows)))),
            total_inflows=float(content.get("total_inflows", 0.0)),
            total_outflows=float(content.get("total_outflows", 0.0)),
            net_total=float(content.get("net_total", 0.0)),
            preview_transactions=preview_rows,
            report_transactions=report_rows,
            semantic_type=str(content.get("semantic_type") or "").strip() or None,
            semantic_confidence=(
                float(content["semantic_confidence"]) if content.get("semantic_confidence") is not None else None
            ),
            semantic_evidence=[
                str(item).strip()
                for item in (content.get("semantic_evidence") or [])
                if str(item).strip()
            ]
            if isinstance(content.get("semantic_evidence"), list)
            else None,
            document_type=str(content.get("document_type") or "").strip() or None,
            financial_document=content.get("financial_document"),
            document_type_confidence=(
                float(content["document_type_confidence"])
                if content.get("document_type_confidence") is not None
                else None
            ),
            document_type_evidence=[
                str(item).strip()
                for item in (content.get("document_type_evidence") or [])
                if str(item).strip()
            ]
            if isinstance(content.get("document_type_evidence"), list)
            else None,
            document_classification_version=(
                str(content.get("document_classification_version") or "").strip() or None
            ),
            document_parts=[dict(part) for part in content.get("document_parts", []) if isinstance(part, dict)]
            if isinstance(content.get("document_parts"), list) else [],
            preview_before_after=self._parse_before_after_rows(content.get("preview_before_after", [])),
            matched_groups=int(content.get("matched_groups", 0)),
            reversed_entries=int(content.get("reversed_entries", 0)),
            potential_duplicates=int(content.get("potential_duplicates", 0)),
            updated_at=str(content.get("updated_at") or "").strip() or None,
            layout_inference_name=str(content.get("layout_inference_name") or "").strip() or None,
            layout_inference_confidence=(
                float(content["layout_inference_confidence"])
                if content.get("layout_inference_confidence") is not None
                else None
            ),
            pdf_processing_metrics=(
                dict(content["pdf_processing_metrics"])
                if isinstance(content.get("pdf_processing_metrics"), dict)
                else None
            ),
            quality_issues=[dict(item) for item in content.get("quality_issues", []) if isinstance(item, dict)]
            if isinstance(content.get("quality_issues"), list)
            else [],
            ofx_account_type=str(content.get("ofx_account_type") or "").strip() or None,
            opening_balance=float(content["opening_balance"]) if content.get("opening_balance") is not None else None,
            closing_balance=float(content["closing_balance"]) if content.get("closing_balance") is not None else None,
            statement_balance_check=parse_balance_check(content.get("statement_balance_check")),
            bank_name=str(content.get("bank_name") or "").strip() or None,
            bank_branch=str(content.get("bank_branch") or "").strip() or None,
            account_number=str(content.get("account_number") or "").strip() or None,
            bank_code=str(content.get("bank_code") or "").strip() or None,
        )

    def _refresh_structured_result_content(
        self,
        content: dict[str, object],
        *,
        preview_rows: list[TransactionRow],
        report_rows: list[TransactionRow],
    ) -> None:
        snapshot = self._build_analysis_snapshot_from_content(
            content,
            preview_rows=preview_rows,
            report_rows=report_rows,
        )
        content["structured_result"] = asdict(
            build_structured_conversion_result_from_analysis_data(
                snapshot,
                expires_at=str(content.get("expires_at") or "").strip() or None,
            )
        )

    def _resolve_row_index(self, row_id_raw: object, max_index: int) -> int:
        row_id = str(row_id_raw or "").strip()
        if not row_id.startswith("row_"):
            raise ValueError("Invalid row_id.")
        suffix = row_id.split("_", 1)[1]
        if not suffix.isdigit():
            raise ValueError("Invalid row_id.")
        idx = int(suffix) - 1
        if idx < 0 or idx >= max_index:
            raise ValueError("row_id out of bounds.")
        return idx

    def _resolve_insert_position(self, position_raw: object, max_index: int) -> int:
        if position_raw is None:
            return max_index
        position = int(position_raw)
        if position < 0 or position > max_index:
            raise ValueError("insert_position out of bounds.")
        return position

    def _build_amount_from_credit_debit(self, credit_raw: object, debit_raw: object) -> float:
        credit = None if credit_raw is None else float(credit_raw)
        debit = None if debit_raw is None else float(debit_raw)
        has_credit = credit is not None and credit > 0
        has_debit = debit is not None and debit > 0
        if has_credit == has_debit:
            raise ValueError("Provide only one of credit or debit.")
        if has_credit:
            return round(float(credit), 2)
        return -round(abs(float(debit)), 2)

    def _set_analysis_owner(
        self,
        *,
        analysis_id: str,
        metadata_filename: str,
        identity_type: str,
        identity_id: str,
    ) -> None:
        content, metadata_path = self._read_analysis_metadata(
            analysis_id=analysis_id,
            metadata_filename=metadata_filename,
        )

        owner_type = str(content.get("owner_identity_type") or "").strip()
        owner_id = str(content.get("owner_identity_id") or "").strip()
        if owner_type and owner_id:
            if owner_type != identity_type or owner_id != identity_id:
                raise AnalysisAccessDeniedError
            return

        content["owner_identity_type"] = identity_type
        content["owner_identity_id"] = identity_id
        self._write_metadata(metadata_path, content)

    def _assert_analysis_owner(
        self,
        *,
        analysis_id: str,
        metadata_filename: str,
        identity_type: str,
        identity_id: str,
        allow_unowned: bool,
    ) -> None:
        content, _metadata_path = self._read_analysis_metadata(
            analysis_id=analysis_id,
            metadata_filename=metadata_filename,
        )

        owner_type = str(content.get("owner_identity_type") or "").strip()
        owner_id = str(content.get("owner_identity_id") or "").strip()
        if not owner_type and not owner_id:
            if allow_unowned:
                return
            raise AnalysisAccessDeniedError
        if owner_type != identity_type or owner_id != identity_id:
            raise AnalysisAccessDeniedError

    def _read_analysis_metadata(self, *, analysis_id: str, metadata_filename: str) -> tuple[dict[str, object], Path]:
        analysis_dir = self._resolve_analysis_dir(analysis_id)
        if self._is_expired(analysis_dir):
            self._cleanup_analysis(analysis_dir)
            raise AnalysisNotFoundError

        metadata_path = analysis_dir / metadata_filename
        if not metadata_path.exists():
            raise AnalysisNotFoundError
        try:
            content = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raise AnalysisNotFoundError from None
        return content, metadata_path

    def _resolve_analysis_dir(self, analysis_id: str) -> Path:
        clean_analysis_id = str(analysis_id or "").strip()
        if not _SAFE_ANALYSIS_ID_RE.fullmatch(clean_analysis_id):
            raise AnalysisNotFoundError
        return self.root_dir / clean_analysis_id

    def _is_expired(self, analysis_dir: Path) -> bool:
        metadata_path = analysis_dir / "analysis.json"
        try:
            content = json.loads(metadata_path.read_text(encoding="utf-8"))
            expires_at_raw = content.get("expires_at")
            if not isinstance(expires_at_raw, str):
                return False
            expires_at = datetime.fromisoformat(expires_at_raw)
            return self.now_provider() > expires_at
        except (OSError, TypeError, json.JSONDecodeError, ValueError):
            return False

    def _cleanup_analysis(self, analysis_dir: Path) -> None:
        if not analysis_dir.exists():
            return
        for file in analysis_dir.glob("**/*"):
            if file.is_file():
                file.unlink(missing_ok=True)
        for directory in sorted((item for item in analysis_dir.glob("**/*") if item.is_dir()), key=lambda x: len(x.parts), reverse=True):
            directory.rmdir()
        analysis_dir.rmdir()

    def _format_transacoes_sheet(self, sheet) -> None:
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

        for column_index in range(1, sheet.max_column + 1):
            column_letter = get_column_letter(column_index)
            max_len = 0
            for row_index in range(1, sheet.max_row + 1):
                cell_value = sheet.cell(row=row_index, column=column_index).value
                max_len = max(max_len, len(str(cell_value)) if cell_value is not None else 0)
            sheet.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 80)

    def _build_analysis_id(self, prefix: str) -> str:
        return f"{prefix}_{uuid4().hex[:12]}"

    def _escape_csv_value(self, value: object) -> str:
        text = "" if value is None else str(value)
        if any(char in text for char in [",", "\"", "\n", "\r"]):
            return f"\"{text.replace('\"', '\"\"')}\""
        return text

    def _translate_source_label(self, source: object) -> str:
        text = "" if source is None else str(source)
        labels = {
            "bank": "Banco",
            "sheet": "Planilha",
            "system": "Sistema",
        }
        return labels.get(text, text or "-")

    def _translate_status_label(self, status: object) -> str:
        text = "" if status is None else str(status)
        labels = {
            "conciliado": "Conciliado",
            "pendente": "Pendente",
            "divergente": "Divergente",
            "none": "-",
        }
        return labels.get(text, text or "-")

    def _translate_match_rule_label(self, match_rule: object) -> str:
        text = "" if match_rule is None else str(match_rule)
        labels = {
            "exact": "Exato",
            "date_tolerance": "Tolerancia de data",
            "description_similarity": "Similaridade de descricao",
            "none": "-",
        }
        return labels.get(text, text or "-")

    def _translate_reason_label(self, reason: object) -> str:
        text = "" if reason is None else str(reason)
        labels = {
            "missing_in_sheet": "Pendente na planilha",
            "missing_in_bank": "Pendente no banco",
            "amount_mismatch": "Diferenca de valor",
            "date_out_of_tolerance_window": "Data fora da tolerancia",
            "matched_equal_amount_same_day": "Valor igual na mesma data",
            "matched_equal_amount_within_2_days": "Valor igual dentro de 2 dias",
            "matched_equal_amount_with_similar_description": "Valor igual com descricao similar",
            "none": "-",
        }
        return labels.get(text, text or "-")

    def _resolve_balance_bounds(
        self,
        rows: list[TransactionRow],
        *,
        opening_balance_hint: float | str | None = None,
        closing_balance_hint: float | str | None = None,
    ) -> tuple[float | None, float | None]:
        opening_balance = self._coerce_balance_value(opening_balance_hint)
        if opening_balance is None:
            for item in rows:
                if self._is_opening_balance_row(item):
                    opening_balance = round(float(item.amount), 2)
                    break
            if opening_balance is None:
                for item in rows:
                    if item.running_balance is None:
                        continue
                    opening_balance = round(float(item.running_balance) - float(item.amount), 2)
                    break

        manual_closing_balance = self._coerce_balance_value(closing_balance_hint)
        if manual_closing_balance is not None:
            return opening_balance, manual_closing_balance

        last_balance_index: int | None = None
        closing_balance: float | None = None
        for index, item in enumerate(rows):
            if item.running_balance is None:
                continue
            last_balance_index = index
            closing_balance = round(float(item.running_balance), 2)
        if closing_balance is not None and last_balance_index is not None:
            trailing_amount = sum(float(item.amount) for item in rows[last_balance_index + 1 :])
            return opening_balance, round(closing_balance + trailing_amount, 2)

        if opening_balance is not None:
            transaction_amounts = sum(float(item.amount) for item in rows if not self._is_opening_balance_row(item))
            return opening_balance, round(opening_balance + transaction_amounts, 2)

        return None, None

    def _resolve_recalculation_opening_balance(
        self,
        *,
        original_rows: list[dict[str, object]] | list[TransactionRow],
        current_rows: list[TransactionRow],
        content: dict[str, object],
        requested_opening_balance: float | str | None,
        persisted_opening_balance: float | str | None,
    ) -> float | None:
        for candidate in (
            requested_opening_balance,
            persisted_opening_balance,
            content.get("opening_balance"),
        ):
            coerced = self._coerce_balance_value(candidate)
            if coerced is not None:
                return coerced

        original_parsed_rows = (
            original_rows
            if original_rows and isinstance(original_rows[0], TransactionRow)
            else self._parse_transaction_rows(original_rows)
        )
        if (content.get("pdf_processing_metrics") or {}).get("balance_bounds_unavailable"):
            return None
        original_opening_balance, _ = self._resolve_balance_bounds(original_parsed_rows)
        if original_opening_balance is not None:
            return original_opening_balance

        current_opening_balance, _ = self._resolve_balance_bounds(current_rows)
        return current_opening_balance

    def _recalculate_running_balances(
        self,
        *,
        preview_rows: list[TransactionRow],
        report_rows: list[TransactionRow],
        opening_balance: float | None,
    ) -> None:
        has_balance_context = opening_balance is not None or any(
            row.running_balance is not None for row in report_rows
        )
        if not has_balance_context:
            return

        cursor = round(float(opening_balance), 2) if opening_balance is not None else None
        for preview_row, report_row in zip(preview_rows, report_rows, strict=False):
            if preview_row.is_deleted or report_row.is_deleted:
                preview_row.running_balance = None
                report_row.running_balance = None
                continue

            if self._is_opening_balance_row(report_row):
                cursor = round(float(report_row.amount), 2)
                preview_row.running_balance = cursor
                report_row.running_balance = cursor
                continue

            if cursor is None:
                preview_row.running_balance = None
                report_row.running_balance = None
                continue

            cursor = round(cursor + float(report_row.amount), 2)
            preview_row.running_balance = cursor
            report_row.running_balance = cursor

    def _coerce_balance_value(self, value: float | str | None) -> float | None:
        if value is None or value == "":
            return None
        try:
            return round(float(value), 2)
        except (TypeError, ValueError):
            return None

    def _is_opening_balance_row(self, row: TransactionRow) -> bool:
        normalized = " ".join(str(row.description or "").strip().upper().split())
        return normalized in {"SALDO ANTERIOR", "SALDO INICIAL"}
