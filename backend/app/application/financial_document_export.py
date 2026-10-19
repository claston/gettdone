"""Exports preserve report semantics, original fields, and validation evidence."""
import csv
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from app.application.errors import InvalidFileContentError
from app.application.financial_event_projection import event_date
from app.application.models import NormalizedTransaction
from app.application.ofx_writer import build_ofx_statement

LABELS = {
    "event_id": "ID do evento", "event_type": "Tipo de evento", "status": "Situação", "reference": "Referência",
    "counterparty": "Contraparte", "amount": "Valor", "amount_basis": "Base do valor", "currency": "Moeda",
    "cash_effect": "Efeito financeiro", "description": "Descrição", "source_refs": "Origem (página e linha)",
    "dates.payment": "Data do pagamento", "dates.due": "Vencimento", "dates.application": "Data da aplicação",
    "dates.redemption": "Data do resgate", "dates.maturity": "Vencimento do investimento", "dates.issue": "Emissão",
    "amounts.original": "Valor original", "amounts.collected": "Valor cobrado", "amounts.adjustment": "Oscilação",
    "amounts.principal": "Principal", "amounts.gross": "Valor bruto", "amounts.net": "Valor líquido",
    "amounts.income": "Renda total", "amounts.iof": "IOF", "amounts.irrf": "IRRF",
    "amounts.period_income": "Renda bruta no período", "amounts.document": "Valor do documento",
    "attributes.access_key": "Chave de acesso", "attributes.protocol": "Protocolo", "attributes.model": "Modelo fiscal",
    "attributes.fiscal_direction": "Entrada/saída fiscal", "attributes.status_code": "Código da situação",
    "attributes.rate_percent": "Taxa (%)", "attributes.source_values": "Valores na ordem da origem", "attributes.car": "Car",
    "observation_type": "Tipo de posição", "field": "Campo", "value": "Conteúdo", "code": "Validação",
    "actual": "Extraído", "expected": "Impresso / esperado",
    "owner_name": "Titular", "owner_tax_id": "CPF/CNPJ do titular", "period_start": "Início do período",
    "period_end": "Fim do período", "source_bank_branch": "Agência de origem", "source_account_number": "Conta de origem",
    "product_account": "Conta do produto",
}


def assert_financial_export(document: dict, file_format: str, *, bank_code=None, bank_branch=None, account_number=None):
    policy = document["export"]
    if file_format not in policy["available_formats"]:
        raise InvalidFileContentError(f"Formato {file_format.upper()} indisponível. {policy['reason']}")
    if file_format == "ofx" and policy["requires_destination_account"]:
        cleaned = [str(value or "").strip().replace("-", "").replace(".", "")
                   for value in (bank_code, bank_branch, account_number)]
        if not all(value.isascii() and value.isdigit() for value in cleaned) or len(cleaned[0]) != 3:
            raise InvalidFileContentError("Informe banco, agência e conta de destino para exportar os recebimentos em OFX.")


def _flatten(event):
    row = {key: value for key, value in event.items() if key not in {"dates", "amounts", "attributes"}}
    for section in ("dates", "amounts", "attributes"):
        row.update({f"{section}.{key}": value for key, value in event.get(section, {}).items()})
    return {key: json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value for key, value in row.items()}


def _sheet(workbook, title, rows, *, default_headers=()):
    sheet = workbook.create_sheet(title)
    headers = list(dict.fromkeys(key for row in rows for key in row)) or list(default_headers)
    sheet.append([LABELS.get(key, key) for key in headers])
    for row in rows:
        sheet.append([row.get(key) for key in headers])
        # Identifiers, including leading zeros and formula-like descriptions, are text.
        for key, cell in zip(headers, sheet[sheet.max_row], strict=True):
            if cell.value is not None and (key in {"amount", "cash_effect"} or key.startswith("amounts.")):
                cell.value = float(Decimal(cell.value))
                cell.number_format = "#,##0.00"
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(70, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    return sheet


def write_financial_workbook(path: Path, document: dict):
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Eventos", [_flatten(event) for event in document["events"]], default_headers=("event_type", "amount", "cash_effect"))
    details = [{"event_id": event.get("event_id"), "field": key, "value": value}
               for event in document["events"] for key, value in _flatten(event).items()]
    details = [dict(item, field=LABELS.get(item["field"], item["field"])) for item in details]
    _sheet(workbook, "Detalhes", details, default_headers=("event_id", "field", "value"))
    _sheet(workbook, "Posições", [_flatten(item) for item in document["observations"]], default_headers=("observation_type",))
    _sheet(workbook, "Validações", [dict(item, source_refs=json.dumps(item["source_refs"])) for item in document["validations"]])
    document_fields = {
        "Tipo": document["document_type"], "Origem": document["origin"], "Observação": document["notice"],
        "Formatos disponíveis": ", ".join(document["export"]["available_formats"]),
        "Extração": document.get("extraction_method"), "Versão": document["schema_version"],
        "Documento SHA-256": document.get("source_document_sha256"),
        "Edição": "Os campos de origem são preservados. A edição de relatórios especializados ainda não está disponível.",
    }
    document_fields.update({key: value for key, value in document.get("metadata", {}).items()})
    _sheet(workbook, "Documento", [{"Campo": key, "Valor": value} for key, value in document_fields.items()])
    workbook.save(path)


def write_financial_artifacts(output_dir: Path, document: dict, *, bank_code=None, bank_branch=None, account_number=None):
    rows = [dict(record_type="event", **_flatten(event)) for event in document["events"]]
    rows.extend(dict(record_type="observation", **_flatten(item)) for item in document["observations"])
    headers = list(dict.fromkeys(key for row in rows for key in row)) or ["event_type", "amount", "cash_effect"]
    with (output_dir / "converted.csv").open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))
                             and key not in {"amount", "cash_effect"} and not key.startswith("amounts.") else value
                             for key, value in row.items()})
    write_financial_workbook(output_dir / "converted.xlsx", document)
    if "ofx" not in document["export"]["available_formats"] or not all((bank_code, bank_branch, account_number)):
        (output_dir / "converted.ofx").unlink(missing_ok=True)
        return
    assert_financial_export(document, "ofx", bank_code=bank_code, bank_branch=bank_branch, account_number=account_number)
    rows = [NormalizedTransaction(date=event_date(event), description=event["description"],
            amount=float(Decimal(event["cash_effect"])), type="inflow") for event in document["events"]]
    (output_dir / "converted.ofx").write_text(build_ofx_statement(rows, bank_id=bank_code, bank_branch=bank_branch,
                                                               account_number=account_number, currency="BRL"), encoding="utf8")
