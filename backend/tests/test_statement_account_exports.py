"""Scoped accounts survive JSON reload, edits and standard financial exports."""

import csv
import json
import re
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.errors import InvalidFileContentError
from app.application.parsers.ofx import parse_ofx_transactions
from app.application.storage_service import TempAnalysisStorage
from tests.fixtures.positioned_cash_statements import positioned_pdf
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_statement_account_blocks import account_page, consolidated_pages, inactive_page, ocr_pages


@pytest.fixture(autouse=True)
def enabled(monkeypatch):
    monkeypatch.setenv("PDF_MULTI_ACCOUNT_ENABLED", "true")
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")


def assert_ofx_accounts(content, *, first_balance=True):
    blocks = re.findall(r"<STMTTRNRS>(.*?)</STMTTRNRS>", content.decode(), re.S)
    assert len(blocks) == 2
    for i, (block, account, balance) in enumerate(zip(blocks, ["123456", "987654"], [130, 230], strict=True)):
        assert f"<ACCTID>{account}" in block
        assert "<BANKID>033" in block
        assert [t.amount for t in parse_ofx_transactions(block.encode())] == [50 if first_balance or i else 60, -20]
        assert ("<LEDGERBAL>" in block) == (i != 0 or first_balance)
        if i != 0 or first_balance:
            assert f"<BALAMT>{balance:.2f}" in block


@pytest.mark.real_http
@pytest.mark.parametrize("mode", ["native", "text", "analysis"])
def test_http_account_scopes_are_visible_and_preserved_in_every_export(mode, tmp_path, monkeypatch):
    pages = consolidated_pages()
    raw = positioned_pdf(pages) if mode == "native" else configure_ocr(monkeypatch, ocr_pages(pages), mode)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-account-scopes-" + mode}
        response = client.post("/convert", data=owner, files={"file": ("consolidado.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        data = response.json()["analysis"]
        assert data["account_number"] is data["bank_branch"] is data["opening_balance"] is data["closing_balance"] is None
        assert [r["statement_account"]["account_number"] for r in data["preview_transactions"]] == ["12345-6"] * 2 + ["98765-4"] * 2
        identifier = data["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200, report.text
            if fmt == "ofx":
                assert_ofx_accounts(report.content)
            elif fmt == "csv":
                rows = list(csv.DictReader(StringIO(report.text)))
                assert [r["account_number"] for r in rows] == ["12345-6"] * 2 + ["98765-4"] * 2
                assert [float(r["amount"]) for r in rows] == [50, -20] * 2
            else:
                book = load_workbook(BytesIO(report.content), data_only=True)
                rows = list(book["Conversao"].iter_rows(min_row=2, values_only=True))
                assert [r[7] for r in rows] == ["12345-6"] * 2 + ["98765-4"] * 2
                assert book["Resumo"].max_row == 3
                book.close()
        edited = client.post(f"/convert-edits/{identifier}", params=owner, json={"edits": [{
            "row_id": data["preview_transactions"][0]["transaction_id"], "date": "2026-09-01",
            "description": "PIX CORRIGIDO", "credit": 60, "debit": None}]})
        assert edited.status_code == 200, edited.text
        accounts = [r["statement_account"]["account_number"] for r in edited.json()["preview_transactions"]]
        assert accounts == ["12345-6"] * 2 + ["98765-4"] * 2
        exported = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx"})
        assert exported.status_code == 200
        assert_ofx_accounts(exported.content, first_balance=False)
        rejected = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx", "account_number": "99999"})
        assert rejected.status_code == 400, rejected.text
        assert client.get("/report/an_accounts_nonexistent", params=owner).status_code == 404


def test_json_reload_and_anchored_insert_preserve_account_identity(tmp_path):
    data = build_default_conversion_pipeline().run(filename="consolidado.pdf", raw_bytes=positioned_pdf(consolidated_pages()),
                                                  analysis_id="an_account_storage").analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    stored = json.loads((tmp_path/data.analysis_id/"analysis.json").read_text())
    assert stored["structured_result"]["transactions"][2]["statement_account"]["account_number"] == "98765-4"
    with pytest.raises(InvalidFileContentError):
        storage.apply_convert_edits(data.analysis_id, [{"action": "insert", "insert_position": 2,
            "date": "2026-09-02", "description": "MANUAL", "credit": 10}])
    result = storage.apply_convert_edits(data.analysis_id, [{"action": "insert", "insert_position": 2,
        "anchor_transaction_id": stored["preview_transactions"][2]["transaction_id"],
        "date": "2026-09-02", "description": "MANUAL", "credit": 10}])
    assert result["preview_transactions"][2]["statement_account"]["account_number"] == "98765-4"
    assert result["opening_balance"] is result["closing_balance"] is None
    for kwargs in ({"opening_balance": 100}, {"closing_balance": 200}, {"bank_branch": "9999"}):
        with pytest.raises(InvalidFileContentError):
            storage.apply_convert_edits(data.analysis_id, [], **kwargs)


@pytest.mark.real_http
def test_worker_preserves_separate_account_exports(tmp_path):
    with worker_case(tmp_path, positioned_pdf(consolidated_pages())) as (client, job):
        data = job["result"]["analysis"]
        assert data["transactions_total"] == 4
        assert len({r["statement_account"]["account_id"] for r in data["preview_transactions"]}) == 2
        exported = client.get(f"/convert-report/{data['analysis_id']}", params={"format": "ofx"})
        assert exported.status_code == 200
        assert_ofx_accounts(exported.content)


def test_an_inactive_account_remains_in_ofx_and_workbook_summary(tmp_path):
    data = build_default_conversion_pipeline().run(filename="consolidado.pdf", raw_bytes=positioned_pdf([account_page(), inactive_page()]),
                                                  analysis_id="an_inactive_account").analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    ofx = storage.get_convert_report_path(data.analysis_id, "ofx").read_text()
    assert ofx.count("<BANKACCTFROM>") == 2
    assert "<ACCTID>987654" in ofx
    assert "<BALAMT>200.00" in ofx
    assert "<DTASOF>20260930" in ofx
    assert ofx.count("<STMTTRN>") == 2
    book = load_workbook(storage.get_convert_report_path(data.analysis_id, "xlsx"), data_only=True)
    assert book["Resumo"].max_row == 3
    book.close()


def test_two_inactive_accounts_do_not_create_a_global_account_or_cash_movements(tmp_path):
    first = [(x, y, "12345-6" if text == "98765-4" else "0017" if text == "0039"
              else "100,00" if text == "200,00" else text) for x, y, text in inactive_page()]
    data = build_default_conversion_pipeline().run(filename="consolidado.pdf", raw_bytes=positioned_pdf([first, inactive_page()]),
                                                  analysis_id="an_inactive_accounts").analysis_data
    assert data.transactions_total == 0
    assert data.bank_branch is data.account_number is data.opening_balance is data.closing_balance is None
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    ofx = storage.get_convert_report_path(data.analysis_id, "ofx").read_text()
    assert ofx.count("<BANKACCTFROM>") == 2
    assert "<STMTTRN>" not in ofx
    with pytest.raises(InvalidFileContentError):
        storage.get_convert_report_path(data.analysis_id, "ofx", account_number="99999")


def test_repeated_blocks_of_one_account_keep_distinct_ofx_transaction_ids(tmp_path):
    pages = [account_page(), account_page(account="98765-4", branch="0039", opening=200), account_page()]
    data = build_default_conversion_pipeline().run(filename="consolidado.pdf", raw_bytes=positioned_pdf(pages),
                                                  analysis_id="an_repeated_account_blocks").analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    ofx = storage.get_convert_report_path(data.analysis_id, "ofx").read_text()
    blocks = re.findall(r"<STMTTRNRS>(.*?)</STMTTRNRS>", ofx, re.S)
    ids = [value for block in blocks if "<ACCTID>123456" in block for value in re.findall(r"<FITID>([^\n<]+)", block)]
    assert len(ids) == len(set(ids)) == 4
    assert ofx.count("<STMTTRN>") == 6
