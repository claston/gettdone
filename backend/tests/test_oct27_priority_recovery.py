"""Fictional regressions for the two priority recovery documents."""
from dataclasses import replace

import pytest

from app.application import pdf_parser
from app.application.errors import PdfFinancialValidationError
from tests.fixtures.recovery_acceptance import grid_pdf


def infinitepay_fragment(*, mutation=None):
    def header():
        return [("InfinitePay Relatório de movimentações", .03, .03),
                ("Data", .04, .12), ("Hora", .13, .12), ("Tipo de transação", .21, .12),
                ("Nome", .4, .12), ("Detalhe", .66, .12), ("Valor (R$)", .91, .12),
                ("Central de ajuda ajuda@infinitepay.io", .03, .94)]

    def row(time, y, amount):
        return [(time, .13, y), ("Pix", .21, y), ("EMPRESA FICTÍCIA", .4, y),
                ("Transferência", .66, y), (amount, .92, y)]

    pages = [header() + [("01 Out, 2026", .04, .20)] + row("10:00", .22, "+50,00")
             + row("11:00", .28, "-20,00") + [("Saldo do dia", .21, .34), ("+ 130,00", .92, .34),
                ("02 Out, 2026", .04, .40)] + row("10:00", .42, "+10,00") + [("Página 11 de 27", .8, .94)],
             header() + [("02 Out, 2026", .04, .20)] + row("11:00", .22, "-5,00")
             + [("Saldo do dia", .21, .28), ("+ 135,00", .92, .28), ("03 Out, 2026", .04, .34)]
             + row("10:00", .36, "+7,00") + [("Página 12 de 27", .8, .94)]]
    replacements = {
        "missing_sign": {"+50,00": "50,00"}, "bad_time": {"10:00": "25:99"},
        "bad_date": {"01 Out, 2026": "32 Out, 2026"},
        "balance_mismatch": {"+ 135,00": "+ 136,00"},
        "not_partial": {"Página 11 de 27": "Página 1 de 2", "Página 12 de 27": "Página 2 de 2"},
    }
    if mutation in replacements:
        pages = [[(replacements[mutation].get(text, text), x, y) for text, x, y in page] for page in pages]
    if mutation == "orphan_amount":
        pages[0].append(("+9,00", .92, .65))
    if mutation == "missing_first_date":
        pages[0] = [cell for cell in pages[0] if cell[0] != "01 Out, 2026"]
    return grid_pdf(pages)


def detail_statement(*, mutation=None):
    header = [("DETALHE DOS MOVIMENTOS", .06, .26), ("Data", .06, .30),
              ("Descrição", .20, .30), ("Valor", .72, .30), ("Saldo", .86, .30)]
    summary = [("EXTRATO DE CONTA", .65, .03), ("Agência: 0001 | Conta: 12345-6", .06, .06),
               ("Período: 01/10/2026 a 31/10/2026", .06, .09), ("Titular: EMPRESA FICTÍCIA", .06, .12),
               ("Saldo inicial", .08, .18), ("Entradas", .30, .18), ("Saídas", .52, .18), ("Saldo final", .74, .18),
               ("R$ 100,00", .08, .21), ("R$ 50,00", .30, .21), ("R$ 30,00", .52, .21), ("R$ 120,00", .74, .21)]
    def row(date, y, description, amount, balance):
        return [(description, .20, y - .011), (date, .06, y), (amount, .71, y), (balance, .85, y),
                ("CONTRAPARTE FICTÍCIA", .20, y + .011)]
    pages = [summary + header + row("02/10/2026", .40, "Pagamento", "-R$ 10,00", "R$ 120,00")
             + row("01/10/2026", .48, "Pix recebido", "R$ 50,00", "R$ 130,00") + [("1/2", .91, .97)],
             header + row("01/10/2026", .40, "Pix enviado", "-R$ 20,00", "R$ 80,00") + [("2/2", .91, .97)]]
    replacements = {
        "balance_mismatch": {"R$ 130,00": "R$ 131,00"},
        "summary_mismatch": {"R$ 30,00": "R$ 31,00"},
        "bad_date": {"02/10/2026": "32/10/2026"},
        "outside_period": {"02/10/2026": "02/11/2026"},
        "missing_summary": {"Saldo inicial": "Controle ausente"},
    }
    if mutation in replacements:
        pages = [[(replacements[mutation].get(text, text), x, y) for text, x, y in page] for page in pages]
    if mutation == "orphan_amount":
        pages[-1].append(("R$ 9,00", .71, .70))
    if mutation == "missing_amount":
        pages[0] = [cell for cell in pages[0] if cell[0] != "-R$ 10,00"]
    if mutation == "partial":
        pages = [[("Página 2 de 3" if text == "1/2" else "Página 3 de 3" if text == "2/2" else text, x, y)
                  for text, x, y in page] for page in pages]
    return grid_pdf(pages)


def test_infinitepay_native_fragment_preserves_partial_days_and_printed_balance_evidence(monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    result = pdf_parser.parse_pdf_transactions(infinitepay_fragment())
    assert [(row.date, row.amount) for row in result.transactions] == [
        ("2026-10-01", 50), ("2026-10-01", -20), ("2026-10-02", 10), ("2026-10-02", -5), ("2026-10-03", 7)]
    assert result.parse_metrics["partial_statement"] == 1
    assert result.parse_metrics["printed_balance_checks"] == 1
    assert result.parse_metrics["balance_consistency_failed"] == 0
    assert "verified_opening_balance" not in result.parse_metrics
    assert "verified_closing_balance" not in result.parse_metrics
    assert [row.running_balance for row in result.canonical_transactions] == [None, 130, None, 135, None]
    assert "partial_statement" in result.canonical_transactions[0].warnings
    assert all(row.source_regions for row in result.canonical_transactions)
    assert result.parse_metrics["textract_attempted"] == 0


@pytest.mark.parametrize("mutation", ["missing_sign", "bad_time", "bad_date", "orphan_amount", "missing_first_date", "not_partial"])
def test_infinitepay_fragment_does_not_invent_missing_evidence(mutation):
    with pytest.raises(PdfFinancialValidationError):
        pdf_parser.parse_pdf_transactions(infinitepay_fragment(mutation=mutation))


def test_infinitepay_fragment_keeps_a_real_daily_balance_difference_visible():
    result = pdf_parser.parse_pdf_transactions(infinitepay_fragment(mutation="balance_mismatch"))
    assert len(result.transactions) == 5
    assert "financial_validation_failed" in result.canonical_transactions[0].warnings


def test_detail_statement_preserves_descending_order_and_both_description_lines(monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    result = pdf_parser.parse_pdf_transactions(detail_statement())
    assert [(row.date, row.amount) for row in result.transactions] == [
        ("2026-10-02", -10), ("2026-10-01", 50), ("2026-10-01", -20)]
    assert [row.description for row in result.transactions] == [
        f"{label} CONTRAPARTE FICTÍCIA" for label in ["Pagamento", "Pix recebido", "Pix enviado"]]
    assert result.parse_metrics["verified_opening_balance"] == 100
    assert result.parse_metrics["verified_closing_balance"] == 120
    assert result.parse_metrics["balance_consistency_failed"] == 0
    assert result.parse_metrics["financial_controls_validated"] == 1
    assert result.parse_metrics["printed_balance_checks"] >= 3
    assert all(row.source_regions for row in result.canonical_transactions)
    assert not any(row.warnings for row in result.canonical_transactions)


@pytest.mark.parametrize("mutation", ["balance_mismatch", "summary_mismatch", "outside_period"])
def test_detail_statement_preserves_readable_rows_and_real_financial_warnings(mutation):
    result = pdf_parser.parse_pdf_transactions(detail_statement(mutation=mutation))
    assert [row.amount for row in result.transactions] == [-10, 50, -20]
    assert "financial_validation_failed" in result.canonical_transactions[0].warnings


@pytest.mark.parametrize("mutation", ["bad_date", "missing_amount", "orphan_amount"])
def test_detail_statement_rejects_ambiguous_or_unpaired_cells(mutation):
    with pytest.raises(PdfFinancialValidationError):
        pdf_parser.parse_pdf_transactions(detail_statement(mutation=mutation))


def test_detail_statement_partial_pagination_is_preserved_for_review():
    result = pdf_parser.parse_pdf_transactions(detail_statement(mutation="partial"))
    assert len(result.transactions) == 3
    assert "partial_statement" in result.canonical_transactions[0].warnings


def test_detail_structure_does_not_claim_an_institution_from_a_counterparty():
    result = pdf_parser.parse_pdf_transactions(detail_statement())
    assert result.layout.layout_name == "account_detail_descending_cash_v1"
    assert result.layout.used_fallback is False


def test_detail_reader_requires_the_financial_summary():
    from app.application.parsers.pdf.native_account_detail import parse_account_detail
    from app.application.parsers.pdf.native_table_source import read_native_table_source

    with pytest.raises(PdfFinancialValidationError):
        parse_account_detail(read_native_table_source(detail_statement(mutation="missing_summary")))


def test_infinitepay_control_and_date_can_continue_on_the_next_page():
    from app.application.parsers.pdf.infinitepay_native import parse_infinitepay_native
    from app.application.parsers.pdf.native_table_source import read_native_table_source

    source = read_native_table_source(infinitepay_fragment())
    page = source.pages[1]
    page = replace(page, lines=[line for line in page.lines if line.text != "02 Out, 2026"])
    statement = parse_infinitepay_native(replace(source, pages=[source.pages[0], page]))
    assert [row.transaction.amount for row in statement.rows] == [50, -20, 10, -5, 7]
    assert statement.partial


def test_fragment_edits_and_regenerated_exports_do_not_invent_statement_balances(tmp_path):
    import json

    from openpyxl import load_workbook

    from app.application.default_conversion_pipeline import build_default_conversion_pipeline
    from app.application.storage_service import TempAnalysisStorage

    analysis = build_default_conversion_pipeline().run(
        filename="fictional.pdf", raw_bytes=infinitepay_fragment(), analysis_id="an_fragment_edits").analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(analysis)
    storage.apply_convert_edits(analysis.analysis_id, [
        {"row_id": "row_1", "date": "2026-10-01", "description": "DESCRIÇÃO CONFERIDA", "credit": 50}])
    folder = tmp_path / analysis.analysis_id
    content = json.loads((folder / "analysis.json").read_text(encoding="utf-8"))
    assert content["opening_balance"] is content["closing_balance"] is None
    assert [row["running_balance"] for row in content["report_transactions"]] == [None, 130, None, 135, None]
    assert "<LEDGERBAL>" not in (folder / "converted.ofx").read_text(encoding="utf-8")
    (folder / "converted.xlsx").unlink()
    storage.get_convert_report_path(analysis.analysis_id, "ofx", bank_code="000")
    workbook = load_workbook(folder / "converted.xlsx", data_only=True)
    summary = dict(workbook["Resumo"].iter_rows(min_row=2, values_only=True))
    assert summary["Saldo anterior"] is summary["Saldo final"] is None
    storage.apply_convert_edits(analysis.analysis_id, [], opening_balance=100, closing_balance=142)
    assert "<BALAMT>142.00" in (folder / "converted.ofx").read_text(encoding="utf-8")


def test_structural_account_detail_does_not_use_counterparty_bank_as_issuer():
    from app.application.bank_identity import resolve_bank_name

    assert resolve_bank_name(layout_inference_name="account_detail_descending_cash_v1",
                             extracted_text="Pix recebido BANCO SANTANDER BRADESCO PAGSEGURO") is None
