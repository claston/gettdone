"""Upload, report, export and worker paths share Sicoob row interpretation."""

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application.parsers.ofx import parse_ofx_transactions
from tests.fixtures.recovery_acceptance import grid_pdf
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_recovery_sweep_ocr import extraction
from tests.test_sicoob_statement_context import modern_cells, sisbr_pending_cells

pytestmark = pytest.mark.real_http


def input_case(kind, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    if kind.startswith("sisbr"):
        raw = grid_pdf(sisbr_pending_cells(idle_balance="151,00C" if kind.endswith("review") else "150,00C"))
        return raw, [50, 200, -20]
    cells = modern_cells(descending=kind == "modern_ocr", mismatch=kind.endswith("review"))
    raw = configure_ocr(monkeypatch, extraction(cells)) if "ocr" in kind else grid_pdf([cells])
    return raw, [50, -10, -10]


@pytest.mark.parametrize("kind", ["modern_native", "modern_ocr", "modern_ocr_review", "sisbr", "sisbr_review"])
def test_sicoob_context_converts_and_exports_without_changing_cash(kind, tmp_path, monkeypatch):
    raw, amounts = input_case(kind, monkeypatch)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-sicoob-context-" + kind}
        response = client.post("/convert", data=owner, files={"file": ("ficticio.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        data = response.json()["analysis"]
        assert [r["amount"] for r in data["preview_transactions"]] == amounts
        assert data["layout_inference_name"].startswith("sicoob_")
        metrics = data["pdf_processing_metrics"]
        assert metrics["financial_controls_validated"] == int(not kind.endswith("review"))
        if kind.endswith("review") or kind == "sisbr":
            assert metrics["export_recommendation"] == "review_recommended"
        if kind.startswith("sisbr"):
            assert "pending_deposits_excluded" in metrics["canonical_warning_types"]
        identifier = data["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200, report.text
            if fmt == "csv":
                rows = csv.DictReader(line for line in StringIO(report.content.decode("utf-8-sig")) if line.strip())
                assert [float(row["amount"]) for row in rows] == amounts
            elif fmt == "ofx":
                assert [row.amount for row in parse_ofx_transactions(report.content)] == amounts
                assert "<BANKID>756" in report.text
            else:
                book = load_workbook(BytesIO(report.content), read_only=True, data_only=True)
                rows = list(book["Conversao"].iter_rows(min_row=2, values_only=True))
                assert [round(float(row[2] or 0) - float(row[3] or 0), 2) for row in rows] == amounts
                book.close()
        assert client.get("/report/an_sicoob_context_nonexistent", params=owner).status_code == 404


@pytest.mark.parametrize("kind", ["modern_ocr", "sisbr_review"])
def test_worker_uses_sicoob_context_and_preserves_review(kind, tmp_path, monkeypatch):
    raw, amounts = input_case(kind, monkeypatch)
    with worker_case(tmp_path, raw) as (client, job):
        data = job["result"]["analysis"]
        assert [row["amount"] for row in data["preview_transactions"]] == amounts
        assert data["pdf_processing_metrics"]["financial_controls_validated"] == int(not kind.endswith("review"))
        report = client.get(f"/convert-report/{data['analysis_id']}", params={"format": "ofx"})
        assert report.status_code == 200
        assert [row.amount for row in parse_ofx_transactions(report.content)] == amounts
