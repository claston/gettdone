"""BB pending-deposit status survives native/OCR HTTP and worker boundaries."""

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application.parsers.ofx import parse_ofx_transactions
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_bb_pending_marker import bb_pending_marker, bb_pending_native

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("mode,mismatch", [("native", False), ("merged_text", False), ("detached", True)])
def test_pending_deposit_is_explained_and_release_is_exported_once(mode, mismatch, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw = (bb_pending_native() if mode == "native"
           else configure_ocr(monkeypatch, bb_pending_marker(mode=mode, mismatch=mismatch)))
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictitious-bb-pending-marker"}
        response = client.post("/convert", data=owner, files={"file": ("bb-ficticio.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert analysis["transactions_total"] == 2
        metrics = analysis["pdf_processing_metrics"]
        assert metrics["non_cash_entries_excluded"] == 1
        assert "pending_deposits_excluded" in metrics["canonical_warning_types"]
        if mismatch:
            assert metrics["financial_controls_validated"] == 0
            assert "financial_validation_failed" in metrics["canonical_warning_types"]
        identifier = analysis["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200
            if fmt == "csv":
                rows = csv.DictReader(line for line in StringIO(report.content.decode("utf-8-sig")) if line.strip())
                amounts = [float(row["amount"]) for row in rows]
            elif fmt == "ofx":
                amounts = [row.amount for row in parse_ofx_transactions(report.content)]
            else:
                book = load_workbook(BytesIO(report.content), read_only=True, data_only=True)
                amounts = [round(float(row[2] or 0)-float(row[3] or 0), 2)
                           for row in book["Conversao"].iter_rows(min_row=2, values_only=True)]
                book.close()
            assert amounts == [50, 200]
        assert client.get("/report/an_bb_pending_nonexistent", params=owner).status_code == 404
        assert client.post("/analyze").status_code == 405


def test_worker_keeps_pending_warning_and_available_cash_movements(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, bb_pending_native()) as (client, job):
        analysis = job["result"]["analysis"]
        assert [t["amount"] for t in analysis["preview_transactions"]] == [50, 200]
        assert analysis["pdf_processing_metrics"]["non_cash_entries_excluded"] == 1
        assert "pending_deposits_excluded" in analysis["pdf_processing_metrics"]["canonical_warning_types"]
        response = client.get(f"/convert-report/{analysis['analysis_id']}", params={"format": "ofx"})
        assert response.status_code == 200
        assert [row.amount for row in parse_ofx_transactions(response.content)] == [50, 200]
