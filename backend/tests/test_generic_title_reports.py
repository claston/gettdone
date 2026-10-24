"""One report model supports payables/receivables without bank projections."""
import csv
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application import pdf_parser
from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.document_type_classifier import classify_document_type
from app.application.errors import InvalidFileContentError
from app.application.storage_service import TempAnalysisStorage
from tests.fixtures.generic_titles import ROWS, title_ocr, title_pdf


@pytest.fixture(autouse=True)
def no_cloud(monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")


@pytest.mark.parametrize("kind,expected", [("payable", "accounts_payable_report"), ("receivable", "accounts_receivable_report")])
@pytest.mark.parametrize("reordered", [False, True])
@pytest.mark.parametrize("ocr", [False, True])
def test_generic_columns_preserve_titles_and_never_create_cash_effect(kind, expected, reordered, ocr):
    raw = title_pdf(kind, reordered=reordered)
    parsed = (pdf_parser._parse_scanned_pdf_with_textract_text_pages(title_ocr(kind, reordered=reordered))
              if ocr else pdf_parser.parse_pdf_transactions(raw))
    assert classify_document_type("report.pdf", raw).document_type == expected
    report = parsed.financial_document
    assert report["document_type"] == expected
    assert report["export"]["available_formats"] == ["csv", "xlsx"]
    first = report["events"][0]
    assert first["reference"] == "0001/01"
    assert first["counterparty"] == "EMPRESA FICTÍCIA"
    assert first["dates"]["due"] == "2026-07-10"
    assert first["amounts"]["original"] == first["amounts"]["outstanding"] == "100.00"
    assert first["amount"] == "100.00" and first["cash_effect"] is None
    assert first["attributes"]["source_fields"]
    assert all(row.type == "non_cash" and row.source_regions for row in parsed.canonical_transactions)
    data = build_default_conversion_pipeline().run(filename="report.pdf", raw_bytes=raw, analysis_id="an_generic_titles",
                                                  pdf_parser=lambda *a, **kw: parsed).analysis_data
    assert data.document_type == expected
    assert data.total_inflows == data.total_outflows == data.net_total == 0
    assert data.matched_groups == data.reversed_entries == data.potential_duplicates == 0
    assert all(row.reconciliation_status == "not_applicable" for row in data.report_transactions)
    assert data.opening_balance is data.closing_balance is data.bank_code is None


def test_generic_report_exports_original_fields_csv_excel_and_explains_ofx_refusal(tmp_path):
    data = build_default_conversion_pipeline().run(
        filename="report.pdf", raw_bytes=title_pdf(), analysis_id="an_titles_exports").analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    rows = list(csv.DictReader(storage.get_convert_report_path(data.analysis_id, "csv").read_text(encoding="utf-8-sig").splitlines()))
    assert len(rows) == 2 and rows[0]["reference"] == "0001/01"
    assert rows[0]["amounts.outstanding"] == "100.00"
    book = load_workbook(BytesIO(storage.get_convert_report_path(data.analysis_id, "xlsx").read_bytes()))
    assert book["Eventos"].max_row == 3
    assert any(cell.value == "Saldo em aberto" for cell in book["Eventos"][1])
    with pytest.raises(InvalidFileContentError, match="OFX|bancári"):
        storage.get_convert_report_path(data.analysis_id, "ofx")


def test_total_difference_is_review_warning_without_dropping_or_changing_titles():
    result = pdf_parser.parse_pdf_transactions(title_pdf(total="181,00"))
    assert [event["amount"] for event in result.financial_document["events"]] == ["100.00", "80.00"]
    assert any(v["status"] == "failed" for v in result.financial_document["validations"])
    assert result.parse_metrics["export_recommendation"] == "review_recommended"


def test_duplicate_titles_and_page_provenance_are_preserved():
    parsed = pdf_parser.parse_pdf_transactions(title_pdf(rows=[ROWS[0], ROWS[0]], total="200,00", split_pages=True))
    assert len(parsed.financial_document["events"]) == 2
    assert [row.source_page for row in parsed.canonical_transactions] == [1, 2]


def test_unreadable_required_amount_does_not_fall_back_to_bank_parser():
    rows = [dict(ROWS[0], original="ilegível"), ROWS[1]]
    with pytest.raises(InvalidFileContentError):
        pdf_parser.parse_pdf_transactions(title_pdf(rows=rows))


@pytest.mark.parametrize("header", ["Extrato de Conta Corrente", "Extrato consolidado inteligente", "Extrato Santander empresarial"])
def test_mention_of_receivables_inside_bank_statement_does_not_change_its_type(header):
    text = header+"\nData Histórico Valor Saldo\n01/07/2026 PIX CONTAS A RECEBER 10,00 C 100,00 C\nNotas a Pagar"
    assert classify_document_type("bank.pdf", b"", extracted_text=text).document_type != "accounts_receivable_report"
    assert classify_document_type("bank.pdf", b"", extracted_text=text).document_type != "accounts_payable_report"


def test_mixed_payable_and_receivable_pages_are_not_given_a_single_wrong_direction():
    payable, receivable = title_ocr(), title_ocr("receivable")
    source = replace(payable, pages=[payable.pages[0], receivable.pages[0]])
    with pytest.raises(InvalidFileContentError, match="mistura"):
        pdf_parser._parse_scanned_pdf_with_textract_text_pages(source)


@pytest.mark.parametrize("title", ["Notas a Pagar", "Contas a Receber"])
def test_recognized_title_report_without_columns_cannot_fall_back_to_bank_rows(title):
    with pytest.raises(InvalidFileContentError, match="relatório de contas"):
        pdf_parser._parse_pdf_transactions_from_page_texts([title+"\n01/07/2026 FORNECEDOR 10,00 100,00"])


def test_unicode_case_expansion_in_supplier_does_not_shift_native_dates_or_values():
    rows = [dict(ROWS[0], party="Straße Serviços"), ROWS[1]]
    result = pdf_parser.parse_pdf_transactions(title_pdf(rows=rows))
    assert [event["dates"]["due"] for event in result.financial_document["events"]] == ["2026-07-10", "2026-07-11"]
    assert result.financial_document["events"][0]["counterparty"] == "Straße Serviços"


def test_supplier_starting_with_total_is_a_dated_title_not_a_summary():
    source = title_ocr(rows=[dict(ROWS[0], party="TOTAL SERVIÇOS"), ROWS[1]])
    result = pdf_parser._parse_scanned_pdf_with_textract_text_pages(source)
    assert len(result.financial_document["events"]) == 2
    assert result.financial_document["events"][0]["counterparty"] == "TOTAL SERVIÇOS"


@pytest.mark.parametrize("native", [False, True])
def test_unknown_header_is_preserved_as_a_column_without_contaminating_counterparty(native, monkeypatch):
    source = title_ocr()
    page = source.pages[0]
    sample = page.lines[0]
    extra = [replace(sample, id=f"extra-{index}", line_index=500+index, text=text,
                     bbox={"left": .7, "top": y, "width": .04, "height": .006})
             for index, (text, y) in enumerate([("Projeto", .12), ("P0001", .2), ("P0002", .28)])]
    source = replace(source, pages=[replace(page, lines=page.lines+extra)])
    if native:
        from tests.fixtures import generic_titles

        monkeypatch.setattr(generic_titles, "title_ocr", lambda *a, **kw: source)
        result = pdf_parser.parse_pdf_transactions(generic_titles.title_pdf())
    else:
        result = pdf_parser._parse_scanned_pdf_with_textract_text_pages(source)
    event = result.financial_document["events"][0]
    assert event["counterparty"] == "EMPRESA FICTÍCIA"
    assert event["attributes"]["source_fields"]["Projeto"] == "P0001"
