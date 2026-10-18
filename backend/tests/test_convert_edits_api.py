from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.application import AccessControlService, ReportService, TempAnalysisStorage
from app.application.models import AnalysisData, TransactionRow
from app.dependencies import get_access_control_service, get_report_service
from app.main import app


def _build_analysis_data(analysis_id: str = "an_convert123", *, first_amount: float = -20.0) -> AnalysisData:
    preview = [
        TransactionRow(
            date="2026-04-01",
            description="DEBIT A",
            amount=first_amount,
            category="Outros",
            reconciliation_status="unmatched",
            running_balance=80.0,
        ),
        TransactionRow(
            date="2026-04-02",
            description="CREDIT B",
            amount=100.0,
            category="Outros",
            reconciliation_status="unmatched",
            running_balance=180.0,
        ),
    ]
    return AnalysisData(
        analysis_id=analysis_id,
        file_type="pdf",
        upload_filename="extrato_nubank.pdf",
        transactions_total=2,
        total_inflows=100.0,
        total_outflows=-20.0,
        net_total=80.0,
        preview_transactions=preview,
        report_transactions=preview,
    )


def build_client(tmp_path: Path, *, first_amount: float = -20.0) -> TestClient:
    storage = TempAnalysisStorage(root_dir=tmp_path, ttl_seconds=3600)
    storage.save_analysis(_build_analysis_data(first_amount=first_amount))
    access_control = AccessControlService(
        state_file=tmp_path / "access-control-state.json",
        token_secret="test-secret",
    )
    owner = access_control.resolve_identity(anonymous_fingerprint="fp-owner", user_token=None)
    storage.set_convert_owner(
        analysis_id="an_convert123",
        identity_type=owner.identity_type,
        identity_id=owner.identity_id,
    )
    app.dependency_overrides[get_report_service] = lambda: ReportService(storage=storage)
    app.dependency_overrides[get_access_control_service] = lambda: access_control
    return TestClient(app)


def build_client_with_user_owner(tmp_path: Path) -> tuple[TestClient, str]:
    storage = TempAnalysisStorage(root_dir=tmp_path, ttl_seconds=3600)
    storage.save_analysis(_build_analysis_data("an_user_owned"))
    access_control = AccessControlService(
        state_file=tmp_path / "access-control-state.json",
        token_secret="test-secret",
    )
    user = access_control.register_user(name="Erica", email="erica@example.com", password="strong-pass")
    storage.set_convert_owner(
        analysis_id="an_user_owned",
        identity_type="user",
        identity_id=user.user_id,
    )
    app.dependency_overrides[get_report_service] = lambda: ReportService(storage=storage)
    app.dependency_overrides[get_access_control_service] = lambda: access_control
    return TestClient(app), user.token


def test_edit_operation_confirms_and_replays_insert_without_duplication(tmp_path: Path) -> None:
    client = build_client(tmp_path)
    url = "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner"
    request = {"operation_id": "save-insert-123", "edits": [{
        "action": "insert", "insert_position": 0, "date": "2026-04-03", "description": "NOVO", "credit": 12,
    }]}
    try:
        first = client.post(url, json=request)
        assert first.status_code == 200
        repeated = client.post(url, json=request)
        assert repeated.status_code == 200
        assert repeated.json() == first.json()
        assert repeated.json()["transactions_total"] == 3
        status = client.get("/convert-edits/an_convert123/operations/save-insert-123?anonymous_fingerprint=fp-owner")
        assert status.status_code == 200
        assert status.json() == first.json()
        pro_read = client.get("/api/pro/convert-edits/an_convert123/operations/save-insert-123?anonymous_fingerprint=fp-owner")
        assert pro_read.status_code == 200
        assert pro_read.json() == first.json()
        denied_pro_read = client.get(
            "/api/pro/convert-edits/an_convert123/operations/save-insert-123?anonymous_fingerprint=fp-other")
        assert denied_pro_read.status_code == 403
        read_with_body = client.request("GET",
            "/convert-edits/an_convert123/operations/save-insert-123?anonymous_fingerprint=fp-owner",
            json={"edits": [{**request["edits"][0], "credit": 99}]})
        assert read_with_body.json() == first.json()
        changed = {**request, "edits": [{**request["edits"][0], "credit": 99}]}
        assert client.post(url, json=changed).status_code == 409
        assert client.get("/convert-edits/an_convert123/operations/save-insert-123?anonymous_fingerprint=fp-other").status_code == 403
        missing = client.get("/convert-edits/an_convert123/operations/not-saved?anonymous_fingerprint=fp-owner")
        assert missing.status_code == 404
        assert missing.json()["detail"] == "Edit operation not found."
    finally:
        app.dependency_overrides.clear()


def test_convert_edits_updates_preview_and_csv_artifact(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "EDITED CREDIT",
                    "credit": 45.75,
                    "debit": None,
                }
            ]
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["processing_id"] == "an_convert123"
    assert payload["preview_transactions"][0]["description"] == "EDITED CREDIT"
    assert payload["preview_transactions"][0]["amount"] == 45.75

    csv_report = client.get("/convert-report/an_convert123?format=csv&anonymous_fingerprint=fp-owner")
    assert csv_report.status_code == 200
    assert "2026-04-05,EDITED CREDIT,45.75" in csv_report.text
    assert isinstance(payload["updated_at"], str)
    app.dependency_overrides.clear()


def test_convert_edits_can_correct_description_of_an_imported_zero_amount(tmp_path: Path) -> None:
    client = build_client(tmp_path, first_amount=0.0)
    try:
        response = client.post("/convert-edits/an_convert123?anonymous_fingerprint=fp-owner", json={"edits": [{
            "row_id": "row_1", "date": "2026-04-01", "description": "PIX VALIDADO", "credit": None, "debit": None,
        }]})
        assert response.status_code == 200
        assert response.json()["preview_transactions"][0]["amount"] == 0.0
        assert response.json()["preview_transactions"][0]["description"] == "PIX VALIDADO"
    finally:
        app.dependency_overrides.clear()


def test_pro_save_syncs_batch_snapshot_and_rejects_mismatched_job(tmp_path: Path, monkeypatch) -> None:
    from app.adapters.conversion.memory_batches import InMemoryConversionBatchRepository
    from app.application.conversion.contracts.documents import ConversionDocumentReference
    from app.application.conversion.contracts.jobs import ConversionJobResultReference
    from app.application.conversion.conversion_batch import ConversionBatch
    from app.application.conversion.conversion_job import ConversionJob
    from app.dependencies import get_conversion_batch_service

    client = build_client(tmp_path)
    monkeypatch.setenv("CORS_ALLOW_ORIGINS", "https://pro.example.test")
    headers = {"X-OFX-Pro-Client": "web", "Origin": "https://pro.example.test"}
    access = app.dependency_overrides[get_access_control_service]()
    owner = access.resolve_identity(anonymous_fingerprint="fp-owner", user_token=None)
    batch = ConversionBatch.create(identity_type=owner.identity_type, identity_id=owner.identity_id,
                                   files_count=1, idempotency_key="sync-test")
    job = ConversionJob.create(batch_id=batch.batch_id, identity=owner, idempotency_key="sync-job",
                               document=ConversionDocumentReference(storage_key="doc_sync", filename="extrato.pdf",
                                                                    file_type="pdf", size_bytes=100, sha256_hex="0" * 64))
    repository = InMemoryConversionBatchRepository()
    repository.create(batch=batch, jobs=[job])
    repository.mark_uploaded(job.job_id)
    repository.mark_queued(job.job_id)
    repository.mark_running(job.job_id)
    repository.mark_completed(job.job_id, result=ConversionJobResultReference(
        analysis_id="an_convert123", payload={"analysis": {"bank_name": "Banco Teste"}},
    ))
    app.dependency_overrides[get_conversion_batch_service] = lambda: SimpleNamespace(repository=repository)
    request = {"batch_id": batch.batch_id, "job_id": job.job_id, "edits": [{
        "row_id": "row_1", "date": "2026-04-01", "description": "SALVO NO LOTE", "credit": 40,
    }]}
    try:
        wrong_job = client.post("/api/pro/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
                                headers=headers, json={**request, "job_id": "job_other"})
        assert wrong_job.status_code == 404
        response = client.post("/api/pro/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
                               headers=headers, json=request)
        assert response.status_code == 200
        snapshot = repository.get_for_owner(batch.batch_id, identity_type=owner.identity_type, identity_id=owner.identity_id)
        analysis = snapshot.jobs[0].result.payload["analysis"]
        assert analysis["bank_name"] == "Banco Teste"
        assert analysis["updated_at"] == response.json()["updated_at"]
        assert analysis["preview_transactions"][0]["description"] == "SALVO NO LOTE"
        request["expected_updated_at"] = analysis["updated_at"]
        request["edits"][0]["credit"] = 45
        assert client.post("/api/pro/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
                           headers=headers, json=request).status_code == 200
    finally:
        app.dependency_overrides.clear()


def test_pro_convert_edits_requires_browser_guard_and_preserves_owner_scope(
    tmp_path: Path,
    monkeypatch,
) -> None:
    client = build_client(tmp_path)
    monkeypatch.setenv("CORS_ALLOW_ORIGINS", "https://pro.example.test")
    payload = {
        "edits": [
            {
                "row_id": "row_1",
                "date": "2026-04-05",
                "description": "EDITED FROM PRO",
                "credit": 45.75,
                "debit": None,
            }
        ]
    }

    missing_guard = client.post(
        "/api/pro/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json=payload,
    )
    accepted = client.post(
        "/api/pro/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        headers={
            "Origin": "https://pro.example.test",
            "X-OFX-Pro-Client": "web",
        },
        json=payload,
    )

    assert missing_guard.status_code == 403
    assert accepted.status_code == 200
    assert accepted.json()["preview_transactions"][0]["description"] == "EDITED FROM PRO"
    app.dependency_overrides.clear()


def test_convert_edits_recalculates_running_balances_after_edit_and_xlsx_export(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "EDITED DEBIT",
                    "credit": None,
                    "debit": 10.0,
                }
            ]
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["preview_transactions"][0]["amount"] == -10.0
    assert payload["preview_transactions"][0]["running_balance"] == 90.0
    assert payload["preview_transactions"][1]["running_balance"] == 190.0
    assert payload["closing_balance"] == 190.0

    xlsx_report = client.get("/convert-report/an_convert123?format=xlsx&anonymous_fingerprint=fp-owner")
    assert xlsx_report.status_code == 200
    workbook = load_workbook(filename=BytesIO(xlsx_report.content), data_only=True)
    sheet = workbook["Conversao"]
    summary = workbook["Resumo"]
    assert sheet.cell(row=2, column=5).value == 90
    assert sheet.cell(row=3, column=5).value == 190
    assert summary.cell(row=3, column=2).value == 190
    app.dependency_overrides.clear()


def test_convert_edits_rejects_invalid_credit_debit_pair(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "INVALID",
                    "credit": 10.0,
                    "debit": 10.0,
                }
            ]
        },
    )

    assert response.status_code == 400
    assert "Provide only one of credit or debit" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_convert_edits_deletes_row_and_updates_csv(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "action": "delete",
                }
            ]
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["transactions_total"] == 1
    assert len(payload["preview_transactions"]) == 2
    assert payload["preview_transactions"][0]["description"] == "DEBIT A"
    assert payload["preview_transactions"][0]["is_deleted"] is True
    assert payload["preview_transactions"][1]["description"] == "CREDIT B"
    assert payload["preview_transactions"][1]["is_deleted"] is False

    csv_report = client.get("/convert-report/an_convert123?format=csv&anonymous_fingerprint=fp-owner")
    assert csv_report.status_code == 200
    assert "DEBIT A" not in csv_report.text
    assert "CREDIT B" in csv_report.text
    app.dependency_overrides.clear()


def test_convert_edits_inserts_row_and_updates_csv(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "action": "insert",
                    "insert_position": 2,
                    "date": "2026-04-03",
                    "description": "INSERTED C",
                    "credit": 30.0,
                    "debit": None,
                }
            ]
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["transactions_total"] == 3
    assert len(payload["preview_transactions"]) == 3
    assert payload["preview_transactions"][2]["description"] == "INSERTED C"
    assert payload["preview_transactions"][2]["amount"] == 30.0
    assert payload["preview_transactions"][2]["is_deleted"] is False

    csv_report = client.get("/convert-report/an_convert123?format=csv&anonymous_fingerprint=fp-owner")
    assert csv_report.status_code == 200
    assert "2026-04-03,INSERTED C,30.0" in csv_report.text
    app.dependency_overrides.clear()


def test_convert_edits_returns_not_found_for_missing_analysis(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_missing?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "ANY",
                    "credit": 10.0,
                    "debit": None,
                }
            ]
        },
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Analysis not found"
    app.dependency_overrides.clear()


def test_convert_edits_rejects_access_from_other_identity(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    response = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-other",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "ANY",
                    "credit": 10.0,
                    "debit": None,
                }
            ]
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Access denied for this analysis."
    app.dependency_overrides.clear()


def test_convert_edits_returns_conflict_for_stale_version(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    first = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "FIRST",
                    "credit": 10.0,
                    "debit": None,
                }
            ]
        },
    )
    assert first.status_code == 200
    stale_version = first.json()["updated_at"]

    second = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "expected_updated_at": stale_version,
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-06",
                    "description": "SECOND",
                    "credit": 15.0,
                    "debit": None,
                }
            ],
        },
    )
    assert second.status_code == 200

    conflict = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "expected_updated_at": stale_version,
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-07",
                    "description": "STALE",
                    "credit": 20.0,
                    "debit": None,
                }
            ],
        },
    )
    assert conflict.status_code == 409
    assert "changed since last load" in conflict.json()["detail"]
    app.dependency_overrides.clear()


def test_convert_edits_accepts_bearer_user_token(tmp_path: Path) -> None:
    client, user_token = build_client_with_user_owner(tmp_path)

    response = client.post(
        "/convert-edits/an_user_owned",
        headers={"authorization": f"Bearer {user_token}"},
        json={
            "edits": [
                {
                    "row_id": "row_1",
                    "date": "2026-04-05",
                    "description": "EDITED VIA HEADER",
                    "credit": 45.75,
                    "debit": None,
                }
            ]
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["preview_transactions"][0]["description"] == "EDITED VIA HEADER"
    app.dependency_overrides.clear()


def test_convert_edits_persists_ofx_export_settings_for_download_without_query(tmp_path: Path) -> None:
    client = build_client(tmp_path)

    save_settings = client.post(
        "/convert-edits/an_convert123?anonymous_fingerprint=fp-owner",
        json={
            "edits": [],
            "opening_balance": 73997.11,
            "closing_balance": 56276.06,
            "bank_branch": "3456-7",
            "account_number": "12345-6",
            "bank_code": "001",
        },
    )
    assert save_settings.status_code == 200
    assert save_settings.json()["opening_balance"] == 73997.11

    ofx_report = client.get("/convert-report/an_convert123?format=ofx&anonymous_fingerprint=fp-owner")
    assert ofx_report.status_code == 200
    assert "<LEDGERBAL>" in ofx_report.text
    assert "<BALAMT>56276.06" in ofx_report.text
    assert "<BRANCHID>34567" in ofx_report.text
    assert "<ACCTID>123456" in ofx_report.text
    assert "<BANKID>001" in ofx_report.text
    app.dependency_overrides.clear()
