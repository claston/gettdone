from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application import pdf_parser
from synthetic_pdf_corpus.native_text_pdf import generate_native_text_pdf
from tests.ocr_conversion_http_helpers import public_client, worker_case
from tests.test_oct20_new_native_statements import banrisul_page, bb_page, xp_page
from tests.test_oct20_ocr_controls import itau_receipt
from tests.test_pdf_representation_classifier import _full_page_raster_pdf
from tests.test_sicoob_investment_summary import summary_page

pytestmark = pytest.mark.real_http


def native_case(kind):
    text = {"bb": bb_page, "xp": xp_page, "banrisul": banrisul_page, "sicoob": summary_page}[kind]()
    return generate_native_text_pdf((tuple(text.splitlines()),))


@pytest.mark.parametrize("kind,amounts,bank", [
    ("bb", [-194.57, 1100], "001"), ("xp", [500, -300], "348"),
    ("banrisul", [200, -50, -300], "041"), ("sicoob", [], None),
])
def test_new_native_forms_http_preserve_semantics_reports_and_exports(kind, amounts, bank, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-newcases-" + kind}
        response = client.post("/convert", data=owner, files={"file": (kind + ".pdf", native_case(kind), "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert [t["amount"] for t in analysis["preview_transactions"]] == amounts
        assert analysis["bank_code"] == bank
        assert analysis["pdf_processing_metrics"]["textract_used"] == 0
        analysis_id = analysis["analysis_id"]
        assert client.get(f"/report/{analysis_id}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{analysis_id}", params={**owner, "format": fmt})
            assert report.status_code == (400 if kind == "sicoob" and fmt == "ofx" else 200), report.text
            if fmt == "ofx" and kind != "sicoob":
                assert report.text.count("<STMTTRN>") == len(amounts)
                assert f"<BANKID>{bank}" in report.text
            if kind == "sicoob" and fmt == "xlsx":
                book = load_workbook(BytesIO(report.content))
                assert book["Posições"].max_row == 12
        if kind == "sicoob":
            assert len(analysis["financial_document"]["observations"]) == 11
            assert analysis["total_inflows"] == analysis["total_outflows"] == 0
        assert client.get("/report/an_newcases_nonexistent", params=owner).status_code == 404


def test_day_only_itau_ocr_http_uses_full_table_and_controls(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "true")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    def cached_gateway(raw):
        return pdf_parser._with_parse_observability(
            pdf_parser._parse_scanned_pdf_with_textract_text_pages(itau_receipt()),
            textract_used=1, extraction_provider="aws_textract")
    monkeypatch.setattr(pdf_parser, "_parse_scanned_pdf_with_textract_gateway", cached_gateway)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-newcases-ocr"}
        response = client.post("/convert", data=owner,
            files={"file": ("itau.pdf", _full_page_raster_pdf(), "application/pdf")})
        assert response.status_code == 200, response.text
        a = response.json()["analysis"]
        assert [t["amount"] for t in a["preview_transactions"]] == [-100, 50]
        assert (a["opening_balance"], a["closing_balance"]) == (1000, 950)
        assert a["pdf_processing_metrics"]["textract_used"] == 1
        report = client.get(f"/convert-report/{a['analysis_id']}", params={**owner, "format": "ofx"})
        assert report.status_code == 200 and report.text.count("<STMTTRN>") == 2


@pytest.mark.parametrize("kind", ["bb", "sicoob"])
def test_worker_uses_new_statement_and_observation_contracts(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, native_case(kind), filename=kind + ".pdf") as (client, job):
        a = job["result"]["analysis"]
        assert a["transactions_total"] == (0 if kind == "sicoob" else 2)
        assert client.get(f"/convert-report/{a['analysis_id']}", params={"format": "xlsx"}).status_code == 200


def test_missing_unsigned_columns_are_rejected_instead_of_exporting_credits(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw = generate_native_text_pdf((("02/06/25", "9079938", "APLIC. INVEST FACIL", "500,00",
                                     "02/06/25", "0001801", "PAGTO ELETRON COBRANCA EXEMPLO", "100,00"),))
    with public_client(tmp_path) as client:
        response = client.post("/convert", data={"anonymous_fingerprint": "fp-newcases-invalid"},
                               files={"file": ("unsigned.pdf", raw, "application/pdf")})
        assert response.status_code == 400, response.text
        assert "débito" in response.text or "sinal" in response.text
