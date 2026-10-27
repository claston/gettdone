"""Headerless OCR ledger fragments with signed movements and daily controls."""
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application import pdf_parser
from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.errors import PdfFinancialValidationError
from tests.ocr_conversion_http_helpers import configure_ocr, public_client
from tests.test_oct20_followup_ocr import ocr


def daily_cd_fragment(*, mutation=None):
    pages = [[("SALDO DO DIA", .32, .05), ("===== >", .49, .05), ("100,00C", .82, .05),
        ("02/01/2025 000001", .12, .15), ("DÉBITO FICTÍCIO", .32, .15), ("30,00D", .82, .15),
        ("02/01/2025 000001", .12, .25), ("ESTORNO DÉBITO FICTÍCIO", .32, .25), ("10,00C", .82, .25),
        ("SALDO DO DIA >", .32, .35), ("80,00C", .82, .35),
        ("03/01/2025", .12, .50), ("000002", .23, .50), ("TARIFA FICTÍCIA", .32, .50),
        ("5,00D", .82, .50), ("Descrição complementar", .32, .52)],
        [("SALDO DO DIA", .32, .05), ("75,00C", .82, .05),
        ("06/01/2025 Pix", .12, .15), ("TRANSF.RECEBIDA - PIX SICOOB", .32, .15), ("20,00C", .82, .15),
        ("Transferência Pix", .32, .17), ("CONTRAPARTE FICTÍCIA", .32, .19),
        ("SALDO DO DIA", .32, .30), ("95,00C", .82, .30),
        ("07/01/2025 000003", .12, .45), ("DÉBITO FICTÍCIO", .32, .45), ("4,00D", .82, .45)]]
    replacements = {"missing_sign": {"30,00D": "30,00"}, "illegible_amount": {"30,00D": "3X,00D"},
        "bad_date": {"02/01/2025 000001": "32/01/2025 000001"},
        "descending_date": {"07/01/2025 000003": "01/01/2025 000003"},
        "wrong_daily_balance": {"80,00C": "81,00C"},
        "missing_opening_sign": {"100,00C": "100,00"}}
    pages = [[(replacements.get(mutation, {}).get(text, text), x, y) for text, x, y in page] for page in pages]
    if mutation == "missing_date":
        pages[0] = [cell for cell in pages[0] if cell[0] != "03/01/2025"]
    if mutation == "missing_amount":
        pages[0] = [cell for cell in pages[0] if cell[0] != "5,00D"]
    if mutation == "orphan_amount":
        pages[1].append(("7,00D", .82, .60))
    if mutation == "missing_description":
        pages[0] = [cell for cell in pages[0] if cell[0] != "DÉBITO FICTÍCIO"]
    if mutation == "blank_page":
        pages.insert(1, [("Página ilegível", .30, .30)])
    return ocr(pages)


def test_headerless_daily_ledger_preserves_every_movement_and_printed_control():
    result = pdf_parser._parse_scanned_pdf_with_textract_text_pages(daily_cd_fragment())
    assert [(row.date, row.amount) for row in result.transactions] == [
        ("2025-01-02", -30), ("2025-01-02", 10), ("2025-01-03", -5), ("2025-01-06", 20), ("2025-01-07", -4)]
    assert result.layout.layout_name == "dated_cd_daily_balance_fragment_ocr_v1"
    assert [row.running_balance for row in result.canonical_transactions] == [None, 80, 75, 95, None]
    assert [row.external_reference_id for row in result.canonical_transactions] == ["000001", "000001", "000002", "Pix", "000003"]
    assert result.transactions[2].description == "TARIFA FICTÍCIA Descrição complementar"
    assert result.parse_metrics["partial_statement"] == 1
    assert result.parse_metrics["balance_bounds_unavailable"] == 1
    assert result.parse_metrics["printed_balance_checks"] == 3
    assert result.parse_metrics["statement_rows_expected"] == result.parse_metrics["statement_rows_parsed"] == 5
    assert result.parse_metrics["balance_consistency_failed"] == 0
    assert all(row.source_regions for row in result.canonical_transactions)
    assert "partial_statement" in result.canonical_transactions[0].warnings


@pytest.mark.parametrize("mutation", ["missing_sign", "illegible_amount", "bad_date", "descending_date", "missing_date",
    "missing_amount", "orphan_amount", "missing_description", "blank_page", "missing_opening_sign"])
def test_fragment_never_swallows_unexplained_or_incomplete_financial_rows(mutation):
    with pytest.raises(PdfFinancialValidationError):
        pdf_parser._parse_scanned_pdf_with_textract_text_pages(daily_cd_fragment(mutation=mutation))


def test_readable_real_balance_difference_keeps_the_movements_and_review():
    result = pdf_parser._parse_scanned_pdf_with_textract_text_pages(daily_cd_fragment(mutation="wrong_daily_balance"))
    assert [row.amount for row in result.transactions] == [-30, 10, -5, 20, -4]
    assert "financial_validation_failed" in result.canonical_transactions[0].warnings
    assert "partial_statement" in result.canonical_transactions[0].warnings


def test_counterparty_does_not_assign_a_bank_or_global_balances(monkeypatch):
    raw = configure_ocr(monkeypatch, daily_cd_fragment())
    a = build_default_conversion_pipeline().run(filename="fictional.pdf", raw_bytes=raw,
        analysis_id="an_daily_fragment_fictional").analysis_data
    assert a.transactions_total == 5
    assert a.bank_name is None
    assert a.bank_code in (None, "000")
    assert a.opening_balance is a.closing_balance is None


@pytest.mark.real_http
@pytest.mark.parametrize("mode", ["text", "tables"])
def test_fragment_http_and_all_exports_preserve_partial_scope(tmp_path, monkeypatch, mode):
    raw = configure_ocr(monkeypatch, daily_cd_fragment(), mode=mode)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-daily-fragment"}
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        payload = response.json()
        a = payload["analysis"]
        identifier = payload["processing_id"]
        assert a["transactions_total"] == 5
        assert a["opening_balance"] is a["closing_balance"] is None
        assert a["pdf_processing_metrics"]["partial_statement"] == 1
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        xlsx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "xlsx"})
        assert xlsx.status_code == 200
        wb = load_workbook(BytesIO(xlsx.content), data_only=True)
        summary = dict(wb["Resumo"].iter_rows(min_row=2, values_only=True))
        assert summary["Saldo anterior"] is summary["Saldo final"] is None
        assert [row[4] for row in wb["Conversao"].iter_rows(min_row=2, values_only=True)] == [None, 80, 75, 95, None]
        ofx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx"})
        assert ofx.status_code == 200 and ofx.text.count("<STMTTRN>") == 5
        assert "<LEDGERBAL>" not in ofx.text
        csv = client.get(f"/convert-report/{identifier}", params={**owner, "format": "csv"})
        assert csv.status_code == 200 and len([line for line in csv.text.splitlines() if line.strip()]) == 6
        assert client.get("/report/an_fragment_nonexistent", params=owner).status_code == 404


def test_unrelated_document_with_a_daily_balance_phrase_does_not_select_the_reader():
    from app.application.parsers.pdf.daily_cd_fragment_ocr import is_daily_cd_fragment

    unrelated = ocr([[("RELATÓRIO FICTÍCIO", .10, .03), ("SALDO DO DIA", .32, .10),
                      ("100,00C", .82, .10), ("02/01/2025", .12, .20)]])
    assert not is_daily_cd_fragment(unrelated)


def test_fragment_with_missing_coordinates_is_not_parsed_as_a_verified_ledger():
    source = daily_cd_fragment()
    source = replace(source, pages=[replace(page, lines=[replace(line, bbox=None) for line in page.lines]) for page in source.pages])
    from app.application.parsers.pdf.daily_cd_fragment_ocr import is_daily_cd_fragment

    assert not is_daily_cd_fragment(source)
