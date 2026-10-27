"""Final-priority examples cross upload, report, download and worker boundaries."""
from functools import partial
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.ocr_conversion_http_helpers import public_client, worker_case
from tests.test_oct27_final_priority import pagbank_daily, santander_limited, sicoob_split_signs

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("factory,amounts,bounds,bank", [
    (pagbank_daily, [20, -10, -10, 5], (100, 105), "290"),
    (partial(pagbank_daily, multiple_first_day=True), [20, 10, -10, -10, 5], (100, 115), "290"),
    (santander_limited, [5, 10], (None, 105), "033"),
    (sicoob_split_signs, [50, -200], (100, -50), "756"),
])
def test_final_priority_http_preserves_amounts_controls_and_downloads(tmp_path, monkeypatch, factory, amounts, bounds, bank):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-final-priority"}
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", factory(), "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert analysis["transactions_total"] == len(amounts)
        assert (analysis["opening_balance"], analysis["closing_balance"], analysis["bank_code"]) == (*bounds, bank)
        metrics = analysis["pdf_processing_metrics"]
        assert metrics["partial_statement"] == int(factory is santander_limited)
        assert metrics["financial_controls_validated"] == int(factory is sicoob_split_signs)
        identifier = analysis["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        csv = client.get(f"/convert-report/{identifier}", params={**owner, "format": "csv"})
        xlsx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "xlsx"})
        ofx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx"})
        assert csv.status_code == xlsx.status_code == ofx.status_code == 200
        workbook = load_workbook(BytesIO(xlsx.content), data_only=True)
        rows = workbook["Conversao"].iter_rows(min_row=2, values_only=True)
        assert [round(float(row[2] or 0)-float(row[3] or 0), 2) for row in rows] == amounts
        assert ofx.text.count("<STMTTRN>") == len(amounts)
        assert f"<BALAMT>{bounds[1]:.2f}" in ofx.text
        assert client.get("/report/an_nonexistent_final_priority", params=owner).status_code == 404


@pytest.mark.parametrize("factory,count", [(pagbank_daily, 4), (santander_limited, 2), (sicoob_split_signs, 2)])
def test_worker_keeps_the_final_priority_signs_and_real_review_flags(tmp_path, monkeypatch, factory, count):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, factory()) as (client, job):
        analysis = job["result"]["analysis"]
        assert analysis["transactions_total"] == count
        assert analysis["pdf_processing_metrics"]["partial_statement"] == int(factory is santander_limited)
        ofx = client.get(f"/convert-report/{analysis['analysis_id']}", params={"format": "ofx"})
        assert ofx.status_code == 200
        if factory is sicoob_split_signs:
            assert "<TRNAMT>-200.00" in ofx.text
            assert "<BALAMT>-50.00" in ofx.text
