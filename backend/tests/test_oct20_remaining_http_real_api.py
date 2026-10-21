from io import BytesIO

import pytest
from openpyxl import load_workbook

from synthetic_pdf_corpus.native_text_pdf import generate_native_text_pdf
from tests.ocr_conversion_http_helpers import public_client, worker_case
from tests.test_itau_monthly_native_controls import monthly_pdf
from tests.test_nubank_rdb_financial_report import rdb_pages
from tests.test_sicredi_financial_columns import matrix_text

pytestmark = pytest.mark.real_http


def case(kind):
    if kind == "itau":
        return monthly_pdf(shuffled=True), [200, -200, -50, 50], 10
    if kind == "itau_cash":
        return monthly_pdf(cash_only=True, appendix=True, shuffled=True), [200, -50], 1150
    if kind == "sicredi":
        text = matrix_text().replace("DEBITOS.............: 7.564,91\n", "").replace("CREDITOS............: 2.600,00", "")
        return generate_native_text_pdf((tuple(text.splitlines()),)), [2600, -250, -11.61, -7303.3], -982.78
    if kind == "rdb":
        return generate_native_text_pdf(tuple(tuple(page.splitlines()) for page in rdb_pages())), [-1000, 100], None
    return b"date,description,amount\n2026-09-01,PIX,100\n", [100], None


def assert_result(analysis, kind, expected, closing):
    assert [row["amount"] for row in analysis["preview_transactions"]] == expected
    assert analysis["closing_balance"] == closing
    if kind != "csv":
        assert analysis["pdf_processing_metrics"]["canonical_warning_count"] == 0
    if kind == "rdb":
        assert analysis["opening_balance"] is None
        assert analysis["financial_document"]["document_type"] == "investment_statement"
        assert analysis["financial_document"]["events"][1]["amounts"]["irrf"] == "4.00"


@pytest.mark.parametrize("kind", ["rdb", "itau", "itau_cash", "sicredi", "csv"])
def test_public_remaining_formats_preserve_financial_semantics_and_downloads(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw, expected, closing = case(kind)
    filename, mime = ("statement.csv", "text/csv") if kind == "csv" else ("statement.pdf", "application/pdf")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-remaining-" + kind}
        response = client.post("/convert", data=owner, files={"file": (filename, raw, mime)})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert_result(analysis, kind, expected, closing)
        analysis_id = analysis["analysis_id"]
        assert client.get(f"/report/{analysis_id}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{analysis_id}", params={**owner, "format": fmt})
            assert report.status_code == (400 if kind == "rdb" and fmt == "ofx" else 200), report.text
            if kind == "rdb" and fmt == "xlsx":
                workbook = load_workbook(BytesIO(report.content))
                assert {"Eventos", "Posições", "Validações"} <= set(workbook.sheetnames)
            if kind != "rdb" and fmt == "ofx":
                assert report.text.count("<STMTTRN>") == len(expected)
                for amount in expected:
                    assert f"<TRNAMT>{amount:.2f}" in report.text
        assert client.get("/report/an_remaining_missing", params=owner).status_code == 404
        assert client.post("/analyze").status_code == 405


@pytest.mark.parametrize("kind", ["rdb", "itau", "itau_cash", "sicredi"])
def test_worker_uses_the_same_remaining_statement_controls(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw, expected, closing = case(kind)
    with worker_case(tmp_path, raw, filename=kind + ".pdf") as (client, job):
        assert_result(job["result"]["analysis"], kind, expected, closing)
        report = client.get(f"/convert-report/{job['analysis_id']}", params={"format": "ofx"})
        assert report.status_code == (400 if kind == "rdb" else 200), report.text


def test_public_inconsistent_monthly_statement_does_not_export_a_partial_ledger(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        result = client.post("/convert", data={"anonymous_fingerprint": "fp-remaining-invalid"},
                             files={"file": ("itau.pdf", monthly_pdf(wrong_total=True), "application/pdf")})
        assert result.status_code == 400, result.text
