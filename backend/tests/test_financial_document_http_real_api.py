import hashlib
from io import BytesIO
from zipfile import ZipFile

import httpx
import pytest
from openpyxl import load_workbook

from app.application.conversion.document_conversion_pipeline import DocumentConversionPipeline
from app.application.conversion.uploaded_document import ingest_uploaded_document
from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.report_service import ReportService
from app.application.storage_service import TempAnalysisStorage
from app.dependencies import (
    get_access_control_service,
    get_analysis_repository,
    get_async_conversion_report_service,
    get_conversion_batch_service,
    get_conversion_processing_pipeline,
    get_report_service,
)
from app.main import app
from app.workers.conversion_lambda import ConversionLambdaProcessor
from tests.test_conversion_workspace_http_real_api import _run_pro_upload_http_server
from tests.test_convert_edits_http_real_api import _run_http_server
from tests.test_financial_document_events import report_pdf

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal", "csv"])
def test_public_upload_and_download_keep_report_contract(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    storage = TempAnalysisStorage(root_dir=tmp_path / "public")

    def configure(_access):
        app.dependency_overrides[get_analysis_repository] = lambda: storage
        app.dependency_overrides[get_conversion_processing_pipeline] = lambda: build_default_conversion_pipeline()
        app.dependency_overrides[get_report_service] = lambda: ReportService(storage)

    raw = report_pdf(kind) if kind != "csv" else b"date,description,amount\n2026-09-01,PIX,100\n"
    filename, mime = (f"{kind}.pdf", "application/pdf") if kind != "csv" else ("statement.csv", "text/csv")
    with _run_http_server(configure) as base_url, httpx.Client(base_url=base_url, timeout=20) as client:
        owner = {"anonymous_fingerprint": "fp-financial-report"}
        response = client.post("/convert", data=owner, files={"file": (filename, raw, mime)})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        analysis_id = analysis["analysis_id"]
        if kind != "csv":
            assert analysis["financial_document"]["events"]
            assert analysis["insights"][0]["type"] == "financial_document_extracted"
        assert client.get(f"/report/{analysis_id}", params=owner).status_code == 200
        assert client.get(f"/convert-report/{analysis_id}", params={**owner, "format": "csv"}).status_code == 200
        exported = client.get(f"/convert-report/{analysis_id}", params={**owner, "format": "ofx"})
        assert exported.status_code == (200 if kind == "csv" else 400)
        assert client.get("/report/an_financial_nonexistent", params=owner).status_code == 404


@pytest.mark.parametrize("kind,count", [("receivables", 2), ("investment", 2), ("fiscal", 3), ("csv", 2)])
def test_worker_and_http_exports_respect_document_semantics(kind, count, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw = report_pdf(kind) if kind != "csv" else b"date,description,amount\n2026-09-01,PIX,100\n2026-09-02,TARIFA,-5\n"
    filename = f"{kind}.pdf" if kind != "csv" else "statement.csv"
    mime = "application/pdf" if kind != "csv" else "text/csv"

    class OriginalStore:
        def load(self, reference):
            return ingest_uploaded_document(reference.filename, raw)

    with _run_pro_upload_http_server(tmp_path) as (base_url, publisher, token):
        service = app.dependency_overrides[get_conversion_batch_service]()
        access = app.dependency_overrides[get_access_control_service]()
        storage = TempAnalysisStorage(root_dir=tmp_path / "analyses", ttl_seconds=3600)
        reports = ReportService(storage)
        processor = ConversionLambdaProcessor(repository=service.repository, document_store=OriginalStore(),
            pipeline=DocumentConversionPipeline(report_service=reports, access_control_service=access,
                processing_pipeline=build_default_conversion_pipeline(), analysis_repository=storage))
        app.dependency_overrides[get_async_conversion_report_service] = lambda: reports
        app.dependency_overrides[get_report_service] = lambda: reports
        with httpx.Client(base_url=base_url, timeout=20, headers={"Authorization": f"Bearer {token}",
                "Origin": "http://localhost:3000", "X-OFX-Pro-Client": "web"}) as client:
            created = client.post("/api/conversion-batches", headers={"Idempotency-Key": kind}, json={"files": [
                {"filename": filename, "content_type": mime, "size_bytes": len(raw), "sha256_hex": hashlib.sha256(raw).hexdigest()},
            ]})
            assert created.status_code == 201, created.text
            batch_id = created.json()["batch_id"]
            assert client.post(f"/api/conversion-batches/{batch_id}/submit", json={}).status_code == 202
            processor.process(job_id=publisher.calls[0], expected_batch_id=batch_id, trace_id=kind,
                              receive_count=1, lease_owner="financial-report-worker")
            job = client.get(f"/api/conversion-batches/{batch_id}").json()["jobs"][0]
            assert job["status"] == "completed", job
            analysis = job["result"]["analysis"]
            assert analysis["transactions_total"] == count
            analysis_id = job["analysis_id"]
            assert client.get(f"/report/{analysis_id}").status_code == 200
            for fmt in ("csv", "xlsx"):
                exported = client.get(f"/convert-report/{analysis_id}", params={"format": fmt})
                assert exported.status_code == 200, exported.text
                if fmt == "xlsx" and kind != "csv":
                    assert "Eventos" in load_workbook(BytesIO(exported.content)).sheetnames
            ofx = client.get(f"/convert-report/{analysis_id}", params={"format": "ofx"})
            if kind == "csv":
                assert ofx.status_code == 200 and ofx.text.count("<STMTTRN>") == 2
            else:
                assert analysis["financial_document"]["events"]
                assert ofx.status_code == 400, ofx.text
                overrides = {"format": "ofx", "bank_code": "237", "bank_branch": "1234", "account_number": "56789"}
                ofx = client.get(f"/convert-report/{analysis_id}", params=overrides)
                if kind == "receivables":
                    assert ofx.status_code == 200 and ofx.text.count("<STMTTRN>") == 2
                    assert "<LEDGERBAL>" not in ofx.text
                else:
                    assert ofx.status_code == 400
                edited = client.post(f"/convert-edits/{analysis_id}", json={"edits": [
                    {"row_id": "row_1", "description": "ALTERADO", "credit": 100},
                ]})
                assert edited.status_code == 400, edited.text
            archive_response = client.get(f"/api/pro/conversion-batches/{batch_id}/export")
            assert archive_response.status_code == 200, archive_response.text
            with ZipFile(BytesIO(archive_response.content)) as archive:
                assert len(archive.namelist()) == 1
                assert archive.namelist()[0].endswith(".ofx" if kind == "csv" else ".xlsx")
            assert client.get("/report/an_financial_missing").status_code == 404
            assert client.post("/analyze").status_code == 405
