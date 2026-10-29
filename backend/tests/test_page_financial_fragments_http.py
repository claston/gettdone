"""Real HTTP exports and worker results preserve split rows and review status."""

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application.parsers.ofx import parse_ofx_transactions
from tests.ocr_conversion_http_helpers import public_client, worker_case
from tests.test_page_financial_fragments import sicoob_split_row

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("kind,mismatch", [("balance", False), ("transaction", False), ("balance", True)])
def test_split_rows_convert_as_sicoob_and_export_all_movements(kind, mismatch, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictitious-page-fragment"}
        response = client.post("/convert", data=owner, files={
            "file": ("ficticio.pdf", sicoob_split_row(kind=kind, mismatch=mismatch), "application/pdf"),
        })
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        identifier = analysis["analysis_id"]
        assert analysis["transactions_total"] == 3
        assert analysis["layout_inference_name"] == "sicoob_sisbr_native_split_cd_v1"
        metrics = analysis["pdf_processing_metrics"]
        assert metrics["financial_controls_validated"] == int(not mismatch)
        if mismatch:
            assert metrics["export_recommendation"] == "review_recommended"
            assert "financial_validation_failed" in metrics["canonical_warning_types"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200, report.text
            if fmt == "csv":
                rows = csv.DictReader(line for line in StringIO(report.content.decode("utf-8-sig")) if line.strip())
                assert [float(row["amount"]) for row in rows] == [50, -20, 10]
            elif fmt == "ofx":
                assert [row.amount for row in parse_ofx_transactions(report.content)] == [50, -20, 10]
            else:
                book = load_workbook(BytesIO(report.content), read_only=True, data_only=True)
                rows = list(book["Conversao"].iter_rows(min_row=2, values_only=True))
                assert [round(float(row[2] or 0) - float(row[3] or 0), 2) for row in rows] == [50, -20, 10]
                book.close()
        assert client.get("/report/an_fragment_nonexistent", params=owner).status_code == 404
        assert client.post("/analyze").status_code == 405


@pytest.mark.parametrize("mismatch", [False, True])
def test_worker_preserves_fragment_control_and_review(mismatch, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, sicoob_split_row(mismatch=mismatch)) as (client, job):
        analysis = job["result"]["analysis"]
        assert [row["amount"] for row in analysis["preview_transactions"]] == [50, -20, 10]
        assert analysis["pdf_processing_metrics"]["financial_controls_validated"] == int(not mismatch)
        if mismatch:
            assert analysis["pdf_processing_metrics"]["export_recommendation"] == "review_recommended"
        report = client.get(f"/convert-report/{analysis['analysis_id']}", params={"format": "ofx"})
        assert report.status_code == 200
        assert [row.amount for row in parse_ofx_transactions(report.content)] == [50, -20, 10]
