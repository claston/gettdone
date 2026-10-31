"""Pending semantics survive public conversion, persistence, export and worker."""
import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application.parsers.ofx import parse_ofx_transactions
from tests.fixtures.recovery_acceptance import grid_pdf
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_embedded_ocr_quality import layered_pdf
from tests.test_sicoob_channel_pending import modern_pending_cells, source_pages
from tests.test_sicoob_embedded_signed_values import SIGNED_CASH
from tests.test_sicoob_statement_context import sisbr_pending_cells

pytestmark = pytest.mark.real_http


def input_case(kind, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    if kind.startswith("modern"):
        pages = [modern_pending_cells()]
        amounts = [50, 200, -20, -20]
    else:
        pages = sisbr_pending_cells(history="DP CH.CANAL ATEN 1D")
        amounts = [50, 200, -20]
    if kind.endswith("ocr") or kind.endswith("mixed"):
        raw = configure_ocr(monkeypatch, source_pages(pages))
        if kind.endswith("mixed"):
            monkeypatch.setenv("TEXTRACT_FORCE", "false")
            raw = layered_pdf(SIGNED_CASH)
    else:
        raw = grid_pdf(pages)
    return raw, amounts


@pytest.mark.parametrize("kind", ["modern_native", "modern_ocr", "modern_mixed", "sisbr_native", "sisbr_ocr"])
def test_channel_pending_converts_with_warning_and_exports_exact_cash(kind, monkeypatch, tmp_path):
    raw, amounts = input_case(kind, monkeypatch)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-channel-"+kind}
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        data = response.json()["analysis"]
        assert [row["amount"] for row in data["preview_transactions"]] == amounts
        metrics = data["pdf_processing_metrics"]
        assert metrics["financial_controls_validated"] == 1
        assert metrics["non_cash_entries_excluded"] == 1
        assert metrics["export_recommendation"] == "review_recommended"
        assert "pending_deposits_excluded" in metrics["canonical_warning_types"]
        if kind.endswith("mixed"):
            assert metrics["textract_used"] == 1
        identifier = data["analysis_id"]
        assert client.get("/report/"+identifier, params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            r = client.get("/convert-report/"+identifier, params={**owner, "format": fmt})
            assert r.status_code == 200, r.text
            if fmt == "csv":
                actual = [float(row["amount"]) for row in csv.DictReader(StringIO(r.text))]
            elif fmt == "xlsx":
                book = load_workbook(BytesIO(r.content), data_only=True, read_only=True)
                actual = [float(row[2] or 0)-float(row[3] or 0) for row in book["Conversao"].iter_rows(min_row=2, values_only=True)]
                book.close()
            else:
                actual = [row.amount for row in parse_ofx_transactions(r.content)]
            assert actual == amounts
        assert client.get("/report/an_channel_missing", params=owner).status_code == 404


@pytest.mark.parametrize("kind", ["modern_mixed", "sisbr_native"])
def test_worker_preserves_channel_pending_warning_and_visual_reread(kind, monkeypatch, tmp_path):
    raw, amounts = input_case(kind, monkeypatch)
    with worker_case(tmp_path, raw) as (client, job):
        data = job["result"]["analysis"]
        assert [row["amount"] for row in data["preview_transactions"]] == amounts
        assert data["pdf_processing_metrics"]["financial_controls_validated"] == 1
        assert data["pdf_processing_metrics"]["non_cash_entries_excluded"] == 1
        assert "pending_deposits_excluded" in data["pdf_processing_metrics"]["canonical_warning_types"]
        ofx = client.get("/convert-report/"+data["analysis_id"], params={"format": "ofx"})
        assert ofx.status_code == 200
        assert [row.amount for row in parse_ofx_transactions(ofx.content)] == amounts
