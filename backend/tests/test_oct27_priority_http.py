"""Public HTTP and stored artifacts for the fictional priority recoveries."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.ocr_conversion_http_helpers import public_client
from tests.test_oct27_bradesco_recovery import bradesco_continuation
from tests.test_oct27_mercado_empty import mercado_empty
from tests.test_oct27_priority_recovery import detail_statement, infinitepay_fragment
from tests.test_oct27_uniprime_recovery import uniprime_with_position

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("factory,amounts,bounds", [
    (infinitepay_fragment, [50, -20, 10, -5, 7], (None, None)),
    (detail_statement, [-10, 50, -20], (100, 120)),
    (bradesco_continuation, [500, -150, -150], (-100, 100)),
    (uniprime_with_position, [-30, 50], (100, 120)),
    (mercado_empty, [], (150, 150)),
])
def test_priority_documents_keep_movements_and_balance_scope_in_all_exports(tmp_path, monkeypatch, factory, amounts, bounds):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-priority-recovery"}
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", factory(), "application/pdf")})
        assert response.status_code == 200, response.text
        payload = response.json()
        analysis = payload["analysis"]
        identifier = payload["processing_id"]
        assert analysis["transactions_total"] == len(amounts)
        assert (analysis["opening_balance"], analysis["closing_balance"]) == bounds
        assert analysis["pdf_processing_metrics"]["balance_consistency_failed"] == 0
        assert analysis["pdf_processing_metrics"]["textract_attempted"] == 0
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        xlsx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "xlsx"})
        assert xlsx.status_code == 200, xlsx.text
        workbook = load_workbook(BytesIO(xlsx.content), data_only=True)
        summary = dict(workbook["Resumo"].iter_rows(min_row=2, values_only=True))
        assert (summary["Saldo anterior"], summary["Saldo final"]) == bounds
        sheet = workbook["Conversao"]
        assert [float(row[2] or 0) - float(row[3] or 0) for row in sheet.iter_rows(min_row=2, values_only=True)] == amounts
        if factory is infinitepay_fragment:
            assert [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)] == [None, 130, None, 135, None]
            assert analysis["pdf_processing_metrics"]["partial_statement"] == 1
        ofx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx"})
        assert ofx.status_code == 200
        assert ofx.text.count("<STMTTRN>") == len(amounts)
        if bounds[1] is None:
            assert "<LEDGERBAL>" not in ofx.text
        else:
            assert f"<BALAMT>{bounds[1]:.2f}" in ofx.text
        for amount in amounts:
            assert f"<TRNAMT>{amount:.2f}" in ofx.text
        csv = client.get(f"/convert-report/{identifier}", params={**owner, "format": "csv"})
        assert csv.status_code == 200
        assert len([line for line in csv.text.splitlines() if line.strip()]) == len(amounts) + 1
        assert client.get("/report/an_priority_nonexistent", params=owner).status_code == 404


def test_existing_csv_http_path_and_legacy_analyze_status_remain_compatible(tmp_path):
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-priority-csv"}
        raw = b"date,description,amount\n2026-10-01,Fictional credit,25.00\n"
        response = client.post("/convert", data=owner, files={"file": ("fictional.csv", raw, "text/csv")})
        assert response.status_code == 200, response.text
        assert client.get(f"/report/{response.json()['processing_id']}", params=owner).status_code == 200
        assert client.post("/analyze", data=owner, files={"file": ("fictional.csv", raw, "text/csv")}).status_code == 405


def test_bradesco_real_financial_difference_remains_in_the_http_response_and_exports(tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-bradesco-review"}
        raw = bradesco_continuation(undated_opening=True, mutation="wrong_total")
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        payload = response.json()
        analysis = payload["analysis"]
        assert analysis["transactions_total"] == 3
        metrics = analysis["pdf_processing_metrics"]
        assert metrics["financial_controls_validated"] == 0
        assert "financial_validation_failed" in metrics["canonical_warning_types_list"]
        report = client.get(f"/convert-report/{payload['processing_id']}", params={**owner, "format": "ofx"})
        assert report.status_code == 200
        assert report.text.count("<STMTTRN>") == 3
        assert "<TRNAMT>500.00" in report.text and report.text.count("<TRNAMT>-150.00") == 2
