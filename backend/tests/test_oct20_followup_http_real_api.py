from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application import pdf_parser
from synthetic_pdf_corpus.native_text_pdf import generate_native_text_pdf
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_btg_fund_positions import fund_pages
from tests.test_oct20_followup_native import mister_page, portal_pages
from tests.test_oct20_followup_ocr import banrisul_ocr, bradesco, stone_old

pytestmark = pytest.mark.real_http


def native(kind):
    pages = {"banrisul": portal_pages, "santander": lambda: [mister_page()], "btg": fund_pages}[kind]()
    return generate_native_text_pdf(tuple(tuple(page.splitlines()) for page in pages))


@pytest.mark.parametrize("kind,count,bank", [("banrisul", 3, "041"), ("santander", 3, "033"), ("btg", 0, None)])
def test_followup_native_http_reports_and_exports(kind, count, bank, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-followup-"+kind}
        response = client.post("/convert", data=owner, files={"file": (kind+".pdf", native(kind), "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert analysis["transactions_total"] == count
        assert analysis["pdf_processing_metrics"]["textract_used"] == 0
        if bank:
            assert analysis["bank_code"] == bank
        aid = analysis["analysis_id"]
        assert client.get(f"/report/{aid}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            download = client.get(f"/convert-report/{aid}", params={**owner, "format":fmt})
            assert download.status_code == (400 if fmt == "ofx" and kind == "btg" else 200), download.text
            if fmt == "ofx" and bank:
                assert download.text.count("<STMTTRN>") == count
                assert f"<BANKID>{bank}" in download.text
            if fmt == "xlsx" and kind == "btg":
                book = load_workbook(BytesIO(download.content))
                assert book["Posições"].max_row == len(analysis["financial_document"]["observations"])+1
                assert analysis["total_inflows"] == analysis["total_outflows"] == 0


@pytest.mark.parametrize("build,amounts", [(stone_old, [100, -5, -30]), (bradesco, [50, -20]), (banrisul_ocr, [150, -20])])
def test_followup_ocr_http_validated_movements_and_ofx(build, amounts, tmp_path, monkeypatch):
    raw = configure_ocr(monkeypatch, build())
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint":"fp-followup-ocr"}
        response = client.post("/convert", data=owner, files={"file":("ocr.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert [row["amount"] for row in analysis["preview_transactions"]] == amounts
        assert analysis["pdf_processing_metrics"]["textract_used"] == 1
        assert analysis["pdf_processing_metrics"]["balance_consistency_failed"] == 0
        aid = analysis["analysis_id"]
        download = client.get(f"/convert-report/{aid}", params={**owner,"format":"ofx"})
        assert download.status_code == 200 and download.text.count("<STMTTRN>") == len(amounts)


def test_unmapped_native_layer_runs_visual_reader_instead_of_stopping_at_glyph_count(tmp_path, monkeypatch):
    raw = configure_ocr(monkeypatch, banrisul_ocr())
    monkeypatch.setenv("TEXTRACT_FORCE", "false")
    monkeypatch.setattr(pdf_parser, "_read_native_pdf_page_texts", lambda raw: ["□"*100+"abc"*50])
    with public_client(tmp_path) as client:
        response = client.post("/convert", data={"anonymous_fingerprint":"fp-followup-unmapped"},
                               files={"file":("unmapped.pdf",raw,"application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert analysis["transactions_total"] == 2
        assert analysis["pdf_processing_metrics"]["textract_used"] == 1
        assert analysis["opening_balance"] == -100 and analysis["closing_balance"] == 30


@pytest.mark.parametrize("kind", ["banrisul", "btg"])
def test_worker_uses_same_followup_native_contract(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, native(kind), filename=kind+".pdf") as (client, job):
        analysis = job["result"]["analysis"]
        assert analysis["transactions_total"] == (3 if kind == "banrisul" else 0)
        assert client.get(f"/convert-report/{analysis['analysis_id']}", params={"format":"xlsx"}).status_code == 200
