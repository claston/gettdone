"""Synthetic reports: no production documents or customer identifiers."""
import csv
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.document_type_classifier import classify_document_type
from app.application.errors import InvalidFileContentError, PdfFinancialValidationError
from app.application.pdf_parser import parse_pdf_transactions
from app.application.storage_service import TempAnalysisStorage
from app.application.structured_conversion import build_structured_conversion_result_from_analysis_data
from synthetic_pdf_corpus.native_text_pdf import generate_native_text_pdf


def report_pages(kind):
    if kind == "receivables":
        return ["""Títulos Pagos por Conta Crédito
Dados do Título
Tipo Seu Número Pagador Dt de Venc. Dt de Pag. Valor do Título (R$) Valor Cobrado (R$) Car Valor da Osci. (R$)
2 0000123/1 CLIENTE
EXEMPLO LTDA
01/09/2026 02/09/2026 100,00 102,00 Não 2,00
2 0000124/1 OUTRO CLIENTE 03/09/2026 04/09/2026 50,00 50,00 Não 0,00
Total 150,00 152,00 +2,00"""]
    if kind == "investment":
        # The native text stream lists the section headings before the table.
        return ["""Extrato de Investimentos
Produto:CDB - Certificado de Depósito Bancário
Período de:01/08/2026 até 31/08/2026
Saldo anterior em 31/07/2026
Aplicações
Resgates / Vencimentos
Saldo final em 31/08/2026
Dt. Aplicação Dt. Vencto Dt. Resgate / Carência Taxa (%) Vlr Princ. (R$) Vlr. Bruto (R$)
Renda Total(R$) Vlr. IOF (R$) Vlr. IRRF (R$) Vlr Líquido(R$) Renda Bruta Per
30/07/2026 30/07/2028 - 2,0000 100,00 100,00 0,00 0,00 0,00 100,00 0,00
Total 0,0000 100,00 100,00 0,00 0,00 0,00 100,00 0,00
14/08/2026 14/08/2028 - 2,0000 50,00 0,00 0,00 0,00
Total 0,0000 50,00 0,00 0,00 0,00 0,00 0,00 0,00
30/07/2026 30/07/2028 20/08/2026 2,0000 100,00 101,00 1,00 0,20 0,00 100,80 1,00
Total 0,0000 100,00 101,00 1,00 0,20 0,00 100,80 1,00
14/08/2026 14/08/2028 - 2,0000 50,00 50,03 0,03 0,00 0,00 50,03 0,03
Total 0,0000 50,00 50,03 0,03 0,00 0,00 50,03 0,03
Banco Bradesco S/A"""]
    key = "1" * 44
    return [f"""EMPRESA EXEMPLO LTDA
Relatório emitido em: 19/10/2026 às 15:01 por GDOOR Sistemas
NF-e, NFC-e e SAT (Resumo)
Período de: 01/09/2026 Até: 30/09/2026
Nota Emissão Chave Protocolo ValorSit.Ent./Sai.Mod.
000001 03/09/2026 {key} 000000000000001 R$ 100,00 S55
000002 09/09/2026 {'2' * 44} 000000000000002 R$ 50,00 S65
000003 21/09/2026 R$ 0,00IS65
Total Geral..: R$ 150,00
3Registros: Total Saídas..: R$ 150,00
R$ 0,00Total Entradas..:
Página 1 de 1"""]


def report_pdf(kind, pages=None):
    return generate_native_text_pdf(tuple(tuple(page.splitlines()) for page in (pages or report_pages(kind))))


def convert_report(kind, pages=None):
    return build_default_conversion_pipeline().run(
        filename=f"{kind}.pdf", raw_bytes=report_pdf(kind, pages), analysis_id=f"an_{kind}",
    ).analysis_data


@pytest.fixture(autouse=True)
def disable_cloud(monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")


@pytest.mark.parametrize("kind,expected", [
    ("receivables", "receivables_report"), ("investment", "investment_statement"), ("fiscal", "fiscal_summary"),
])
def test_classification_and_extraction_remain_distinct_from_export(kind, expected):
    raw = report_pdf(kind)
    classified = classify_document_type(f"{kind}.pdf", raw)
    assert classified.document_type == expected
    assert classified.confidence >= 0.9 and classified.evidence
    parsed = parse_pdf_transactions(raw)
    assert parsed.financial_document["document_type"] == expected
    result = convert_report(kind)
    assert result.document_type == expected
    assert result.opening_balance is result.closing_balance is None
    assert result.matched_groups == result.reversed_entries == result.potential_duplicates == 0
    assert result.financial_document["schema_version"] == "financial_events_v1"
    assert result.financial_document["representation"] == "native_text"
    assert all(e["source_refs"] for e in result.financial_document["events"])
    assert all(v["status"] == "passed" for v in result.financial_document["validations"])
    assert all(row.source_page and row.source_line and row.source_regions for row in result.preview_transactions)


def test_receivables_preserve_two_dates_three_values_and_title_identity():
    result = convert_report("receivables")
    first = result.financial_document["events"][0]
    assert first["event_type"] == "receipt" and first["status"] == "paid"
    assert first["reference"] == "0000123/1"
    assert first["counterparty"] == "CLIENTE EXEMPLO LTDA"
    assert first["dates"] == {"due": "2026-09-01", "payment": "2026-09-02"}
    assert first["amounts"] == {"original": "100.00", "collected": "102.00", "adjustment": "2.00"}
    assert first["cash_effect"] == "102.00"
    assert result.transactions_total == 2 and result.total_inflows == 152
    assert result.preview_transactions[0].date == "2026-09-02"
    assert result.financial_document["export"]["ofx_profile"] == "receipts"


def test_investments_keep_positions_outside_cash_movements():
    result = convert_report("investment")
    doc = result.financial_document
    assert result.transactions_total == 2
    assert [e["event_type"] for e in doc["events"]] == ["investment_application", "investment_redemption"]
    assert [e["cash_effect"] for e in doc["events"]] == ["-50.00", "100.80"]
    assert len(doc["observations"]) == 2
    assert doc["observations"][0]["amounts"]["principal"] == "100.00"
    assert doc["observations"][1]["amounts"]["net"] == "50.03"
    redemption = doc["events"][1]
    assert redemption["dates"]["application"] == "2026-07-30"
    assert redemption["dates"]["redemption"] == "2026-08-20"
    assert redemption["amounts"]["iof"] == "0.20"
    assert redemption["amounts"]["gross"] == "101.00"
    assert result.net_total == 50.8
    assert doc["export"]["available_formats"] == ["csv", "xlsx"]


def test_fiscal_records_preserve_voided_zero_record_without_claiming_receipts():
    result = convert_report("fiscal")
    doc = result.financial_document
    assert result.transactions_total == 3
    assert result.total_inflows == result.total_outflows == result.net_total == 0
    assert doc["origin"] == "GDOOR"
    assert all(e["cash_effect"] is None for e in doc["events"])
    assert doc["events"][0]["attributes"]["access_key"] == "1" * 44
    assert doc["events"][0]["attributes"]["protocol"] == "000000000000001"
    assert doc["events"][2]["status"] == "voided"
    assert doc["events"][2]["attributes"]["status_code"] == "I"
    assert doc["export"]["available_formats"] == ["csv", "xlsx"]


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_details_survive_persistence_structured_result_and_both_spreadsheets(kind, tmp_path):
    result = convert_report(kind)
    storage = TempAnalysisStorage(root_dir=tmp_path)
    storage.save_analysis(result)
    saved = json.loads((tmp_path / result.analysis_id / "analysis.json").read_text(encoding="utf8"))
    assert saved["financial_document"] == result.financial_document
    structured = build_structured_conversion_result_from_analysis_data(result)
    assert structured.financial_document == result.financial_document
    assert structured.export.available_formats == result.financial_document["export"]["available_formats"]
    assert saved["structured_result"]["financial_document"] == result.financial_document
    csv_path = storage.get_convert_report_path(result.analysis_id, "csv")
    rows = list(csv.DictReader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
    assert len(rows) == len(result.financial_document["events"]) + len(result.financial_document["observations"])
    assert sum(row["record_type"] == "event" for row in rows) == result.transactions_total
    assert "event_type" in rows[0] and "cash_effect" in rows[0] and "source_refs" in rows[0]
    for path in [storage.get_report_path(result.analysis_id), storage.get_convert_report_path(result.analysis_id, "xlsx")]:
        workbook = load_workbook(BytesIO(path.read_bytes()))
        assert {"Eventos", "Detalhes", "Validações", "Documento"}.issubset(workbook.sheetnames)
        assert workbook["Eventos"].max_row == result.transactions_total + 1
        headers = [cell.value for cell in workbook["Eventos"][1]]
        assert workbook["Eventos"].cell(2, headers.index("Valor") + 1).data_type == "n"
        if kind != "investment":
            assert workbook["Eventos"].cell(2, headers.index("Referência") + 1).value.startswith("000")
    snapshot = storage._build_analysis_snapshot_from_content(saved, preview_rows=result.preview_transactions,
                                                           report_rows=result.report_transactions)
    assert snapshot.financial_document == result.financial_document
    storage.set_convert_owner(result.analysis_id, "anonymous", "report-owner")
    history = storage.list_convert_history("anonymous", "report-owner")
    assert history[0]["status"] == "Sucesso" and history[0]["conversion_type"] == "pdf-xlsx"


def test_receipts_ofx_requires_destination_and_never_claims_bank_balance(tmp_path):
    result = convert_report("receivables")
    storage = TempAnalysisStorage(root_dir=tmp_path)
    storage.save_analysis(result)
    with pytest.raises(InvalidFileContentError, match="conta"):
        storage.get_convert_report_path(result.analysis_id, "ofx")
    path = storage.get_convert_report_path(result.analysis_id, "ofx", bank_code="237", bank_branch="1234", account_number="56789")
    ofx = path.read_text(encoding="utf8")
    assert ofx.count("<STMTTRN>") == 2 and "<TRNAMT>102.00" in ofx
    assert "<DTPOSTED>20260902" in ofx and "<ACCTID>56789" in ofx
    assert "<LEDGERBAL>" not in ofx


@pytest.mark.parametrize("kind", ["investment", "fiscal"])
def test_disallowed_ofx_cannot_be_created_by_account_or_balance_overrides(kind, tmp_path):
    result = convert_report(kind)
    storage = TempAnalysisStorage(root_dir=tmp_path)
    storage.save_analysis(result)
    for overrides in [{}, {"bank_code": "237", "bank_branch": "1234", "account_number": "56789", "closing_balance": 1}]:
        with pytest.raises(InvalidFileContentError, match="OFX"):
            storage.get_convert_report_path(result.analysis_id, "ofx", **overrides)
    assert not (tmp_path / result.analysis_id / "converted.ofx").exists()


def test_total_difference_keeps_extraction_but_disables_receipts_ofx(tmp_path):
    pages = [report_pages("receivables")[0].replace("Total 150,00 152,00", "Total 150,00 999,00")]
    result = convert_report("receivables", pages)
    assert result.transactions_total == 2
    assert any(v["status"] == "failed" for v in result.financial_document["validations"])
    assert result.financial_document["export"]["available_formats"] == ["csv", "xlsx"]
    storage = TempAnalysisStorage(root_dir=tmp_path)
    storage.save_analysis(result)
    with pytest.raises(InvalidFileContentError, match="OFX"):
        storage.get_convert_report_path(result.analysis_id, "ofx", bank_code="237", bank_branch="1234", account_number="56789")


@pytest.mark.parametrize("kind,before,after", [
    ("receivables", "02/09/2026", "31/02/2026"),
    ("investment", "20/08/2026", "31/02/2026"),
    ("fiscal", "03/09/2026", "31/02/2026"),
])
def test_invalid_source_date_never_falls_back_to_generic_bank_parser(kind, before, after):
    with pytest.raises(PdfFinancialValidationError):
        convert_report(kind, [report_pages(kind)[0].replace(before, after)])


def test_bank_counterparty_names_cannot_reclassify_statement():
    text = """Extrato de conta corrente
Agência: 1234 Conta: 56789-0
Data Histórico Valor Saldo
01/09/2026 PIX GDOOR Sistemas 100,00 200,00
02/09/2026 Títulos Pagos por Conta Crédito 50,00 250,00"""
    assert classify_document_type("statement.pdf", b"", extracted_text=text).document_type == "bank_statement"


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_bank_statement_and_report_annex_keep_separate_identity_and_rows(kind):
    from app.application.pdf_document_segmentation import segment_pdf_pages

    bank = """Extrato de conta corrente
Agência: 1234 Conta: 56789-0
Data Histórico Valor Saldo
01/09/2026 PIX RECEBIDO 100,00 200,00"""
    for pages in [[bank, *report_pages(kind)], [*report_pages(kind), bank]]:
        segmented = segment_pdf_pages(pages)
        assert len(segmented.parts) == 2
        assert [page.strip() for page in segmented.page_texts if page.strip()] == [bank]
        report_part = next(part for part in segmented.parts if not part["included_in_statement"])
        assert report_part["document_type"] != "unknown"


def test_negative_collected_amount_is_preserved_but_not_exported_as_credit():
    pages = [report_pages("receivables")[0].replace("100,00 102,00 Não 2,00", "100,00 -102,00 Não -202,00")
             .replace("Total 150,00 152,00 +2,00", "Total 150,00 -52,00 -202,00")]
    result = convert_report("receivables", pages)
    assert result.financial_document["events"][0]["amounts"]["collected"] == "-102.00"
    assert "ofx" not in result.financial_document["export"]["available_formats"]


def test_fiscal_direction_totals_are_checked_without_becoming_cash():
    pages = [report_pages("fiscal")[0].replace("R$ 50,00 S65", "R$ 50,00 E65")]
    result = convert_report("fiscal", pages)
    assert result.transactions_total == 3 and result.net_total == 0
    assert any(v["status"] == "failed" for v in result.financial_document["validations"])


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_report_must_not_hide_malformed_date_leading_rows(kind):
    pages = [report_pages(kind)[0] + "\n01/09/2026 CAMPO SEM VALOR"]
    with pytest.raises(PdfFinancialValidationError):
        convert_report(kind, pages)


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_classification_is_kept_in_operational_history(kind):
    from app.application.document_classification_history import prepare_document_classification_record

    result = convert_report(kind)
    record = prepare_document_classification_record(document_type=result.document_type,
        document_type_confidence=result.document_type_confidence,
        document_classification_version=result.document_classification_version,
        document_processing_decision="accepted", document_classification_evidence=result.document_type_evidence)
    assert record[0] == result.document_type and record[3] == "accepted"


def test_report_context_keeps_owner_and_period_separate_from_event_dates():
    pages = [report_pages("investment")[0].replace("Produto:CDB", "Empresa/CNPJ: EMPRESA EXEMPLO | 00.000.000/0001-00\n"
             "Agência/conta: 1234 | 000055-8\nProduto:CDB")]
    result = convert_report("investment", pages)
    metadata = result.financial_document["metadata"]
    assert metadata["owner_name"] == "EMPRESA EXEMPLO"
    assert metadata["owner_tax_id"] == "00.000.000/0001-00"
    assert metadata["source_account_number"] == "000055-8"
    assert metadata["period_start"] == "2026-08-01" and metadata["period_end"] == "2026-08-31"
    assert result.financial_document["events"][1]["dates"]["application"] == "2026-07-30"


def test_unrecognized_zero_title_cannot_disappear_behind_matching_totals():
    extra_row = "2 OUTRO CLIENTE 01/09/2026 02/09/2026 0,00 0,00 Não 0,00\nTotal 150,00"
    pages = [report_pages("receivables")[0].replace("Total 150,00", extra_row)]
    with pytest.raises(PdfFinancialValidationError):
        convert_report("receivables", pages)


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_processing_metrics_count_extracted_events(kind):
    result = convert_report(kind)
    assert result.pdf_processing_metrics["canonical_transactions_count"] == result.transactions_total


@pytest.mark.parametrize("kind", ["receivables", "investment", "fiscal"])
def test_events_survive_s3_reload_and_export_policy_is_enforced(kind, tmp_path):
    from app.application.s3_analysis_storage import S3AnalysisStorage
    from tests.test_s3_analysis_storage import FakeS3Client

    result = convert_report(kind)
    client = FakeS3Client()
    worker = S3AnalysisStorage(root_dir=tmp_path / "worker", bucket="private", s3_client=client)
    worker.save_analysis(result)
    api = S3AnalysisStorage(root_dir=tmp_path / "api", bucket="private", s3_client=client)
    assert api.get_convert_report_path(result.analysis_id, "xlsx").exists()
    stored = json.loads((tmp_path / "api" / result.analysis_id / "analysis.json").read_text(encoding="utf8"))
    assert stored["financial_document"] == result.financial_document
    if kind != "receivables":
        with pytest.raises(InvalidFileContentError, match="OFX"):
            api.get_convert_report_path(result.analysis_id, "ofx", bank_code="237", bank_branch="1234", account_number="56789")
    else:
        path = api.get_convert_report_path(result.analysis_id, "ofx", bank_code="237", bank_branch="1234", account_number="56789")
        assert "<TRNAMT>102.00" in path.read_text(encoding="utf8")
        restarted = S3AnalysisStorage(root_dir=tmp_path / "restarted", bucket="private", s3_client=client)
        assert "<ACCTID>56789" in restarted.get_convert_report_path(result.analysis_id, "ofx").read_text(encoding="utf8")
