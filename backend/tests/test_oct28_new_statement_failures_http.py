"""Real HTTP coverage for native/OCR readers, worker, reports and exports."""

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_oct28_new_statement_failures import (
    append_image_page,
    banrisul_month_source,
    bb_duplicate_opening,
    brb_lot_statement,
    caixa_tilted_rows,
    mercado_without_continuation_header,
    sicoob_descending,
)

pytestmark = pytest.mark.real_http


def source_case(kind, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    if kind == "lot_rotated":
        return brb_lot_statement(rotated=True), [-20, 50, 0.01]
    if kind == "mercado":
        return mercado_without_continuation_header(), [10, -5]
    if kind == "blank":
        return append_image_page(sicoob_descending()), [50, -200]
    source, amounts = {
        "caixa": (caixa_tilted_rows(), [-20, 10]),
        "bb": (bb_duplicate_opening(), [50]),
        "banrisul": (banrisul_month_source(), [-20, 50, 0.01]),
    }[kind]
    return configure_ocr(monkeypatch, source), amounts


@pytest.mark.parametrize("kind", ["lot_rotated", "mercado", "blank", "caixa", "bb", "banrisul"])
def test_new_cases_convert_over_http_and_export_every_transaction(kind, tmp_path, monkeypatch):
    raw, amounts = source_case(kind, monkeypatch)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-oct28-new-" + kind}
        response = client.post("/convert", data=owner, files={"file": (kind + ".pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        data = response.json()["analysis"]
        assert [t["amount"] for t in data["preview_transactions"]] == amounts
        identifier = data["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200, report.text
            if fmt == "ofx":
                assert report.text.count("<STMTTRN>") == len(amounts)
            elif fmt == "csv":
                assert len([r for r in csv.reader(StringIO(report.content.decode("utf-8-sig"))) if r]) - 1 == len(amounts)
            else:
                book = load_workbook(BytesIO(report.content), read_only=True)
                assert book.active.max_row - 1 == len(amounts)
                book.close()
        assert client.get("/report/an_oct28_new_nonexistent", params=owner).status_code == 404


@pytest.mark.parametrize("kind", ["lot_rotated", "banrisul"])
def test_worker_uses_the_same_readers_and_balances(kind, tmp_path, monkeypatch):
    raw, amounts = source_case(kind, monkeypatch)
    with worker_case(tmp_path, raw) as (client, job):
        data = job["result"]["analysis"]
        assert [t["amount"] for t in data["preview_transactions"]] == amounts
        assert data["pdf_processing_metrics"]["balance_consistency_failed"] == 0
        assert client.get(f"/convert-report/{data['analysis_id']}", params={"format": "ofx"}).status_code == 200


def test_csv_happy_path_and_analyze_remains_disabled(tmp_path):
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-oct28-new-csv"}
        response = client.post(
            "/convert",
            data=owner,
            files={"file": ("statement.csv", b"date,description,amount\n2026-10-28,Recebimento teste,100\n", "text/csv")},
        )
        assert response.status_code == 200, response.text
        assert client.get(f"/report/{response.json()['analysis']['analysis_id']}", params=owner).status_code == 200
        assert client.post("/analyze").status_code == 405
