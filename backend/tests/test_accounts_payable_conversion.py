from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from app.application import pdf_parser
from app.application.default_conversion_pipeline import build_default_conversion_pipeline
from app.application.document_extraction_models import ExtractedLine, ExtractedPage, RawDocumentExtraction
from app.application.document_routing_policy import DocumentRoutingAction, enforce_document_routing
from app.application.document_type_classifier import DocumentTypeClassification
from app.application.errors import InvalidFileContentError
from app.application.storage_service import TempAnalysisStorage
from tests.fixtures.accounts_payable import ACCOUNTS_PAYABLE_REPORT
from tests.fixtures.accounts_payable_pdf import ROWS, accounts_payable_pdf


@pytest.mark.parametrize("rotation,split_pages", [(0, False), (90, False), (180, False), (270, False), (0, True), (90, True)])
def test_native_report_converts_paid_and_pending_with_correct_dates_without_ocr(monkeypatch, rotation, split_pages):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: True)
    monkeypatch.setattr(pdf_parser, "is_textract_forced", lambda: False)
    monkeypatch.setattr(pdf_parser, "is_pdf_ocr_enabled", lambda: True)

    def no_ocr(*args, **kwargs):
        pytest.fail("Native report does not require OCR")

    monkeypatch.setattr(pdf_parser, "_parse_scanned_pdf_with_textract_gateway", no_ocr)
    monkeypatch.setattr(pdf_parser, "_extract_pdf_page_texts_with_ocr", no_ocr)
    result = pdf_parser.parse_pdf_transactions(accounts_payable_pdf(rotation=rotation, split_pages=split_pages))
    assert [row.amount for row in result.transactions] == [90, 30, 80]
    assert [row.date for row in result.transactions] == ["2026-09-01", "2026-09-03", "2026-09-04"]
    assert [event["status"] for event in result.financial_document["events"]] == ["settled", "pending", "partial"]
    assert all(row.type == "non_cash" for row in result.canonical_transactions)
    assert all(row.running_balance is None for row in result.canonical_transactions)
    assert result.parse_metrics["balance_consistency_failed"] == 0
    assert result.parse_metrics["textract_attempted"] == 0
    assert result.financial_document["export"]["available_formats"] == ["csv", "xlsx"]
    if split_pages:
        assert [row.source_page for row in result.canonical_transactions] == [1, 2, 2]


def test_pipeline_keeps_report_identity_and_descriptions_without_guessing_bank(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    data = build_default_conversion_pipeline().run(
        filename="report.pdf", raw_bytes=accounts_payable_pdf(), analysis_id="synthetic-payables",
    ).analysis_data
    assert data.document_type == "accounts_payable_report"
    assert data.transactions_total == 3
    assert data.total_outflows == data.total_inflows == 0
    assert all(value is None for value in (data.bank_name, data.bank_code, data.bank_branch, data.account_number,
                                          data.opening_balance, data.closing_balance))
    assert "UBER SERVIÇO" in data.preview_transactions[0].description
    assert data.financial_document["events"][0]["status"] == "settled"
    assert data.financial_document["events"][0]["counterparty"] == "BANCO INTER FICTÍCIO"


def test_recognized_report_is_allowed_by_routing():
    assert enforce_document_routing(DocumentTypeClassification("accounts_payable_report", 0.99)).action == DocumentRoutingAction.CONTINUE


def test_excel_csv_include_pending_and_ofx_is_not_a_title_export(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    data = build_default_conversion_pipeline().run(
        filename="report.pdf", raw_bytes=accounts_payable_pdf(), analysis_id="synthetic-exports",
    ).analysis_data
    storage = TempAnalysisStorage(tmp_path)
    storage.save_analysis(data)
    with pytest.raises(InvalidFileContentError, match="OFX"):
        storage.get_convert_report_path(data.analysis_id, "ofx")
    csv = storage.get_convert_report_path(data.analysis_id, "csv").read_text(encoding="utf-8")
    assert "pending" in csv and "partial" in csv
    workbook = load_workbook(BytesIO(storage.get_convert_report_path(data.analysis_id, "xlsx").read_bytes()))
    assert workbook["Eventos"].max_row == 4


@pytest.mark.parametrize("mutation", ["total", "missing_amount", "missing_date", "bad_date", "conflicting_status"])
def test_ambiguous_report_does_not_fall_back_to_partial_generic_reading(monkeypatch, mutation):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    monkeypatch.setattr(pdf_parser, "is_pdf_ocr_enabled", lambda: True)

    def no_ocr(*args, **kwargs):
        pytest.fail("Financial ambiguity must not trigger OCR")

    monkeypatch.setattr(pdf_parser, "_extract_pdf_page_texts_with_ocr", no_ocr)
    rows = [list(row) for row in ROWS]
    total = "200,00"
    if mutation == "total":
        total = "201,00"
    elif mutation == "missing_amount":
        rows[1][5] = ""
    elif mutation == "missing_date":
        rows[0][1] = ""
    elif mutation == "bad_date":
        rows[0][1] = "31/02/2026"
    else:
        rows[0][6] = "10,00"
    if mutation in {"total", "missing_date", "conflicting_status"}:
        parsed = pdf_parser.parse_pdf_transactions(accounts_payable_pdf(rows, total=total))
        assert len(parsed.financial_document["events"]) == 3
        if mutation == "missing_date":
            assert "payment" not in parsed.financial_document["events"][0]["dates"]
        else:
            assert parsed.parse_metrics["export_recommendation"] == "review_recommended"
    else:
        with pytest.raises(InvalidFileContentError, match="contas a pagar"):
            pdf_parser.parse_pdf_transactions(accounts_payable_pdf(rows, total=total))


def test_wrapped_supplier_description_and_duplicate_items_are_preserved(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    row = list(ROWS[0])
    row[2] = "SERVIÇO\nCOMPLEMENTO"
    row[3] = "FORNECEDOR\nFICTÍCIO"
    parsed = pdf_parser.parse_pdf_transactions(accounts_payable_pdf([row, row], total="180,00"))
    assert len(parsed.transactions) == 2
    assert all("COMPLEMENTO" in item.description for item in parsed.transactions)
    assert all("FICTÍCIO" in event["counterparty"] for event in parsed.financial_document["events"])
    assert parsed.canonical_transactions[0].source_line != parsed.canonical_transactions[1].source_line


def test_empty_report_with_zero_printed_total_is_identified_without_inventing_transactions(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    data = build_default_conversion_pipeline().run(
        filename="empty.pdf", raw_bytes=accounts_payable_pdf([], total="0,00"), analysis_id="synthetic-empty-report",
    ).analysis_data
    assert data.document_type == "accounts_payable_report" and data.transactions_total == 0


def test_period_filter_dates_and_user_footer_do_not_become_payables(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    parsed = pdf_parser.parse_pdf_transactions(accounts_payable_pdf(footer=True, rotation=90))
    assert len(parsed.transactions) == 3
    assert all("Usuário" not in row.description for row in parsed.transactions)


def test_filtered_paid_rows_reconcile_paid_subtotal_and_explain_different_global_total(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    parsed = pdf_parser.parse_pdf_transactions(accounts_payable_pdf([ROWS[0], ROWS[0]], total="200,00", paid="180,00", footer=True))
    assert [item.amount for item in parsed.transactions] == [90, 90]
    assert parsed.parse_metrics["export_recommendation"] == "review_recommended"


def test_incomplete_header_does_not_fall_back_to_generic_bank_movements(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    with pytest.raises(InvalidFileContentError, match="contas a pagar"):
        pdf_parser.parse_pdf_transactions(accounts_payable_pdf(missing_header=True))


@pytest.mark.parametrize("scale", [0.75, 3])
def test_column_and_footer_geometry_follow_page_scale(monkeypatch, scale):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    reader = PdfReader(BytesIO(accounts_payable_pdf(footer=True)))
    writer = PdfWriter()
    for page in reader.pages:
        page.scale_by(scale)
        writer.add_page(page)
    raw = BytesIO()
    writer.write(raw)
    parsed = pdf_parser.parse_pdf_transactions(raw.getvalue())
    assert [item.amount for item in parsed.transactions] == [90, 30, 80]


def test_report_without_usable_ocr_geometry_does_not_become_generic_bank_transactions(monkeypatch):
    extraction = RawDocumentExtraction(provider="textract", document_hash="synthetic", metrics={}, pages=[
        ExtractedPage(1, lines=[ExtractedLine(str(index), 1, index, text)
                               for index, text in enumerate(ACCOUNTS_PAYABLE_REPORT.splitlines())]),
    ])

    def unexpected_adapter(*args, **kwargs):
        pytest.fail("A recognized report must not fall back to bank transaction heuristics")

    monkeypatch.setattr(pdf_parser, "adapt_textract_extraction_to_transactions", unexpected_adapter)
    with pytest.raises(InvalidFileContentError, match="contas a pagar"):
        pdf_parser._parse_scanned_pdf_with_textract_text_pages(extraction)


def test_paying_a_card_invoice_does_not_turn_payables_into_a_credit_card_statement(monkeypatch):
    monkeypatch.setattr(pdf_parser, "is_textract_enabled", lambda: False)
    row = list(ROWS[0])
    row[2] = "FATURA CARTÃO DE CRÉDITO TOTAL A PAGAR"
    data = build_default_conversion_pipeline().run(
        filename="report.pdf", raw_bytes=accounts_payable_pdf([row], total="90,00"), analysis_id="synthetic-card-payable",
    ).analysis_data
    assert data.document_type == "accounts_payable_report" and data.ofx_account_type is None
