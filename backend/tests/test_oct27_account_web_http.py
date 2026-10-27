"""The batch's public and worker exports retain each document's meaning."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.ocr_conversion_http_helpers import public_client, worker_case
from tests.test_oct27_account_web_batch import descending_cards, itau_complete, xp_web

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("factory,amounts,bounds", [
    (xp_web, [100, -30], (None, None)),
    (itau_complete, [-5, 15, -15, -5, 20], (10, 20)),
    (descending_cards, [-5, 20, -10, 0], (100, 105)),
])
def test_web_account_documents_export_over_real_http(tmp_path, monkeypatch, factory, amounts, bounds):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fictional-web-account-batch"}
        response = client.post("/convert", data=owner, files={"file": ("fictional.pdf", factory(), "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert analysis["transactions_total"] == len(amounts)
        assert (analysis["opening_balance"], analysis["closing_balance"]) == bounds
        identifier = analysis["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        csv = client.get(f"/convert-report/{identifier}", params={**owner, "format": "csv"})
        xlsx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "xlsx"})
        assert csv.status_code == xlsx.status_code == 200
        workbook = load_workbook(BytesIO(xlsx.content), data_only=True)
        ofx = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx", "bank_code": "237",
            "bank_branch": "1234", "account_number": "12345-6"})
        if factory is xp_web:
            assert ofx.status_code == 400 and "OFX" in ofx.text
            assert "dates.settlement" in csv.text and "dates.transaction" in csv.text
            assert "Data de liquidação" in [cell.value for cell in workbook["Eventos"][1]]
            assert analysis["document_type"] == "investment_statement"
            assert analysis["bank_code"] is None
        else:
            assert ofx.status_code == 200
            assert ofx.text.count("<STMTTRN>") == len(amounts)
            rows = workbook["Conversao"].iter_rows(min_row=2, values_only=True)
            assert [float(row[2] or 0)-float(row[3] or 0) for row in rows] == amounts
            assert f"<BALAMT>{bounds[1]:.2f}" in ofx.text
            if factory is descending_cards:
                assert analysis["bank_code"] is None and analysis["bank_name"] is None
                assert "COBRANCA DE JUROS" in csv.text
            manual = client.get(f"/convert-report/{identifier}", params={**owner, "format": "ofx", "bank_code": "237",
                "bank_branch": "1234", "account_number": "12345-6", "closing_balance": 0})
            assert manual.status_code == 200
            assert "<BALAMT>0.00" in manual.text
        assert client.get("/report/an_nonexistent_web_account", params=owner).status_code == 404


@pytest.mark.parametrize("factory,count", [(xp_web, 2), (itau_complete, 5), (descending_cards, 4)])
def test_worker_preserves_the_same_document_policy(tmp_path, monkeypatch, factory, count):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, factory()) as (client, job):
        analysis = job["result"]["analysis"]
        assert analysis["transactions_total"] == count
        for fmt in ("csv", "xlsx"):
            assert client.get(f"/convert-report/{analysis['analysis_id']}", params={"format": fmt}).status_code == 200
        if factory is xp_web:
            assert analysis["financial_document"]["export"]["available_formats"] == ["csv", "xlsx"]
            assert analysis["document_type"] == "investment_statement"
