"""Real HTTP acceptance for page fragments, OCR evidence and financial review."""

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook

from app.application import pdf_parser
from app.application.parsers.ofx import parse_ofx_transactions
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case
from tests.test_oct28_late_statement_recovery import (
    bb_overlapping_sign,
    bb_page_break,
    bb_pending_deposit,
    bb_trailing_mark,
    bradesco_reordered_period,
    santander_available,
    sicoob_page_break,
)

pytestmark = pytest.mark.real_http


def source_case(kind, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    if kind == "sicoob":
        return sicoob_page_break(blocked=True), [50, -20], "pending_deposits_excluded"
    if kind == "bb_native":
        return bb_page_break(), [10, 20], None
    if kind == "bradesco":
        return bradesco_reordered_period(mismatch=True), [50, -20], "printed_totals_differ_from_extracted_movements"
    source, amounts, warning = {
        "santander": (santander_available(), [5000, 200, 200, 200], None),
        "bb_stroke": (bb_overlapping_sign(), [-4.45], "ocr_sign_stroke_overlap_confirmed"),
        "bb_mark": (bb_trailing_mark(skew=True), [50], "ocr_trailing_mark_confirmed"),
        "bb_pending": (bb_pending_deposit(), [50], "pending_deposits_excluded"),
    }[kind]
    raw = configure_ocr(monkeypatch, source)
    response = pdf_parser.TextractGateway.analyze_pdf(raw)
    by_id = {line.id: line for page in source.pages for line in page.lines}
    for block in response["blocks"]:
        ids = by_id[block["Id"]].word_ids
        if ids:
            block["Relationships"] = [{"Type": "CHILD", "Ids": list(ids)}]
    response["blocks"].extend({
        "Id": word.id, "BlockType": "WORD", "Page": word.page_number, "Text": word.text,
        "Confidence": word.confidence, "TextType": word.text_type,
        "Geometry": {"BoundingBox": {key.title(): value for key, value in word.bbox.items()}},
    } for page in source.pages for word in page.words)
    monkeypatch.setattr(pdf_parser.TextractGateway, "analyze_pdf", lambda *a, **kw: response)
    return raw, amounts, warning


@pytest.mark.parametrize("kind", ["santander", "sicoob", "bb_native", "bb_stroke", "bb_mark", "bb_pending", "bradesco"])
def test_late_cases_convert_and_preserve_all_export_values_over_http(kind, tmp_path, monkeypatch):
    raw, amounts, warning = source_case(kind, monkeypatch)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-oct28-late-" + kind}
        response = client.post("/convert", data=owner, files={"file": (kind + ".pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        assert [row["amount"] for row in analysis["preview_transactions"]] == amounts
        if warning:
            assert warning in analysis["pdf_processing_metrics"]["canonical_warning_types"]
        identifier = analysis["analysis_id"]
        assert client.get(f"/report/{identifier}", params=owner).status_code == 200
        for fmt in ("csv", "xlsx", "ofx"):
            report = client.get(f"/convert-report/{identifier}", params={**owner, "format": fmt})
            assert report.status_code == 200, report.text
            if fmt == "csv":
                rows = csv.DictReader(line for line in StringIO(report.content.decode("utf-8-sig")) if line.strip())
                assert [float(row["amount"]) for row in rows] == amounts
            elif fmt == "ofx":
                assert [row.amount for row in parse_ofx_transactions(report.content)] == amounts
            else:
                book = load_workbook(BytesIO(report.content), read_only=True, data_only=True)
                rows = list(book["Conversao"].iter_rows(min_row=2, values_only=True))
                assert [round(float(row[2] or 0) - float(row[3] or 0), 2) for row in rows] == amounts
                book.close()
        assert client.get("/report/an_oct28_late_nonexistent", params=owner).status_code == 404


@pytest.mark.parametrize("kind", ["sicoob", "bb_stroke"])
def test_worker_preserves_recovered_values_and_review_over_http(kind, tmp_path, monkeypatch):
    raw, amounts, warning = source_case(kind, monkeypatch)
    with worker_case(tmp_path, raw) as (client, job):
        analysis = job["result"]["analysis"]
        assert [row["amount"] for row in analysis["preview_transactions"]] == amounts
        assert warning in analysis["pdf_processing_metrics"]["canonical_warning_types"]
        report = client.get(f"/convert-report/{analysis['analysis_id']}", params={"format": "ofx"})
        assert report.status_code == 200
        assert [row.amount for row in parse_ofx_transactions(report.content)] == amounts
