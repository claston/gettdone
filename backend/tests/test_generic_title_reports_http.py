"""Real uploads, worker processing and safe exports of generic title reports."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.fixtures.generic_titles import title_ocr, title_pdf
from tests.ocr_conversion_http_helpers import configure_ocr, public_client, worker_case

pytestmark = pytest.mark.real_http


@pytest.mark.parametrize("kind", ["payable", "receivable"])
@pytest.mark.parametrize("scanned", [False, True])
def test_generic_reports_public_http_exports_and_ofx_policy(kind, scanned, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    raw = configure_ocr(monkeypatch, title_ocr(kind, reordered=True)) if scanned else title_pdf(kind, reordered=True)
    with public_client(tmp_path) as client:
        owner = {"anonymous_fingerprint": "fp-generic-titles"}
        response = client.post("/convert", data=owner, files={"file": ("report.pdf", raw, "application/pdf")})
        assert response.status_code == 200, response.text
        analysis = response.json()["analysis"]
        report = analysis["financial_document"]
        assert report["export"]["available_formats"] == ["csv", "xlsx"]
        assert analysis["total_inflows"] == analysis["total_outflows"] == 0
        assert analysis["transactions_total"] == 2
        assert all(e["cash_effect"] is None for e in report["events"])
        aid = analysis["analysis_id"]
        assert client.get(f"/report/{aid}", params=owner).status_code == 200
        exported = client.get(f"/convert-report/{aid}", params={**owner, "format": "csv"})
        assert exported.status_code == 200 and "0001/01" in exported.text
        assert "amounts.outstanding" in exported.text
        exported = client.get(f"/convert-report/{aid}", params={**owner, "format": "xlsx"})
        assert exported.status_code == 200
        book = load_workbook(BytesIO(exported.content))
        assert book["Eventos"].max_row == book["Campos originais"].max_row == 3
        rejected = client.get(f"/convert-report/{aid}", params={**owner, "format": "ofx", "bank_code": "237",
                              "bank_branch": "1234", "account_number": "123456"})
        assert rejected.status_code == 400 and "OFX" in rejected.text
        assert client.get("/report/an_nonexistent", params=owner).status_code == 404


@pytest.mark.parametrize("kind", ["payable", "receivable"])
def test_worker_uses_generic_title_model_and_excel_download(kind, tmp_path, monkeypatch):
    monkeypatch.setenv("TEXTRACT_ENABLED", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "false")
    with worker_case(tmp_path, title_pdf(kind, reordered=True)) as (client, job):
        analysis = job["result"]["analysis"]
        assert analysis["document_type"] == ("accounts_payable_report" if kind == "payable" else "accounts_receivable_report")
        assert analysis["transactions_total"] == 2
        assert analysis["net_total"] == 0
        aid = analysis["analysis_id"]
        for fmt in ("csv", "xlsx"):
            assert client.get(f"/convert-report/{aid}", params={"format": fmt}).status_code == 200
        assert client.get(f"/convert-report/{aid}", params={"format": "ofx"}).status_code == 400
